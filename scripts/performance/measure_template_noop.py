"""Compare an installed baseline with a source tree using the same native build.

Example:
  python scripts/performance/measure_template_noop.py \
    --fixture /tmp/plain-200000.xlsx --baseline-python /tmp/baseline/bin/python \
    --optimized-pythonpath ./python --output /tmp/template-noop.json

The fixture is copied outside operation timing; independent verification is
also reported separately. Both variants run in alternating fresh processes.
"""

from __future__ import annotations

import argparse
import gc
import hashlib
import json
import os
from pathlib import Path
import platform
import shutil
import statistics
import subprocess
import sys
import tempfile
import time


def measure(fixture: Path) -> dict[str, object]:
    import openpyxl
    import wolfxl

    with tempfile.TemporaryDirectory() as temporary:
        destination = Path(temporary) / "edit.xlsx"
        start = time.perf_counter()
        shutil.copyfile(fixture, destination)
        copy_s = time.perf_counter() - start
        gc.collect()
        gc.disable()
        try:
            start = time.perf_counter()
            workbook = wolfxl.load_workbook(destination, modify=True)
            load_s = time.perf_counter() - start
            start = time.perf_counter()
            workbook.active["B2"] = "changed"
            workbook.active["C3"] = 12345
            edit_s = time.perf_counter() - start
            start = time.perf_counter()
            workbook.save(destination)
            save_s = time.perf_counter() - start
            workbook.close()
        finally:
            gc.enable()
        start = time.perf_counter()
        check = openpyxl.load_workbook(destination, read_only=True, data_only=False)
        try:
            values = list(check.active.iter_rows(min_row=2, max_row=3, min_col=2, max_col=3, values_only=True))
            assert values[0][0] == "changed" and values[1][1] == 12345
        finally:
            check.close()
        verify_s = time.perf_counter() - start
        return dict(copy_s=copy_s, load_s=load_s, edit_s=edit_s, save_s=save_s,
                    operation_s=load_s + edit_s + save_s, verify_s=verify_s, values=values)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fixture", required=True, type=Path)
    parser.add_argument("--baseline-python", default=sys.executable)
    parser.add_argument("--optimized-pythonpath", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--samples", type=int, default=5)
    parser.add_argument("--worker", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.worker:
        print(json.dumps(measure(args.fixture)))
        return
    if args.optimized_pythonpath is None or args.output is None:
        parser.error("--optimized-pythonpath and --output are required")
    records: dict[str, list[dict[str, object]]] = {"baseline": [], "optimized": []}
    for index in range(args.samples + 1):
        order = ("baseline", "optimized") if index % 2 == 0 else ("optimized", "baseline")
        for name in order:
            environment = os.environ.copy()
            environment.pop("PYTHONPATH", None)
            if name == "optimized":
                environment["PYTHONPATH"] = str(args.optimized_pythonpath.resolve())
            process = subprocess.run(
                [args.baseline_python, str(Path(__file__).resolve()), "--worker", "--fixture", str(args.fixture.resolve())],
                env=environment, text=True, check=True, capture_output=True,
            )
            if index:
                records[name].append(json.loads(process.stdout))
    keys = ("copy_s", "load_s", "edit_s", "save_s", "operation_s", "verify_s")
    medians = {name: {key: statistics.median(sample[key] for sample in samples) for key in keys}
               for name, samples in records.items()}
    receipt = dict(python=sys.version, platform=platform.platform(),
                   fixture_sha256=hashlib.sha256(args.fixture.read_bytes()).hexdigest(),
                   measurement=dict(warmups=1, samples=args.samples, fresh_process_per_sample=True,
                                    alternating_order=True, verification_outside_operation=True),
                   samples=records, medians=medians,
                   speedup={key: medians["baseline"][key] / medians["optimized"][key]
                            for key in ("save_s", "operation_s")})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(dict(medians=medians, speedup=receipt["speedup"]), indent=2))


if __name__ == "__main__":
    main()
