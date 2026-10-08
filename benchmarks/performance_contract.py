"""Run labelled v2 styled/edit benchmarks, or compare matched cumulative receipts.

PYTHONPATH=python python benchmarks/performance_contract.py --label baseline \
    --fixture-dir /tmp/wolfxl-perf-fixtures --output /tmp/baseline.json
python benchmarks/performance_contract.py --compare /tmp/baseline.json /tmp/final.json \
    --output /tmp/cumulative.json

Original benchmark files and committed receipts remain unchanged. Engine-operation,
bounded verification, end-to-end total, and full validation are separate metrics.
"""
from __future__ import annotations

import argparse
import gc
import hashlib
import importlib
import json
import os
import platform
import shutil
import statistics
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter

import openpyxl

from performance_fixtures import (CONTRACT, edits_for, fixture_kind, modification_cases,
                                  prepare_fixture, sha256)
from performance_validation import (bounded_verify, comparison, expected_digest,
                                    full_verify, package_verify)

STYLED_KINDS = ("small", "large", "sparse", "high_cardinality", "merged")
ALL_CASES = tuple(f"edit_{case}" for case in modification_cases()) + tuple(
    f"styled_{kind}_{mode}" for kind in STYLED_KINDS for mode in ("eager", "read_only"))


def timed(fn):
    started = perf_counter()
    result = fn()
    return perf_counter() - started, result


def phase_summary(samples: list[dict]) -> dict:
    phases = sorted(samples[0]["phase_seconds"])
    return {"raw_samples": samples,
            "phase_medians_seconds": {phase: statistics.median(
                sample["phase_seconds"][phase] for sample in samples) for phase in phases},
            "phase_ranges_seconds": {phase: [min(sample["phase_seconds"][phase] for sample in samples),
                                             max(sample["phase_seconds"][phase] for sample in samples)]
                                      for phase in phases}}


def measure_edit(module, engine: str, fixture: Path, output: Path, edits: list[tuple]) -> dict:
    phases = {}
    total_started = perf_counter()
    phases["copy"], _ = timed(lambda: shutil.copy2(fixture, output))
    operation_started = perf_counter()
    kwargs = {"modify": True} if engine == "wolfxl" else {}
    phases["load"], wb = timed(lambda: module.load_workbook(str(output), data_only=False, **kwargs))
    try:
        def assign():
            for sheet, row, col, value in edits:
                wb[sheet].cell(row, col, value)
        phases["assignment"], _ = timed(assign)
        phases["save"], _ = timed(lambda: wb.save(str(output)))
    finally:
        phases["close"], _ = timed(wb.close)
    phases["operation"] = perf_counter() - operation_started
    phases["bounded_verification"], verified = timed(lambda: bounded_verify(output, edits))
    phases["total"] = perf_counter() - total_started
    return {"phase_seconds": phases, "verification": verified}


def styled_signature(module, path: Path, *, read_only: bool) -> tuple[dict, dict]:
    phases = {}
    total_started = perf_counter()
    phases["load"], wb = timed(lambda: module.load_workbook(
        str(path), read_only=read_only, data_only=True))
    try:
        def scan():
            rows, checksum, styled_cells = 0, 0.0, 0
            for row in wb.active.iter_rows():
                rows += 1
                first = row[0].value if row else None
                if type(first) in (int, float):
                    checksum += float(first)
                for cell in row:
                    # Identical public Cell/font/fill/number-format loop for both engines.
                    fill = getattr(cell.fill, "fill_type", None) or getattr(cell.fill, "patternType", None)
                    if getattr(cell.font, "bold", False) or fill or (cell.number_format or "General") != "General":
                        styled_cells += 1
            return {"rows": rows, "first_column_checksum": checksum, "styled_cells": styled_cells}
        phases["iterate_styled_cells"], signature = timed(scan)
    finally:
        phases["close"], _ = timed(wb.close)
    phases["operation"] = phases["total"] = perf_counter() - total_started
    return phases, signature


def measure_styled(module, fixture: Path, read_only: bool, expected: dict) -> dict:
    phases, signature = styled_signature(module, fixture, read_only=read_only)
    if signature != expected:
        raise AssertionError(f"Styled-read signature differs: {signature} != {expected}")
    return {"phase_seconds": phases, "signature": signature}


def package_fingerprint(module) -> dict:
    directory = Path(module.__file__).resolve().parent
    source = hashlib.sha256()
    count = 0
    for path in sorted(directory.rglob("*.py")):
        source.update(str(path.relative_to(directory)).encode())
        source.update(path.read_bytes())
        count += 1
    return {"version": module.__version__, "import_path": str(module.__file__),
            "python_source_sha256": source.hexdigest(), "python_source_files": count}


def source_state(root: Path) -> dict:
    def git(*args):
        return subprocess.check_output(["git", "-C", str(root), *args], text=True).strip()
    try:
        return {"root": str(root.resolve()), "commit": git("rev-parse", "HEAD"),
                "dirty": bool(git("status", "--porcelain"))}
    except (OSError, subprocess.CalledProcessError):
        return {"root": str(root.resolve()), "commit": None, "dirty": None}


def metadata(args, wolfxl) -> dict:
    native = importlib.import_module("wolfxl._rust")
    native_path = Path(native.__file__).resolve()
    scripts = ("performance_contract.py", "performance_fixtures.py", "performance_validation.py")
    digest = hashlib.sha256()
    for name in scripts:
        digest.update(name.encode())
        digest.update(Path(__file__).with_name(name).read_bytes())
    return {"contract": CONTRACT, "label": args.label,
            "created_utc": datetime.now(timezone.utc).isoformat(),
            "harness_sha256": digest.hexdigest(),
            "environment": {"python": sys.version, "platform": platform.platform(),
                            "machine": platform.machine(), "processor": platform.processor(),
                            "cpu_count": os.cpu_count(), "executable": sys.executable,
                            "gc_during_operation": "disabled",
                            "engine_order": "alternates each round and warmup"},
            "packages": {"wolfxl": package_fingerprint(wolfxl),
                         "openpyxl": package_fingerprint(openpyxl)},
            "native": {"path": str(native_path), "sha256": sha256(native_path)},
            "source": source_state(args.source_root),
            "timing_contract": {
                "operation": "load + assignment + save + close (edit); load + Cell/style loop + close (read)",
                "total": "copy + operation + same bounded openpyxl verification (edit); operation (read)",
                "deep_validation": "full eager openpyxl reopen, all stored values/styles/merges and package checks; outside total",
                "bounded_read_note": "one iterator per edited sheet; late rows still require streaming the XML prefix",
                "legacy_receipts": "unchanged; this v2 contract is not retroactively comparable"}}


def workload_spec(case: str, args) -> tuple[str, int, int, bool]:
    if case.startswith("edit_"):
        return fixture_kind(case[5:]), args.edit_rows, 0, False
    kind = next(kind for kind in STYLED_KINDS if case.startswith(f"styled_{kind}_"))
    rows = args.small_rows if kind == "small" else args.styled_rows
    fixture = {"sparse": "styled_sparse", "merged": "merged"}.get(kind, "styled")
    return fixture, rows, args.style_cardinality if kind == "high_cardinality" else 0, case.endswith("read_only")


def run(args) -> dict:
    wolfxl = importlib.import_module("wolfxl")
    receipt = metadata(args, wolfxl)
    receipt["status"] = "in_progress"
    receipt["workload_config"] = {"edit_rows": args.edit_rows, "small_rows": args.small_rows,
                                  "styled_rows": args.styled_rows, "style_cardinality": args.style_cardinality,
                                  "rounds": args.rounds, "warmups": args.warmups,
                                  "cases": args.cases, "engines": args.engines}
    receipt["fixtures"], receipt["results"] = {}, []
    modules = {"openpyxl": openpyxl, "wolfxl": wolfxl}
    with tempfile.TemporaryDirectory(prefix="wolfxl-perf-v2-") as temporary:
        for case in args.cases:
            kind, rows, cardinality, read_only = workload_spec(case, args)
            fixture, manifest = prepare_fixture(args.fixture_dir, rows=rows, kind=kind, cardinality=cardinality)
            receipt["fixtures"][case] = manifest
            edits = edits_for(case[5:], rows) if case.startswith("edit_") else []
            reference_seconds, expected = timed(lambda: expected_digest(fixture, edits)) if edits else timed(
                lambda: styled_signature(openpyxl, fixture, read_only=read_only)[1])
            samples = {engine: [] for engine in args.engines}
            warmups = {engine: [] for engine in args.engines}
            outputs = {engine: Path(temporary) / f"{case}-{engine}.xlsx" for engine in args.engines}
            for round_index in range(args.warmups + args.rounds):
                order = args.engines if round_index % 2 == 0 else list(reversed(args.engines))
                for order_index, engine in enumerate(order):
                    gc.collect()
                    enabled = gc.isenabled()
                    gc.disable()
                    try:
                        sample = measure_edit(modules[engine], engine, fixture, outputs[engine], edits) if edits else measure_styled(
                            modules[engine], fixture, read_only, expected)
                    finally:
                        if enabled:
                            gc.enable()
                    if round_index >= args.warmups:
                        sample.update({"round": round_index - args.warmups, "order_index": order_index})
                        samples[engine].append(sample)
                    else:
                        sample.update({"warmup_round": round_index, "order_index": order_index})
                        warmups[engine].append(sample)
            for engine in args.engines:
                result = {"case": case, "engine": engine, **phase_summary(samples[engine]),
                          "warmup_samples": warmups[engine],
                          "reference_validation_seconds": reference_seconds}
                if edits:
                    validation_seconds, semantic = timed(lambda: full_verify(outputs[engine], expected))
                    sheet_indexes = {1 if name == "Data" else 2 for name, *_ in edits}
                    package_seconds, package = timed(lambda: package_verify(fixture, outputs[engine], sheet_indexes, engine))
                    result["deep_validation"] = {"full_reopen_seconds": validation_seconds,
                                                  "package_seconds": package_seconds, **semantic, **package}
                receipt["results"].append(result)
            print(f"Completed {case}", flush=True)
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n")
    receipt["status"] = "complete"
    args.output.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n")
    return receipt


def verifier_comparison(args) -> dict:
    """Attribute the verifier change explicitly, without claiming an engine speedup."""
    wolfxl = importlib.import_module("wolfxl")
    receipt = metadata(args, wolfxl)
    fixture, manifest = prepare_fixture(args.fixture_dir, rows=args.edit_rows, kind="plain")
    edits = edits_for("top", args.edit_rows)
    samples = {"legacy_eager": [], "bounded_one_pass": []}
    def eager(path):
        wb = openpyxl.load_workbook(path, data_only=False)
        try:
            for sheet, row, col, expected in edits:
                if wb[sheet].cell(row, col).value != expected:
                    raise AssertionError("Legacy verifier mismatch")
        finally:
            wb.close()
    with tempfile.TemporaryDirectory(prefix="wolfxl-verifier-v2-") as temporary:
        output = Path(temporary) / "modified.xlsx"
        measure_edit(wolfxl, "wolfxl", fixture, output, edits)
        verifiers = {"legacy_eager": lambda: eager(output),
                     "bounded_one_pass": lambda: bounded_verify(output, edits)}
        for index in range(args.warmups + args.rounds):
            order = list(verifiers) if index % 2 == 0 else list(reversed(verifiers))
            for name in order:
                gc.collect()
                seconds, _ = timed(verifiers[name])
                if index >= args.warmups:
                    samples[name].append(seconds)
    medians = {name: statistics.median(values) for name, values in samples.items()}
    receipt.update({"fixture": manifest, "rounds": args.rounds, "warmups": args.warmups,
                    "status": "complete",
                    "verifier_comparison": {"raw_samples_seconds": samples, "medians_seconds": medians,
                                            "verification_speedup": medians["legacy_eager"] / medians["bounded_one_pass"]},
                    "claim_boundary": "Same two assertions; bounded verification gain only, not an engine or historical-total speedup"})
    return receipt


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--label", default="unlabelled")
    parser.add_argument("--fixture-dir", type=Path, default=Path(".perf-fixtures-v2"))
    parser.add_argument("--source-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--edit-rows", type=int, default=200_000)
    parser.add_argument("--styled-rows", type=int, default=20_000)
    parser.add_argument("--small-rows", type=int, default=2_000)
    parser.add_argument("--style-cardinality", type=int, default=1_024)
    parser.add_argument("--rounds", type=int, default=5)
    parser.add_argument("--warmups", type=int, default=1)
    parser.add_argument("--cases", default=",".join(ALL_CASES))
    parser.add_argument("--engines", default="openpyxl,wolfxl")
    parser.add_argument("--compare", type=Path, nargs=2)
    parser.add_argument("--verifier-comparison", action="store_true")
    args = parser.parse_args(argv)
    args.cases = args.cases.split(",")
    args.engines = args.engines.split(",")
    if not 4 <= args.edit_rows <= 200_000 or not 4 <= args.styled_rows <= 200_000 or not 4 <= args.small_rows <= 200_000:
        parser.error("row counts must be between 4 and 200000")
    if not 1 <= args.style_cardinality <= 4096 or not 1 <= args.rounds <= 20 or not 0 <= args.warmups <= 5:
        parser.error("styles must be 1..4096, rounds 1..20, warmups 0..5")
    if not args.cases or set(args.cases) - set(ALL_CASES) or len(set(args.cases)) != len(args.cases):
        parser.error(f"cases must be a unique subset of {ALL_CASES}")
    if not args.engines or set(args.engines) - {"openpyxl", "wolfxl"} or len(set(args.engines)) != len(args.engines):
        parser.error("engines must be a unique subset of openpyxl,wolfxl")
    if args.compare and args.verifier_comparison:
        parser.error("choose --compare or --verifier-comparison")
    return args


def main():
    args = parse_args()
    if args.compare:
        receipt = comparison(*(json.loads(path.read_text()) for path in args.compare))
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n")
    elif args.verifier_comparison:
        receipt = verifier_comparison(args)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n")
    else:
        run(args)


if __name__ == "__main__":
    main()
