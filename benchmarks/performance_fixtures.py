"""Deterministic, bounded fixtures for the separately labelled perf contract v2."""
from __future__ import annotations

import hashlib
import json
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side

CONTRACT = "wolfxl-styled-edit-v2"
FIXTURE_VERSION = 2


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def plain_row(index: int) -> list:
    # The historic nominal eight-column fixture actually populated five columns.
    return [index, f"customer-{index % 997}", index * 1.25, index % 17,
            f"region-{index % 4}"]


def style_cell(cell, row: int, col: int, cardinality: int) -> None:
    if col == 1:
        cell.font = Font(bold=True)
    if col == 2 and row % 2 == 0:
        cell.fill = PatternFill(patternType="solid", fgColor="FFD966")
    if col == 3:
        cell.number_format = "#,##0.00"
    if cardinality > 0:
        index = ((row - 1) * 5 + col - 1) % cardinality
        cell.number_format = f'0.00 "style-{index}"'


def _normalize_zip(path: Path) -> None:
    rewritten = path.with_suffix(".normalized.xlsx")
    with ZipFile(path) as source, ZipFile(rewritten, "w") as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            info.create_system = 3
            data = source.read(name)
            if name == "docProps/core.xml":
                for prefix, namespace in (("cp", "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"),
                                          ("dc", "http://purl.org/dc/elements/1.1/"),
                                          ("dcterms", "http://purl.org/dc/terms/"),
                                          ("xsi", "http://www.w3.org/2001/XMLSchema-instance")):
                    ElementTree.register_namespace(prefix, namespace)
                core = ElementTree.fromstring(data)
                for field in ("created", "modified"):
                    element = core.find(f"{{http://purl.org/dc/terms/}}{field}")
                    if element is not None:
                        element.text = "2000-01-01T00:00:00Z"
                data = ElementTree.tostring(core, encoding="utf-8")
            target.writestr(info, data)
    rewritten.replace(path)


def prepare_fixture(directory: Path, *, rows: int, kind: str,
                    cardinality: int = 0) -> tuple[Path, dict]:
    config = {"fixture_version": FIXTURE_VERSION, "rows": rows, "kind": kind,
              "style_cardinality": cardinality, "builder_openpyxl": openpyxl.__version__}
    key = hashlib.sha256(json.dumps(config, sort_keys=True).encode()).hexdigest()[:20]
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{kind}-{key}.xlsx"
    manifest = path.with_suffix(".json")
    if path.exists() or manifest.exists():
        if not path.exists() or not manifest.exists():
            raise ValueError(f"Incomplete cached fixture: {path}")
        metadata = json.loads(manifest.read_text())
        if metadata["config"] != config or metadata["sha256"] != sha256(path):
            raise ValueError(f"Cached fixture changed: {path}")
        return path, metadata
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    ws = wb.active
    ws.title = "Data"
    styled = kind.startswith("styled") or kind == "merged"
    sparse = kind == "styled_sparse"
    populated = 0
    for row in range(1, rows + 1):
        if sparse and (row - 1) % 16:
            continue
        for col, value in enumerate(plain_row(row), start=1):
            cell = ws.cell(row, col, value)
            if styled:
                style_cell(cell, row, col, cardinality)
            populated += 1
    if sparse:
        cell = ws.cell(rows, 32, rows)
        style_cell(cell, rows, 3, cardinality)
        populated += 1
    if kind == "formula":
        ws["C3"] = "=A3*2"
        ws["D4"] = "=SUM(A1:A3)"
    if kind == "cross_sheet":
        extra = wb.create_sheet("Other")
        for row in range(1, rows + 1):
            extra.append(plain_row(row))
        populated *= 2
    if kind == "merged":
        ws.cell(rows, 5).border = Border(right=Side(style="thin"), bottom=Side(style="thin"))
        ws.merge_cells(start_row=rows - 1, end_row=rows, start_column=4, end_column=5)
        populated -= 3
    wb.save(path)
    wb.close()
    _normalize_zip(path)
    metadata = {"config": config, "sha256": sha256(path), "bytes": path.stat().st_size,
                "rows_per_sheet": rows, "requested_plain_columns": 8,
                "populated_plain_columns": 5, "max_column": 32 if sparse else 5,
                "populated_value_cells": populated,
                "sheet_count": 2 if kind == "cross_sheet" else 1}
    manifest.write_text(json.dumps(metadata, indent=2, sort_keys=True) + "\n")
    return path, metadata


def modification_cases() -> tuple[str, ...]:
    return ("top", "middle", "bottom", "missing", "cross_sheet", "formula", "styled", "merged")


def edits_for(case: str, rows: int) -> list[tuple[str, int, int, object]]:
    if case == "middle":
        row = max(2, rows // 2)
        return [("Data", row, 2, "changed"), ("Data", row + 1, 3, 12345)]
    if case == "bottom":
        return [("Data", rows - 1, 2, "changed"), ("Data", rows, 3, 12345)]
    if case == "missing":
        return [("Data", rows + 1, 2, "changed"), ("Data", rows + 2, 3, 12345)]
    if case == "cross_sheet":
        return [("Data", 2, 2, "changed"), ("Other", rows, 3, 12345)]
    return [("Data", 2, 2, "changed"), ("Data", 3, 3, 12345)]


def fixture_kind(case: str) -> str:
    return case if case in ("cross_sheet", "formula", "merged") else "styled" if case == "styled" else "plain"
