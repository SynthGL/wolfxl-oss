"""Independent verification. None of these checks are engine-operation timing."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
from openpyxl.xml.functions import tostring


def bounded_verify(path: Path, edits: list[tuple]) -> dict:
    """One bounded values-only iterator per edited sheet, identical for both engines."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    windows = []
    try:
        for sheet in sorted({edit[0] for edit in edits}):
            expected = {(row, col): value for name, row, col, value in edits if name == sheet}
            rows = [coord[0] for coord in expected]
            cols = [coord[1] for coord in expected]
            bounds = {"min_row": min(rows), "max_row": max(rows),
                      "min_col": min(cols), "max_col": max(cols)}
            seen = set()
            for row_index, values in enumerate(wb[sheet].iter_rows(values_only=True, **bounds),
                                               start=bounds["min_row"]):
                for col_index, value in enumerate(values, start=bounds["min_col"]):
                    coord = row_index, col_index
                    if coord in expected:
                        if value != expected[coord]:
                            raise AssertionError(f"{sheet}!{coord}: {value!r} != {expected[coord]!r}")
                        seen.add(coord)
            if seen != expected.keys():
                raise AssertionError(f"Missing verification cells on {sheet}")
            windows.append({"sheet": sheet, **bounds})
    finally:
        wb.close()
    return {"verified_cells": len(edits), "windows": windows,
            "verifier": "openpyxl-read-only-one-pass-per-sheet"}


def _style_xml(component) -> str:
    # Style IDs can legitimately be renumbered by a writer; compare their meaning.
    return tostring(component.to_tree()).decode()


def semantic_digest(wb) -> str:
    digest = hashlib.sha256()
    styles = {}
    for ws in wb.worksheets:
        header = (ws.title, ws.max_row, ws.max_column,
                  sorted(str(item) for item in ws.merged_cells.ranges))
        digest.update(json.dumps(header).encode())
        # _cells avoids manufacturing empty cells during this independent eager check.
        for coord, cell in sorted(ws._cells.items()):
            if cell.value is None and not cell.has_style:
                continue
            style_id = cell.style_id
            if style_id not in styles:
                meaning = (cell.number_format, _style_xml(cell.font), _style_xml(cell.fill),
                           _style_xml(cell.border), _style_xml(cell.alignment),
                           _style_xml(cell.protection))
                styles[style_id] = hashlib.sha256(json.dumps(meaning).encode()).hexdigest()
            record = (coord, cell.value, cell.data_type, styles[style_id])
            digest.update(json.dumps(record, default=str, separators=(",", ":")).encode())
    return digest.hexdigest()


def expected_digest(path: Path, edits: list[tuple]) -> str:
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        for sheet, row, col, value in edits:
            wb[sheet].cell(row, col, value)
        return semantic_digest(wb)
    finally:
        wb.close()


def full_verify(path: Path, expected: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        actual = semantic_digest(wb)
    finally:
        wb.close()
    if actual != expected:
        raise AssertionError(f"Full value/style/merge signature differs: {actual} != {expected}")
    return {"semantic_sha256": actual, "full_reopen": "openpyxl-eager"}


def package_verify(source: Path, output: Path, edited_sheets: set[int], engine: str) -> dict:
    """Validate ZIP CRCs and all XML. WolfXL must retain untouched package parts."""
    with ZipFile(source) as before, ZipFile(output) as after:
        if after.testzip() is not None:
            raise AssertionError("Output ZIP CRC failure")
        names_before, names_after = set(before.namelist()), set(after.namelist())
        for name in names_after:
            if name.endswith((".xml", ".rels")):
                with after.open(name) as stream:
                    events = ElementTree.iterparse(stream, events=("start", "end"))
                    _, root = next(events)
                    for event, element in events:
                        if event == "end":
                            element.clear()
                            root.clear()
        changed = sorted(name for name in names_before & names_after
                         if before.read(name) != after.read(name))
        missing, added = sorted(names_before - names_after), sorted(names_after - names_before)
        if engine == "wolfxl":
            allowed = {f"xl/worksheets/sheet{index}.xml" for index in edited_sheets}
            # Workbook/calc metadata may change when edits invalidate formula caches.
            allowed |= {"xl/workbook.xml", "xl/calcChain.xml", "[Content_Types].xml",
                        "xl/_rels/workbook.xml.rels"}
            forbidden = (set(changed) | set(missing) | set(added)) - allowed
            if forbidden:
                raise AssertionError(f"Untouched package parts changed: {sorted(forbidden)}")
    return {"zip_crc": "pass", "xml_parse": "pass", "changed_parts": changed,
            "removed_parts": missing, "added_parts": added,
            "untouched_part_byte_check": "pass" if engine == "wolfxl" else "reported-full-rewrite"}


def comparison(before: dict, after: dict) -> dict:
    """Compare matched receipts directly; never multiply isolated optimization gains."""
    if before.get("status") != "complete" or after.get("status") != "complete":
        raise ValueError("Comparison requires two complete receipts")
    for key in ("contract", "workload_config", "fixtures", "harness_sha256"):
        if before[key] != after[key]:
            raise ValueError(f"Receipts have different {key}")
    for key in ("python", "platform", "machine", "processor", "cpu_count"):
        if before["environment"][key] != after["environment"][key]:
            raise ValueError(f"Receipts have different environment.{key}")
    for key in ("version", "python_source_sha256"):
        if before["packages"]["openpyxl"][key] != after["packages"]["openpyxl"][key]:
            raise ValueError(f"Receipts have different openpyxl {key}")
    old = {(item["case"], item["engine"]): item for item in before["results"]}
    new = {(item["case"], item["engine"]): item for item in after["results"]}
    if old.keys() != new.keys():
        raise ValueError("Receipts have different case/engine coverage")
    expected_cases = set(before["workload_config"]["cases"])
    if {case for case, _ in old} != expected_cases or len(old) != len(expected_cases) * 2:
        raise ValueError("Comparison requires complete declared case/engine coverage")
    if {engine for _, engine in new} != {"openpyxl", "wolfxl"}:
        raise ValueError("Comparison requires both openpyxl and wolfxl")
    rows = []
    for case, engine in sorted(new):
        if engine != "wolfxl":
            continue
        baseline = old[case, engine]["phase_medians_seconds"]
        final = new[case, engine]["phase_medians_seconds"]
        openpyxl_final = new[case, "openpyxl"]["phase_medians_seconds"]
        rows.append({"case": case, "wolfxl_before_seconds": baseline["operation"],
                     "wolfxl_after_seconds": final["operation"],
                     "cumulative_engine_speedup": baseline["operation"] / final["operation"],
                     "final_speedup_vs_openpyxl": openpyxl_final["operation"] / final["operation"],
                     "before_total_seconds": baseline["total"],
                     "after_total_seconds": final["total"],
                     "cumulative_total_speedup": baseline["total"] / final["total"]})
    return {"contract": before["contract"], "baseline_label": before["label"],
            "final_label": after["label"], "comparisons": rows,
            "method": "matched baseline/final medians; no multiplication of isolated gains"}
