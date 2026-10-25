"""Large streamed sheets inflate from the ZIP entry without a temp-file spool."""

from __future__ import annotations

import io
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import openpyxl
import pytest

import wolfxl
from wolfxl import _rust

SHEET_PART = "xl/worksheets/sheet1.xml"
ROWS = 20_000
EXPECTED = [(row, f"text {row}") for row in range(1, ROWS + 1)]


def _sheet_xml() -> bytes:
    rows = "".join(
        f'<row r="{row}"><c r="A{row}"><v>{row}</v></c>'
        f'<c r="B{row}" t="inlineStr"><is><t>text {row}</t></is></c></row>'
        for row in range(1, ROWS + 1)
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<dimension ref="A1:B{ROWS}"/><sheetData>{rows}</sheetData></worksheet>'
    ).encode()


def _package(sheet_compression: int) -> bytes:
    template = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(template)
    sheet = _sheet_xml()
    # Above the reader's 1 MiB in-memory limit, so the entry-streaming path runs.
    assert len(sheet) > 1024 * 1024
    result = io.BytesIO()
    with ZipFile(template) as source, ZipFile(result, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name == SHEET_PART:
                target.writestr(name, sheet, compress_type=sheet_compression)
            else:
                target.writestr(name, source.read(name))
    return result.getvalue()


def _open(kind: str, data: bytes, tmp_path: Path):
    if kind == "path":
        path = tmp_path / "large.xlsx"
        path.write_bytes(data)
        return _rust.StreamingSheetReader.open(str(path), "Sheet")
    return _rust.StreamingSheetReader.open_from_bytes(data, "Sheet")


def _drain(reader) -> None:
    while reader.read_next_values() is not None:
        pass


@pytest.fixture
def private_tmpdir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    spool = tmp_path / "private-tmp"
    spool.mkdir()
    for name in ("TMPDIR", "TMP", "TEMP"):
        monkeypatch.setenv(name, str(spool))
    monkeypatch.setattr(tempfile, "tempdir", str(spool))
    return spool


@pytest.mark.parametrize("source_kind", ["path", "fileobj"])
def test_large_sheet_streams_without_temp_file(
    tmp_path: Path, private_tmpdir: Path, source_kind: str
) -> None:
    data = _package(ZIP_DEFLATED)
    if source_kind == "path":
        source = tmp_path / "large.xlsx"
        source.write_bytes(data)
    else:
        source = io.BytesIO(data)

    wb = wolfxl.load_workbook(source, read_only=True)
    rows = wb.active.iter_rows(values_only=True)
    first = next(rows)
    assert list(private_tmpdir.iterdir()) == []
    seen = [first, *rows]
    wb.close()

    assert seen == EXPECTED
    assert list(private_tmpdir.iterdir()) == []


@pytest.mark.parametrize("source_kind", ["path", "fileobj"])
def test_large_sheet_checksum_is_verified_while_streaming(
    tmp_path: Path, source_kind: str
) -> None:
    data = _package(ZIP_STORED)
    assert data.count(b"<v>12345</v>") == 1
    corrupt = data.replace(b"<v>12345</v>", b"<v>12346</v>")

    reader = _open(source_kind, corrupt, tmp_path)
    with pytest.raises(OSError, match="Invalid checksum"):
        _drain(reader)


@pytest.mark.parametrize("source_kind", ["path", "fileobj"])
@pytest.mark.parametrize(
    ("limit", "value", "message"),
    [
        ("WOLFXL_MAX_ZIP_ENTRY_BYTES", str(1024 * 1024), "too large"),
        ("WOLFXL_MAX_ZIP_COMPRESSION_RATIO", "2", "compression ratio"),
    ],
)
def test_large_sheet_keeps_zip_limits(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    source_kind: str,
    limit: str,
    value: str,
    message: str,
) -> None:
    data = _package(ZIP_DEFLATED)
    _drain(_open(source_kind, data, tmp_path))

    monkeypatch.setenv(limit, value)
    with pytest.raises(OSError, match=message):
        _open(source_kind, data, tmp_path)
