"""pandas' openpyxl engine under ``wolfxl.install_as_openpyxl()``.

pandas gates its ``openpyxl`` engine on the module's ``__version__``. The
alias must report the openpyxl API release it provides to pandas' probe, or
``read_excel`` and ``ExcelWriter`` raise ImportError. Every scenario runs in a
fresh interpreter so import order matches a real one-line migration, and the
workbooks the alias writes are graded by upstream openpyxl.
"""

from __future__ import annotations

import json
import subprocess
import sys
import textwrap
from pathlib import Path

import pytest

pytest.importorskip("pandas")
pytest.importorskip("openpyxl")
pytest.importorskip("wolfxl._rust")

ALIAS_FIRST = """
import wolfxl
wolfxl.install_as_openpyxl()
import pandas as pd
"""

PANDAS_FIRST = """
import pandas as pd
import wolfxl
wolfxl.install_as_openpyxl()
"""

UPSTREAM = """
import pandas as pd
import openpyxl
import wolfxl
assert openpyxl is not wolfxl
"""

FRAME = """
import datetime
import json
frame = pd.DataFrame(
    {
        "name": ["alpha", "beta", "gamma"],
        "count": [1, 2, 3],
        "ratio": [0.5, 1.25, None],
        "when": [datetime.datetime(2024, 1, 1), datetime.datetime(2024, 2, 29), None],
    }
)

def dump(path):
    frames = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    return {
        name: json.loads(sheet.to_json(orient="split", date_format="iso"))
        for name, sheet in frames.items()
    }
"""


def _run(cwd: Path, *parts: str) -> object:
    script = "\n".join(textwrap.dedent(part) for part in parts)
    completed = subprocess.run(
        [sys.executable, "-c", script],
        capture_output=True,
        text=True,
        check=False,
        cwd=cwd,
    )
    assert completed.returncode == 0, completed.stderr
    return json.loads(completed.stdout.strip().splitlines()[-1])


@pytest.mark.parametrize("imports", [ALIAS_FIRST, PANDAS_FIRST], ids=["alias-first", "pandas-first"])
def test_excel_writer_and_read_excel_through_alias(imports: str, tmp_path: Path) -> None:
    written = tmp_path / "alias.xlsx"
    oracle = tmp_path / "upstream.xlsx"
    read_back = _run(
        tmp_path,
        imports,
        "import openpyxl\nassert openpyxl is wolfxl",
        FRAME,
        f"""
        with pd.ExcelWriter({str(written)!r}, engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name="data", index=False)
        print(json.dumps(dump({str(written)!r})))
        """,
    )
    graded = _run(
        tmp_path,
        UPSTREAM,
        FRAME,
        f"""
        frame.to_excel({str(oracle)!r}, sheet_name="data", index=False, engine="openpyxl")
        print(json.dumps([dump({str(oracle)!r}), dump({str(written)!r})]))
        """,
    )

    expected, alias_written = graded
    assert read_back == expected
    assert alias_written == expected


APPEND = """
base = pathlib.Path({base!r})
extra = pd.DataFrame({{"name": ["delta", "epsilon"], "count": [4, 5], "ratio": [2.5, None]}})
frame.to_excel(base, sheet_name="data", index=False, engine="openpyxl")
with pd.ExcelWriter(base, engine="openpyxl", mode="a") as writer:
    extra.to_excel(writer, sheet_name="more", index=False)
with pd.ExcelWriter(base, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
    extra.to_excel(writer, sheet_name="more", index=False, startrow=1)
with pd.ExcelWriter(base, engine="openpyxl", mode="a", if_sheet_exists="overlay") as writer:
    extra.to_excel(writer, sheet_name="data", index=False, header=False, startrow=4)
print(json.dumps(dump(base)))
"""


def test_excel_writer_append_mode_matches_upstream(tmp_path: Path) -> None:
    """``ExcelWriter(mode="a")`` reloads the workbook through the alias and re-saves it."""
    aliased = _run(
        tmp_path,
        ALIAS_FIRST,
        "import pathlib\nimport openpyxl\nassert openpyxl is wolfxl",
        FRAME,
        APPEND.format(base=str(tmp_path / "alias-append.xlsx")),
    )
    upstream = _run(
        tmp_path,
        UPSTREAM,
        "import pathlib",
        FRAME,
        APPEND.format(base=str(tmp_path / "upstream-append.xlsx")),
    )

    assert set(aliased) == {"data", "more"}
    assert aliased == upstream


def test_alias_reports_api_version_to_pandas_only_and_uninstall_restores(
    tmp_path: Path,
) -> None:
    result = _run(
        tmp_path,
        """
        import json
        import numpy
        import wolfxl
        from pandas.compat import _optional
        from wolfxl.openpyxl_compat import OPENPYXL_API_VERSION

        original = _optional.get_version
        version = wolfxl.__version__
        wolfxl.install_as_openpyxl()
        aliased = _optional.import_optional_dependency("openpyxl")
        installed = {
            "openpyxl": _optional.get_version(aliased),
            "numpy": _optional.get_version(numpy) == numpy.__version__,
            "wolfxl_version": wolfxl.__version__ == version,
        }
        wolfxl.uninstall_as_openpyxl()
        print(json.dumps({
            "api": OPENPYXL_API_VERSION,
            "installed": installed,
            "restored": _optional.get_version is original,
            "after": _optional.get_version(wolfxl) == version,
        }))
        """,
    )

    assert result == {
        "api": "3.1.5",
        "installed": {"openpyxl": "3.1.5", "numpy": True, "wolfxl_version": True},
        "restored": True,
        "after": True,
    }


def _three_by_two(tmp_path: Path) -> Path:
    import openpyxl

    path = tmp_path / "dims.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in (("a", 1), ("b", 2), ("c", 3)):
        sheet.append(row)
    workbook.save(path)
    return path


def test_read_only_dimensions_restore_after_reset(tmp_path: Path) -> None:
    """pandas 3 restores ``_max_row``/``_max_column`` after reset_dimensions()."""
    import wolfxl

    workbook = wolfxl.load_workbook(_three_by_two(tmp_path), read_only=True)
    sheet = workbook.active
    saved = (sheet.max_row, sheet.max_column)
    assert saved == (3, 2)

    sheet.reset_dimensions()
    assert (sheet.max_row, sheet.max_column) == (None, None)
    assert [row for row in sheet.iter_rows(values_only=True)] == [("a", 1), ("b", 2), ("c", 3)]
    sheet._max_row, sheet._max_column = saved
    assert (sheet.max_row, sheet.max_column) == saved
    assert (sheet._max_row, sheet._max_column) == saved

    sheet.reset_dimensions()
    assert (sheet.max_row, sheet.max_column) == (None, None)
    sheet.calculate_dimension(force=True)
    assert (sheet.max_row, sheet.max_column) == saved
    workbook.close()


def test_dimension_attributes_are_read_only_outside_read_only_mode(tmp_path: Path) -> None:
    import wolfxl

    workbook = wolfxl.load_workbook(_three_by_two(tmp_path))
    sheet = workbook.active
    assert (sheet._max_row, sheet._max_column) == (3, 2)
    with pytest.raises(AttributeError, match="read_only=True"):
        sheet._max_row = 10
    assert sheet.max_row == 3
