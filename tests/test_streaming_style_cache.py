"""Styled streaming reuses source IDs without eager cell-format lookups."""

from __future__ import annotations

from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side
from openpyxl.utils.datetime import CALENDAR_MAC_1904

import wolfxl


class StyleReaderGuard:
    def __init__(self, inner: object) -> None:
        self.inner = inner
        self.calls: Counter[int] = Counter()

    def __getattr__(self, name: str) -> object:
        return getattr(self.inner, name)

    def read_format_for_style_id(self, style_id: int) -> object:
        self.calls[style_id] += 1
        return self.inner.read_format_for_style_id(style_id)

    def read_cell_format(self, *args: object) -> object:
        raise AssertionError("streamed styles must not resolve a cell coordinate")

    def read_cell_border(self, *args: object) -> object:
        raise AssertionError("unmerged borders must not materialize sheet cells")


def _styled_fixture(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.epoch = CALENDAR_MAC_1904
    for ws in (wb.active, wb.create_sheet("Second")):
        for row in range(1, 5):
            ws.cell(row, 1, row)
            cell = ws.cell(row, 2, row + 0.25)
            cell.font = Font(name="Aptos", bold=True, italic=True, color="FF123456")
            cell.fill = PatternFill(fill_type="solid", fgColor="FF445566")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = Border(left=Side(style="thin", color="FF112233"))
            cell.number_format = "0.00"
            ws.cell(row, 3, datetime(2024, 1, row, 12, 30))
            duration = ws.cell(row, 4, timedelta(days=2, hours=row))
            duration.number_format = "[h]:mm:ss"
            gradient = ws.cell(row, 5, "gradient" if row % 2 else None)
            gradient.fill = GradientFill(degree=30, stop=("FF000000", "FFFFFFFF"))
    wb.save(path)
    return path


def test_styles_shared_across_sheets_and_iterations_without_coordinate_reads(tmp_path: Path) -> None:
    path = _styled_fixture(tmp_path / "styled.xlsx")
    wb = wolfxl.load_workbook(path, read_only=True)
    guard = StyleReaderGuard(wb._rust_reader)
    wb._rust_reader = guard
    reference = openpyxl.load_workbook(path, read_only=True)
    for _ in range(2):
        for ws, reference_ws in zip(wb, reference):
            for row, reference_row in zip(ws.iter_rows(), reference_ws.iter_rows()):
                for cell, reference_cell in zip(row, reference_row):
                    assert cell.value == reference_cell.value
                    assert cell.number_format == reference_cell.number_format
                    assert bool(cell.font.bold) == bool(reference_cell.font.bold)
                    assert cell.alignment.horizontal == reference_cell.alignment.horizontal
                    assert cell.border.left.style == reference_cell.border.left.style
                    if cell.column == 2:
                        assert cell.font.name == reference_cell.font.name
                        assert cell.font.color.rgb == reference_cell.font.color.rgb
                        assert cell.fill.fgColor.rgb == reference_cell.fill.fgColor.rgb
                    if cell.column == 5:
                        assert cell.fill.degree == reference_cell.fill.degree
                        assert [(s.position, s.color.rgb) for s in cell.fill.stop] == [
                            (s.position, s.color.rgb) for s in reference_cell.fill.stop
                        ]
    assert guard.calls
    assert set(guard.calls.values()) == {1}
    assert len(wb._streaming_style_cache.styles) <= len(wb._cell_styles)
    assert all("_format_payload" not in cell.__slots__ for cell in row)
    reference.close()
    wb.close()
    assert wb._streaming_style_cache is None


def test_date_values_only_uses_style_ids_and_workbook_epoch(tmp_path: Path) -> None:
    path = _styled_fixture(tmp_path / "dates.xlsx")
    wb = wolfxl.load_workbook(path, read_only=True)
    guard = StyleReaderGuard(wb._rust_reader)
    wb._rust_reader = guard
    reference = openpyxl.load_workbook(path, read_only=True)
    assert list(wb.active.iter_rows(values_only=True)) == list(
        reference.active.iter_rows(values_only=True)
    )
    assert not guard.calls
    wb.close()
    reference.close()


@pytest.mark.parametrize("date1904", [False, True])
def test_locale_date_format_remains_datetime_in_value_rows(tmp_path: Path, date1904: bool) -> None:
    source = openpyxl.Workbook()
    if date1904:
        source.epoch = CALENDAR_MAC_1904
    source.active["A1"] = datetime(2024, 1, 15, 12, 30)
    source.active["A1"].number_format = "[$-409]m/d/yy"
    path = tmp_path / "locale-date.xlsx"
    source.save(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    wb._rust_reader = StyleReaderGuard(wb._rust_reader)
    expected = source.active["A1"].value
    assert next(wb.active.iter_rows(values_only=True))[0] == expected
    assert next(wb.active.iter_rows())[0].value == expected
    wb.close()


def test_native_reader_replacement_rebuilds_cache(tmp_path: Path) -> None:
    path = _styled_fixture(tmp_path / "replace.xlsx")
    wb = wolfxl.load_workbook(path, read_only=True)
    cell = next(wb.active.iter_rows())[1]
    first = StyleReaderGuard(wb._rust_reader)
    wb._rust_reader = first
    assert cell.font.bold
    original_cache = wb._streaming_style_cache
    second = StyleReaderGuard(first.inner)
    wb._rust_reader = second
    assert cell.font.bold
    assert second.calls == first.calls
    assert wb._streaming_style_cache is not original_cache
    wb.close()


def test_mutable_gradient_access_does_not_change_shared_style(tmp_path: Path) -> None:
    path = _styled_fixture(tmp_path / "gradient.xlsx")
    wb = wolfxl.load_workbook(path, read_only=True)
    first, second = list(wb.active.iter_rows(max_row=2, min_col=5, max_col=5))
    fill = first[0].fill
    fill.degree = 91
    fill.stop[0].position = 0.25
    fill.stop[0].color.rgb = "FF112233"
    assert first[0].fill.degree == 30
    assert second[0].fill.stop[0].position == 0
    assert second[0].fill.stop[0].color.rgb == "FF000000"
    wb.close()


def test_default_blank_style_and_closed_workbook(tmp_path: Path) -> None:
    path = _styled_fixture(tmp_path / "blank.xlsx")
    wb = wolfxl.load_workbook(path, read_only=True)
    blank = next(wb.active.iter_rows(min_row=6, max_row=6, min_col=1, max_col=1))[0]
    assert blank.value is None
    assert blank.number_format == "General"
    assert not blank.font.bold
    assert blank.border.left.style is None
    assert blank.fill.patternType is None
    with pytest.raises(RuntimeError, match="read_only"):
        blank.number_format = "0.00"
    wb.close()
    assert blank.number_format is None
    assert blank.border.left.style is None


def test_style_ids_are_scoped_to_each_workbook(tmp_path: Path) -> None:
    paths = []
    for bold in (False, True):
        reference = openpyxl.Workbook()
        reference.active["A1"] = 42
        reference.active["A1"].font = Font(bold=bold, italic=not bold)
        path = tmp_path / f"workbook-{bold}.xlsx"
        reference.save(path)
        paths.append(path)
    workbooks = [wolfxl.load_workbook(path, read_only=True) for path in paths]
    cells = [next(wb.active.iter_rows())[0] for wb in workbooks]
    assert [cell._style_id for cell in cells] == [1, 1]
    assert [bool(cell.font.bold) for cell in cells] == [False, True]
    assert [bool(cell.font.italic) for cell in cells] == [True, False]
    for wb in workbooks:
        wb.close()


def test_streamed_source_cache_does_not_change_modify_writeback(tmp_path: Path) -> None:
    from wolfxl._streaming import stream_iter_rows

    source = _styled_fixture(tmp_path / "source.xlsx")
    output = tmp_path / "edited.xlsx"
    wb = wolfxl.load_workbook(source, modify=True)
    streamed = next(stream_iter_rows(wb.active, max_row=1))[1]
    assert streamed.font.bold
    streamed.border.left.color.rgb = "FF998877"
    assert wb._borders[wb._cell_styles[streamed._style_id].borderId].left.color.rgb == "FF112233"
    wb.active["B1"].font = wolfxl.Font(italic=True)
    wb.active["B1"] = 123
    wb.save(output)
    wb.close()
    reference = openpyxl.load_workbook(output)
    assert reference.active["B1"].value == 123
    assert reference.active["B1"].font.italic
    assert not reference.active["B1"].font.bold
    assert reference.active["B2"].font.bold
    assert reference.active["B2"].fill.fgColor.rgb == "FF445566"
    assert reference.active["B2"].border.left.color.rgb == "FF112233"
    reference.close()


def test_merged_borders_match_eager_reference_without_coordinate_reads(tmp_path: Path) -> None:
    source = openpyxl.Workbook()
    ws = source.active
    ws["A1"] = "anchor"
    ws["A1"].font = Font(bold=True)
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor="FF445566")
    ws["B2"].border = Border(
        right=Side(style="thin", color="FF112233"),
        bottom=Side(style="double", color="FF223344"),
    )
    ws.merge_cells("A1:B2")
    path = tmp_path / "merged.xlsx"
    source.save(path)
    reference = openpyxl.load_workbook(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    wb._rust_reader = StyleReaderGuard(wb._rust_reader)
    for row in wb.active.iter_rows():
        for cell in row:
            expected = reference.active[cell.coordinate]
            assert bool(cell.font.bold) == bool(expected.font.bold)
            assert cell.number_format == expected.number_format
            assert cell.fill.patternType == expected.fill.patternType
            for side in ("left", "right", "top", "bottom"):
                actual_side = getattr(cell.border, side)
                expected_side = getattr(expected.border, side)
                assert actual_side.style == (expected_side.style if expected_side else None)
                if expected_side is not None and expected_side.color is not None:
                    assert actual_side.color.rgb == expected_side.color.rgb
    assert wb.active._cells == {}
    wb.close()
    reference.close()


def test_existing_stream_cell_cache_tracks_merge_and_unmerge(tmp_path: Path) -> None:
    from wolfxl._streaming import stream_iter_rows

    source = openpyxl.Workbook()
    source.active["A1"] = "anchor"
    source.active["B2"] = 42
    source.active["B2"].font = Font(bold=True)
    source.active["B2"].number_format = "0.00"
    path = tmp_path / "mutation.xlsx"
    source.save(path)
    wb = wolfxl.load_workbook(path, modify=True)
    cell = next(stream_iter_rows(wb.active, min_row=2, max_row=2, min_col=2, max_col=2))[0]
    assert cell.font.bold
    original_cache = wb._streaming_style_cache
    wb.active.merge_cells("A1:B2")
    assert not cell.font.bold
    assert cell.number_format == "General"
    assert wb._streaming_style_cache is not original_cache
    wb.active.unmerge_cells("A1:B2")
    assert cell.font.bold
    assert cell.number_format == "0.00"
    wb.close()


def test_merged_border_color_access_does_not_change_authoring_styles(tmp_path: Path) -> None:
    from wolfxl._streaming import stream_iter_rows

    source = openpyxl.Workbook()
    source.active["A1"] = "anchor"
    source.active["A1"].border = Border(left=Side(style="thin", color="FF112233"))
    source.active["B2"].border = Border(right=Side(style="double", color="FF223344"))
    source.active.merge_cells("A1:B2")
    path = tmp_path / "merged-color.xlsx"
    output = tmp_path / "merged-color-copy.xlsx"
    source.save(path)
    wb = wolfxl.load_workbook(path, modify=True)
    colors_before = [
        (border.left.color.rgb if border.left.color else None,
         border.right.color.rgb if border.right.color else None)
        for border in wb._borders
    ]
    anchor = next(stream_iter_rows(wb.active, max_row=1, max_col=1))[0]
    anchor.border.left.color.rgb = "FF998877"
    anchor.border.right.color.rgb = "FF887766"
    assert [
        (border.left.color.rgb if border.left.color else None,
         border.right.color.rgb if border.right.color else None)
        for border in wb._borders
    ] == colors_before
    assert anchor.border.left.color.rgb == "FF112233"
    assert anchor.border.right.color.rgb == "FF223344"
    wb.save(output)
    wb.close()
    assert wb._merged_border_metadata_cache is None
    reference = openpyxl.load_workbook(output)
    assert reference.active["A1"].border.left.color.rgb == "FF112233"
    assert reference.active["A1"].border.right.color.rgb == "FF223344"
    reference.close()
