"""Compact date rows retain the openpyxl 3.0.10 value/type contract."""

from __future__ import annotations

import io
from datetime import datetime, time, timedelta
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

import wolfxl
from wolfxl import _streaming


FORMATS = ("yyyy-mm-dd hh:mm:ss.000", "yyyy-mm-dd", "h:mm:ss.000", "[h]:mm:ss", "[mm]:ss", "[$-409]m/d/yy")
SERIALS = (
    0, 0.25, 0.5 / 86_400_000, 1.5 / 86_400_000, 2.5 / 86_400_000,
    0.999999999, 1, 1.5, 59, 59.5, 60, 60.5, 61, 61.5,
    41224.99998836806, 45292.123456789, -1, -0.25, -1000.123456789,
)
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _workbook(epoch, *, invalid=False):
    wb = openpyxl.Workbook()
    wb.epoch = epoch
    ws = wb.active
    for fmt in FORMATS:
        for serial in SERIALS:
            ws.append([serial])
            ws.cell(ws.max_row, 1).number_format = fmt
        for text in ("text", "45292", ""):
            ws.append([text])
            ws.cell(ws.max_row, 1).number_format = fmt
        ws.append([True])
        ws.cell(ws.max_row, 1).number_format = fmt
    caches = {}
    for fmt in FORMATS:
        for value, cell_type in ((45292.5, "n"), (0, "n"), (-0.25, "n"), ("45292", "str"), (1, "b"), (None, "n")):
            ws.append(["=1+1"])
            cell = ws.cell(ws.max_row, 1)
            cell.number_format = fmt
            caches[cell.coordinate] = (value, cell_type)
    # More than one dense compact chunk, without formula/error normalization.
    for _ in range(1030):
        ws.append([45292.5])
        ws.cell(ws.max_row, 1).number_format = FORMATS[0]
    if invalid:
        for serial in (10**20, -(10**20)):
            for fmt in (FORMATS[0], FORMATS[3]):
                ws.append([serial])
                ws.cell(ws.max_row, 1).number_format = fmt
    buffer = io.BytesIO()
    wb.save(buffer)
    result = io.BytesIO()
    with ZipFile(buffer) as source, ZipFile(result, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                for cell in root.iter(NS + "c"):
                    if cell.get("r") not in caches:
                        continue
                    value, cell_type = caches[cell.get("r")]
                    cell.set("t", cell_type)
                    node = cell.find(NS + "v")
                    if node is None:
                        node = ET.SubElement(cell, NS + "v")
                    node.text = None if value is None else str(value)
                data = ET.tostring(root)
            target.writestr(name, data)
    return result.getvalue()


def _source(data, tmp_path, source):
    if source == "fileobj":
        return io.BytesIO(data)
    target = tmp_path / "dates.xlsx"
    target.write_bytes(data)
    return target


def _assert_values_and_types(actual, expected):
    assert actual == expected
    assert [[type(value) for value in row] for row in actual] == [
        [type(value) for value in row] for row in expected
    ]


def _spy_compact(monkeypatch):
    original = _streaming._open_stream_reader
    calls = []

    class Reader:
        def __init__(self, inner):
            self.inner = inner

        def __getattr__(self, name):
            return getattr(self.inner, name)

        def read_next_row(self):
            pytest.fail("date values must not use per-cell read_next_row")

        def read_next_values_compact_chunk(self, *args):
            calls.append(True)
            return self.inner.read_next_values_compact_chunk(*args)

    monkeypatch.setattr(_streaming, "_open_stream_reader", lambda *args: Reader(original(*args)))
    return calls


@pytest.mark.parametrize("epoch", (CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904))
@pytest.mark.parametrize("source", ("path", "fileobj"))
@pytest.mark.parametrize("read_only", (True, False), ids=("read-only", "auto-stream"))
@pytest.mark.parametrize("data_only", (True, False))
def test_date_values_match_openpyxl_on_compact_path(tmp_path, monkeypatch, epoch, source, read_only, data_only):
    data = _workbook(epoch)
    oracle = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=data_only)
    try:
        expected = list(oracle.active.iter_rows(values_only=True))
    finally:
        oracle.close()
    # 3.0.10's read-only parser forgets to pass timedelta_formats, unlike
    # its normal parser. Keep WolfXL's existing duration semantics, using the
    # same pinned oracle's normal parser only for numeric duration results.
    oracle = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=data_only)
    try:
        for index, row in enumerate(oracle.active.iter_rows(values_only=True)):
            if isinstance(row[0], timedelta):
                expected[index] = row
    finally:
        oracle.close()
    calls = _spy_compact(monkeypatch)
    # Trigger the real auto-stream dispatch on this small deterministic fixture.
    monkeypatch.setattr(_streaming, "AUTO_STREAM_ROW_THRESHOLD", 1)
    wb = wolfxl.load_workbook(_source(data, tmp_path, source), read_only=read_only, data_only=data_only)
    try:
        assert wb._date_formats
        actual = list(wb.active.iter_rows(values_only=True))
        _assert_values_and_types(actual, expected)
        assert calls
        assert isinstance(actual[0][0], time) and actual[0][0] == time(0, 0)
        assert any(isinstance(row[0], datetime) for row in actual)
        assert any(isinstance(row[0], timedelta) for row in actual)
        calls.clear()
        chunks = _streaming.stream_value_chunks(wb.active, chunk_size=31)
        _assert_values_and_types([row for chunk in chunks for row in chunk], expected)
        assert calls
    finally:
        wb.close()


@pytest.mark.parametrize("epoch", (CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904))
@pytest.mark.parametrize("source", ("path", "fileobj"))
def test_unrepresentable_serials_keep_existing_numeric_fallback(tmp_path, monkeypatch, epoch, source):
    # openpyxl's worksheet parser emits #VALUE! here; WolfXL already kept the
    # serial numeric in 2.0.14. This optimization deliberately preserves that.
    data = _workbook(epoch, invalid=True)
    calls = _spy_compact(monkeypatch)
    wb = wolfxl.load_workbook(_source(data, tmp_path, source), read_only=True, data_only=True)
    try:
        actual = list(wb.active.iter_rows(values_only=True))
        assert actual[-4:] == [(10**20,), (10**20,), (-(10**20),), (-(10**20),)]
        assert all(isinstance(row[0], (int, float)) for row in actual[-4:])
        assert calls
    finally:
        wb.close()
