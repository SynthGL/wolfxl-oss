"""Generated workbook resource regressions; no binary fixtures required."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
import wolfxl


NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def workbook_bytes(*, text="shared text", cells=4, inline=False, second_sheet=False):
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/></Types>')
        sheets = 2 if second_sheet else 1
        zf.writestr("_rels/.rels", f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        zf.writestr("xl/workbook.xml", f'<workbook xmlns="{NS}" xmlns:r="{REL}"><sheets>' + ''.join(f'<sheet name="Sheet{i}" sheetId="{i}" r:id="rId{i}"/>' for i in range(1, sheets + 1)) + '</sheets></workbook>')
        zf.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + ''.join(f'<Relationship Id="rId{i}" Type="{REL}/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1, sheets + 1)) + '</Relationships>')
        if not inline:
            zf.writestr("xl/sharedStrings.xml", f'<sst xmlns="{NS}"><si><t>{text}</t></si></sst>')
        for i in range(1, sheets + 1):
            rows = ''.join(f'<row r="{r}"><c r="A{r}" t="inlineStr"><is><t>{text}</t></is></c></row>' if inline else f'<row r="{r}"><c r="A{r}" t="s"><v>0</v></c></row>' for r in range(1, cells + 1))
            zf.writestr(f"xl/worksheets/sheet{i}.xml", f'<worksheet xmlns="{NS}"><dimension ref="A1:A{cells}"/><sheetData>{rows}</sheetData></worksheet>')
    return out.getvalue()


def generate_zip_bomb(path, *, uncompressed_bytes=1024**3, members=None):
    """Create either a high-ratio part or a many-small-members admission case."""
    with ZipFile(path, "w", ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", '<Types>spreadsheetml</Types>')
        if members is not None:
            for index in range(members):
                zf.writestr(f"parts/{index}.xml", "<part/>")
        else:
            with zf.open("xl/worksheets/sheet1.xml", "w") as part:
                prefix = f'<worksheet xmlns="{NS}"><sheetData>'.encode()
                suffix = b'</sheetData></worksheet>'
                part.write(prefix)
                remaining = uncompressed_bytes - len(prefix) - len(suffix)
                chunk = b" " * (1024 * 1024)
                while remaining:
                    size = min(remaining, len(chunk))
                    part.write(chunk[:size])
                    remaining -= size
                part.write(suffix)


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("source_kind", ["bytes", "path", "file"])
@pytest.mark.parametrize("limit", ["max_package_bytes", "max_part_bytes", "max_cells", "max_shared_string_bytes", "max_materialized_string_bytes"])
def test_each_limit_refuses_before_materialization(tmp_path, limit, source_kind, read_only):
    data = workbook_bytes()
    source = data
    if source_kind == "path":
        source = tmp_path / "bounded.xlsx"
        source.write_bytes(data)
    elif source_kind == "file":
        source = BytesIO(data)
    budgets = wolfxl.Limits(**{limit: 1})
    with pytest.raises(wolfxl.LimitExceeded) as exc:
        wolfxl.load_workbook(source, read_only=read_only, limits=budgets)
    message = str(exc.value)
    assert limit in message
    assert "observed " in message
    assert "configured 1" in message
    assert "limits=" in message


@pytest.mark.parametrize("read_only", [False, True])
def test_exact_decoded_string_budgets_and_repeated_references(read_only):
    # Six decoded UTF-8 bytes: a + ampersand + a four-byte code point.
    data = workbook_bytes(text="a&amp;&#x1f600;", cells=100)
    limits = replace(wolfxl.Limits(), max_shared_string_bytes=6, max_materialized_string_bytes=6, max_cells=100)
    wb = wolfxl.load_workbook(data, read_only=read_only, limits=limits)
    values = [row[0] for row in wb.active.iter_rows(values_only=True)]
    assert values == ["a&😀"] * 100
    assert all(value is values[0] for value in values)
    wb.close()
    with pytest.raises(wolfxl.LimitExceeded, match="max_shared_string_bytes"):
        wolfxl.load_workbook(data, read_only=read_only, limits=replace(limits, max_shared_string_bytes=5))


@pytest.mark.parametrize("read_only", [False, True])
def test_shared_strings_have_one_python_identity_across_access_paths(read_only):
    text = "shared payload " * 4096
    wb = wolfxl.load_workbook(workbook_bytes(text=text, cells=128, second_sheet=True), read_only=read_only)
    first = wb.active["A1"].value
    assert first == text
    assert wb.active.cell(2, 1).value is first
    for sheet in wb.worksheets:
        assert all(row[0] is first for row in sheet.iter_rows(values_only=True))
        assert all(row[0].value is first for row in sheet.iter_rows())
    wb.close()


def test_cell_and_inline_text_budgets_are_workbook_wide():
    data = workbook_bytes(text="four", cells=2, inline=True, second_sheet=True)
    for limits, name in [(wolfxl.Limits(max_cells=3), "max_cells"), (wolfxl.Limits(max_materialized_string_bytes=15), "max_materialized_string_bytes")]:
        with pytest.raises(wolfxl.LimitExceeded, match=name):
            wolfxl.load_workbook(data, limits=limits)
    wb = wolfxl.load_workbook(data, limits=wolfxl.Limits(max_cells=4, max_materialized_string_bytes=16))
    assert wb.active["A1"].value == "four"
    wb.close()


def test_limits_apply_to_modify_mode(tmp_path):
    path = tmp_path / "modify.xlsx"
    path.write_bytes(workbook_bytes())
    with pytest.raises(wolfxl.LimitExceeded, match="max_cells"):
        wolfxl.load_workbook(path, modify=True, limits=wolfxl.Limits(max_cells=1))


@pytest.mark.parametrize("value", [-1, True, 1.5, "1", 2**64])
def test_invalid_limit_values(value):
    with pytest.raises(ValueError, match="max_cells"):
        wolfxl.Limits(max_cells=value)


def test_limits_are_immutable_and_require_a_limits_instance():
    from dataclasses import FrozenInstanceError
    with pytest.raises(FrozenInstanceError):
        wolfxl.Limits().max_cells = 1
    with pytest.raises(TypeError, match="limits="):
        wolfxl.load_workbook(workbook_bytes(), limits={"max_cells": 1})


@pytest.mark.parametrize("read_only", [False, True])
def test_generated_high_ratio_part_is_refused(tmp_path, read_only):
    path = tmp_path / "ratio.xlsx"
    generate_zip_bomb(path, uncompressed_bytes=4 * 1024**2)
    with pytest.raises((OSError, ValueError), match="compression ratio"):
        wolfxl.load_workbook(path, read_only=read_only)


def replace_parts(data, transform):
    out = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            name, value = transform(info.filename, source.read(info))
            target.writestr(name, value)
    return out.getvalue()


def test_limits_follow_relationship_targets_not_part_names_or_root_tags():
    def transform(name, value):
        if name == "xl/_rels/workbook.xml.rels":
            value = value.replace(b"worksheets/sheet1.xml", b"custom/data.part")
        elif name == "xl/worksheets/sheet1.xml":
            name = "xl/custom/data.part"
            value = value.replace(b"worksheet", b"legacyWrapper")
        return name, value

    data = replace_parts(workbook_bytes(), transform)
    with pytest.raises(wolfxl.LimitExceeded, match="max_cells"):
        wolfxl.load_workbook(data, limits=wolfxl.Limits(max_cells=3))


def test_rich_shared_run_text_is_materialized_once():
    def transform(name, value):
        if name == "xl/sharedStrings.xml":
            value = value.replace(b"<si><t>", b"<si><r><rPr><b/></rPr><t>")
            value = value.replace(b"</t></si>", b"</t></r></si>")
        return name, value

    data = replace_parts(workbook_bytes(text="rich payload " * 4096, cells=128), transform)
    wb = wolfxl.load_workbook(data, rich_text=True)
    first = wb.active["A1"].value[0].text
    for row in wb.active.iter_rows():
        assert row[0].value[0].text is first
    wb.close()


def test_eager_and_streaming_share_one_workbook_string_cache():
    from wolfxl._streaming import _open_stream_reader, stream_source

    wb = wolfxl.load_workbook(workbook_bytes(text="cross-path payload " * 1024))
    first = wb.active["A1"].value
    reader = _open_stream_reader(wb, stream_source(wb), "Sheet1", None, None, None, None)
    rows = [reader.read_next_row() for _ in range(4)]
    assert all(cell[1] is first for _, cells in rows for cell in cells)
    wb.close()


def test_literal_xml_line_endings_count_after_normalization():
    data = workbook_bytes(text="a\r\nb\rc", cells=2)
    limits = wolfxl.Limits(max_shared_string_bytes=5, max_materialized_string_bytes=5)
    wb = wolfxl.load_workbook(data, limits=limits)
    assert wb.active["A1"].value == "a\nb\nc"
    wb.close()


def test_preexisting_package_errors_keep_their_load_context():
    def transform(name, value):
        return ("xl\\styles.xml", value) if name == "xl/sharedStrings.xml" else (name, value)

    data = replace_parts(workbook_bytes(), transform)
    with pytest.raises(OSError, match=r"^native xlsx open failed: unsafe OOXML package part path:"):
        wolfxl.load_workbook(data)


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}])
def test_sheet_relationship_type_does_not_bypass_cell_budget(mode):
    def transform(name, value):
        if name == "xl/_rels/workbook.xml.rels":
            value = value.replace(b"/worksheet", b"/custom-sheet")
        return name, value

    data = replace_parts(workbook_bytes(), transform)
    with pytest.raises(wolfxl.LimitExceeded, match="max_cells"):
        wolfxl.load_workbook(data, limits=wolfxl.Limits(max_cells=3), **mode)


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}])
@pytest.mark.parametrize("field,budget", [("max_cells", 7), ("max_materialized_string_bytes", 31)])
def test_sheet_aliases_are_counted_per_reference(mode, field, budget):
    def transform(name, value):
        if name == "xl/_rels/workbook.xml.rels":
            value = value.replace(b"worksheets/sheet2.xml", b"worksheets/sheet1.xml")
        return name, value

    data = replace_parts(workbook_bytes(text="four", inline=True, second_sheet=True), transform)
    with pytest.raises(wolfxl.LimitExceeded, match=field):
        wolfxl.load_workbook(data, limits=wolfxl.Limits(**{field: budget}), **mode)


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}])
@pytest.mark.parametrize("master_last", [False, True])
def test_shared_formula_templates_are_charged_for_each_child(mode, master_last):
    formula = "1+2+3+4+5"
    master_row = 4 if master_last else 1
    cells = "".join(
        f'<row r="{r}"><c r="A{r}">'
        + (f'<f t="&#115;hared" si="&#48;">{formula}</f>' if r == master_row else '<f t="shared" si="0"/>')
        + "</c></row>"
        for r in range(1, 5)
    )

    def transform(name, value):
        if name == "xl/worksheets/sheet1.xml":
            value = f'<worksheet xmlns="{NS}"><sheetData>{cells}</sheetData></worksheet>'.encode()
        return name, value

    data = replace_parts(workbook_bytes(text=""), transform)
    with pytest.raises(wolfxl.LimitExceeded, match="max_materialized_string_bytes"):
        wolfxl.load_workbook(
            data, limits=wolfxl.Limits(max_materialized_string_bytes=4 * len(formula) - 1), **mode
        )


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}])
def test_escaped_cell_type_is_counted_as_string(mode):
    def transform(name, value):
        if name == "xl/worksheets/sheet1.xml":
            value = value.replace(b't="s"', b't="&#115;tr"').replace(b"<v>0</v>", b"<v>12345</v>")
        return name, value

    data = replace_parts(workbook_bytes(text=""), transform)
    with pytest.raises(wolfxl.LimitExceeded, match="max_materialized_string_bytes"):
        wolfxl.load_workbook(data, limits=wolfxl.Limits(max_materialized_string_bytes=19), **mode)


def test_permissive_synthesized_sheets_are_counted():
    def transform(name, value):
        if name == "xl/workbook.xml":
            value = value[:value.index(b"<sheets>")] + b"</workbook>"
        return name, value

    data = replace_parts(workbook_bytes(), transform)
    with pytest.raises(wolfxl.LimitExceeded, match="max_cells"):
        wolfxl.load_workbook(data, permissive=True, limits=wolfxl.Limits(max_cells=3))


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}])
def test_shared_strings_part_used_as_sheet_counts_each_alias(mode):
    def transform(name, value):
        if name == "xl/_rels/workbook.xml.rels":
            value = value.replace(b"worksheets/sheet1.xml", b"sharedStrings.xml")
            value = value.replace(b"worksheets/sheet2.xml", b"sharedStrings.xml")
        elif name == "xl/worksheets/sheet1.xml":
            name = "xl/sharedStrings.xml"
        return name, value

    data = replace_parts(workbook_bytes(text="four", inline=True, second_sheet=True), transform)
    with pytest.raises(wolfxl.LimitExceeded, match="max_materialized_string_bytes"):
        wolfxl.load_workbook(data, limits=wolfxl.Limits(max_materialized_string_bytes=31), **mode)


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}, {"password": "unused"}])
def test_path_replacement_after_admission_loads_the_admitted_snapshot(tmp_path, monkeypatch, mode):
    path = tmp_path / "source.xlsx"
    admitted = workbook_bytes(text="admitted", cells=4)
    path.write_bytes(admitted)
    validate = wolfxl.Limits._validate
    scanned = []

    def swap_after_scan(self, source, *, permissive=False):
        assert isinstance(source, bytes)
        scanned.append(source)
        validate(self, source, permissive=permissive)
        path.write_bytes(workbook_bytes(text="replacement", cells=8))

    monkeypatch.setattr(wolfxl.Limits, "_validate", swap_after_scan)
    wb = wolfxl.load_workbook(path, limits=wolfxl.Limits(max_cells=4), **mode)
    assert scanned and all(source == admitted for source in scanned)
    assert list(wb.active.iter_rows(values_only=True)) == [("admitted",)] * 4
    wb.close()


@pytest.mark.parametrize("read_only", [False, True])
def test_replacement_before_first_stream_read_keeps_admitted_string_table(tmp_path, monkeypatch, read_only):
    path = tmp_path / "source.xlsx"
    admitted = workbook_bytes(text="admitted shared value")
    path.write_bytes(admitted)
    wb = wolfxl.load_workbook(path, read_only=read_only)
    from wolfxl import _streaming

    first = wb.active["A1"].value
    assert first == "admitted shared value"
    monkeypatch.setattr(_streaming, "AUTO_STREAM_ROW_THRESHOLD", 1)
    path.write_bytes(workbook_bytes(text="replacement shared value"))
    assert list(wb.active.iter_rows(values_only=True)) == [("admitted shared value",)] * 4
    assert wb.active["A1"].value == "admitted shared value"
    wb.close()


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}])
@pytest.mark.parametrize("source_kind", ["bytes", "path"])
@pytest.mark.parametrize("field,inline,budget", [
    ("max_cells", False, 3),
    ("max_shared_string_bytes", False, 3),
    ("max_materialized_string_bytes", True, 15),
])
def test_nonstandard_workbook_part_is_admitted_before_normalization(
    tmp_path, mode, source_kind, field, inline, budget
):
    def transform(name, value):
        if name == "xl/workbook.xml":
            name = "xl/alternate.xml"
        elif name == "xl/_rels/workbook.xml.rels":
            name = "xl/_rels/alternate.xml.rels"
        elif name in ("_rels/.rels", "[Content_Types].xml"):
            value = value.replace(b"xl/workbook.xml", b"xl/alternate.xml")
        return name, value

    data = replace_parts(workbook_bytes(text="four", inline=inline), transform)
    limits = wolfxl.Limits(**{field: budget})
    # Exercise native admission directly, without the Python normalization path.
    with pytest.raises(wolfxl.LimitExceeded, match=field):
        limits._validate(data)
    source = data
    if source_kind == "path":
        source = tmp_path / "alternate.xlsx"
        source.write_bytes(data)
    with pytest.raises(wolfxl.LimitExceeded, match=field):
        wolfxl.load_workbook(source, limits=limits, **mode)
    wb = wolfxl.load_workbook(source, **mode)
    assert wb.active["A1"].value == "four"
    wb.close()


@pytest.mark.parametrize("read_only", [False, True])
def test_path_snapshot_preserves_named_and_differential_styles(tmp_path, read_only):
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import NamedStyle, PatternFill

    path = tmp_path / "styles.xlsx"
    original = openpyxl.Workbook()
    original.add_named_style(NamedStyle(name="SnapshotStyle"))
    original.active["A1"] = 1
    original.active.conditional_formatting.add(
        "A1:A4", CellIsRule(
            operator="greaterThan", formula=["0"],
            fill=PatternFill(fill_type="solid", fgColor="FFFF0000"),
        )
    )
    original.save(path)
    wb = wolfxl.load_workbook(path, read_only=read_only)
    assert "SnapshotStyle" in wb.style_names
    ranges = list(wb.active.conditional_formatting)
    rule = wb.active.conditional_formatting[ranges[0]][0]
    assert rule.operator == "greaterThan"
    assert rule.dxf.fill.fgColor == "FFFF0000"
    wb.close()


def test_keep_vba_uses_admitted_snapshot_after_path_replacement(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsm"
    admitted = workbook_bytes(text="admitted")
    path.write_bytes(admitted)
    validate = wolfxl.Limits._validate

    def swap_after_scan(self, source, *, permissive=False):
        validate(self, source, permissive=permissive)
        path.write_bytes(workbook_bytes(text="replacement"))

    monkeypatch.setattr(wolfxl.Limits, "_validate", swap_after_scan)
    wb = wolfxl.load_workbook(path, keep_vba=True)
    assert b"admitted" in wb.vba_archive.read("xl/sharedStrings.xml")
    assert b"replacement" not in wb.vba_archive.read("xl/sharedStrings.xml")
    wb.vba_archive.close()
    wb.close()


@pytest.mark.parametrize("mode", [{}, {"read_only": True}, {"modify": True}])
@pytest.mark.parametrize("field,budget", [
    ("max_cells", 3), ("max_materialized_string_bytes", 15),
])
def test_nonstandard_normalization_cannot_activate_unscanned_orphan_rels(mode, field, budget):
    def transform(name, value):
        if name == "xl/workbook.xml":
            name = "custom/book.xml"
        elif name in ("_rels/.rels", "[Content_Types].xml"):
            value = value.replace(b"xl/workbook.xml", b"custom/book.xml")
        # Deliberately retain canonical rels without custom/_rels/book.xml.rels.
        return name, value

    data = replace_parts(workbook_bytes(text="four", inline=True), transform)
    with pytest.raises(wolfxl.LimitExceeded, match=field):
        wolfxl.load_workbook(data, limits=wolfxl.Limits(**{field: budget}), **mode)


def test_path_snapshot_tempfile_is_removed_when_workbook_is_collected(tmp_path):
    import gc
    from pathlib import Path
    import weakref

    path = tmp_path / "source.xlsx"
    path.write_bytes(workbook_bytes())
    wb = wolfxl.load_workbook(path)
    snapshot_path = Path(wb._tempfile_path)
    assert snapshot_path.exists()
    reference = weakref.ref(wb)
    del wb
    gc.collect()
    assert reference() is None
    assert not snapshot_path.exists()


def test_decrypted_path_read_does_not_write_a_plaintext_snapshot(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace
    from wolfxl._workbook_sources import from_encrypted

    decrypted = workbook_bytes()

    class EncryptedOffice:
        def __init__(self, source):
            assert source.read() == b"encrypted source"

        def is_encrypted(self):
            return True

        def load_key(self, *, password):
            assert password == "password"

        def decrypt(self, target):
            target.write(decrypted)

    monkeypatch.setitem(sys.modules, "msoffcrypto", SimpleNamespace(OfficeFile=EncryptedOffice))
    path = tmp_path / "encrypted.xlsx"
    path.write_bytes(b"encrypted source")
    wb = from_encrypted(wolfxl.Workbook, str(path), password="password")
    assert wb._source_bytes == decrypted
    assert wb._source_path is None
    assert getattr(wb, "_tempfile_path", None) is None
    wb.close()
