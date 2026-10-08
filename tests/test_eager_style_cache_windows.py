"""Eager style batching remains bounded and preserves normal Cell behavior."""

from __future__ import annotations

from datetime import datetime, timedelta
import gc
import weakref
from pathlib import Path

import openpyxl
import wolfxl
from wolfxl._cell import _STYLE_PAYLOAD_CACHE_CELL_LIMIT, _STYLE_PAYLOAD_CACHE_DISABLED
from wolfxl._cell_style_window import StylePayloadWindow
from wolfxl.styles import Font


class CountingReader:
    def __init__(self, reader: object) -> None:
        self.reader = reader
        self.windows: list[str] = []
        self.styles: list[int] = []
        self.per_cell_calls = 0

    def __getattr__(self, name: str) -> object:
        return getattr(self.reader, name)

    def read_sheet_style_ids(self, sheet: str, range_str: str) -> object:
        self.windows.append(range_str)
        return self.reader.read_sheet_style_ids(sheet, range_str)

    def read_format_for_style_id(self, style_id: int) -> object:
        self.styles.append(style_id)
        return self.reader.read_format_for_style_id(style_id)

    def read_cell_format_rc(self, sheet: str, row: int, col: int) -> object:
        self.per_cell_calls += 1
        coordinate = f"{openpyxl.utils.get_column_letter(col)}{row}"
        return self.reader.read_cell_format(sheet, coordinate)


def test_sparse_large_dimensions_keep_bounded_style_windows(tmp_path: Path) -> None:
    source = tmp_path / "sparse.xlsx"
    initial = openpyxl.Workbook()
    for address in ("A1", "XFD1000000"):
        initial.active[address] = 7
        initial.active[address].font = openpyxl.styles.Font(bold=True)
        initial.active[address].fill = openpyxl.styles.PatternFill("solid", fgColor="FFD966")
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source)
    try:
        reader = CountingReader(workbook._rust_reader)
        workbook._rust_reader = reader
        worksheet = workbook.active
        first = worksheet["A1"]
        last = worksheet["XFD1000000"]
        assert first.font.bold and last.font.bold
        assert first.fill is last.fill
        assert first.font is last.font
        assert len(reader.windows) == 2
        assert len(reader.styles) == 1
        assert reader.per_cell_calls == 0
        cache = worksheet._style_payload_cache
        assert isinstance(cache, StylePayloadWindow)
        assert len(cache._entries) == 1
        for window in reader.windows:
            min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(window)
            assert (max_col - min_col + 1) * (max_row - min_row + 1) <= _STYLE_PAYLOAD_CACHE_CELL_LIMIT
    finally:
        workbook.close()


def test_medium_sheet_shares_style_payload_across_windows(tmp_path: Path) -> None:
    source = tmp_path / "medium.xlsx"
    initial = openpyxl.Workbook(write_only=True)
    worksheet = initial.create_sheet()
    for row in range(25000):
        cells = [openpyxl.cell.WriteOnlyCell(worksheet, value=row) for _ in range(5)]
        cells[0].font = openpyxl.styles.Font(bold=True)
        worksheet.append(cells)
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source)
    try:
        reader = CountingReader(workbook._rust_reader)
        workbook._rust_reader = reader
        worksheet = workbook.active
        assert worksheet["A1"].font.bold
        assert worksheet["A20000"].font.bold
        assert worksheet["A20001"].font.bold
        assert worksheet["A25000"].font.bold
        assert len(reader.windows) == 2
        assert len(reader.styles) == 1
        assert reader.per_cell_calls == 0
        assert len(worksheet._style_payload_cache._entries) <= _STYLE_PAYLOAD_CACHE_CELL_LIMIT
    finally:
        workbook.close()


def test_window_failure_falls_back_to_per_cell_format(tmp_path: Path) -> None:
    source = tmp_path / "failure.xlsx"
    initial = openpyxl.Workbook()
    initial.active["A30000"] = 1
    initial.active["E30000"] = 2
    initial.active["A30000"].font = openpyxl.styles.Font(bold=True)
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source)

    class FailingReader(CountingReader):
        def read_sheet_style_ids(self, sheet: str, range_str: str) -> object:
            raise ValueError("simulated unsupported batch read")

    try:
        reader = FailingReader(workbook._rust_reader)
        workbook._rust_reader = reader
        assert workbook.active["A30000"].font.bold
        assert reader.per_cell_calls == 1
        assert workbook.active._style_payload_cache is _STYLE_PAYLOAD_CACHE_DISABLED
    finally:
        workbook.close()


def test_prefilled_default_cells_share_one_component_entry(tmp_path: Path) -> None:
    source = tmp_path / "default.xlsx"
    initial = openpyxl.Workbook()
    for row in range(1, 21):
        initial.active.append([row, "text", True])
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source, data_only=True)
    try:
        cells = [cell for row in workbook.active.iter_rows() for cell in row]
        assert [cell.value for cell in cells[:3]] == [1, "text", True]
        assert all(cell.number_format == "General" for cell in cells)
        assert all(cell.font is cells[0].font for cell in cells)
        cache = workbook.active._style_component_cache
        assert cache is None or len(cache) <= 1
    finally:
        workbook.close()


def test_prefilled_cells_keep_extended_styles_and_two_save_edits(tmp_path: Path) -> None:
    source = tmp_path / "styles.xlsx"
    initial = openpyxl.Workbook()
    worksheet = initial.active
    worksheet["A1"] = 17
    worksheet["A1"].font = openpyxl.styles.Font(name="Arial", bold=True, italic=True, color="FF123456")
    worksheet["A1"].fill = openpyxl.styles.GradientFill(stop=("FFFF0000", "FF00FF00"), degree=45)
    worksheet["A1"].alignment = openpyxl.styles.Alignment(horizontal="right", wrap_text=True, text_rotation=15)
    worksheet["A1"].border = openpyxl.styles.Border(bottom=openpyxl.styles.Side(style="thin", color="FF123456"))
    worksheet["A1"].protection = openpyxl.styles.Protection(locked=False, hidden=True)
    worksheet["B1"] = datetime(2026, 10, 8, 7, 30)
    worksheet["C1"] = timedelta(hours=27)
    worksheet["D1"] = "=A1+1"
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source, modify=True)
    try:
        row = next(workbook.active.iter_rows())
        cell = row[0]
        assert cell.font.name == "Arial" and cell.font.bold and cell.font.italic
        assert cell.fill.type == "linear" and cell.fill.degree == 45
        assert len(cell.fill.stop) == 2
        assert cell.alignment.horizontal == "right" and cell.alignment.wrap_text
        assert cell.border.bottom.style == "thin"
        assert cell.protection.locked is False and cell.protection.hidden is True
        assert row[1].value == datetime(2026, 10, 8, 7, 30)
        assert row[2].value == timedelta(hours=27)
        assert row[3].value == "=A1+1" and row[3].data_type == "f"
        assert cell.comment is None and cell.hyperlink is None
        workbook.active["E1"].font = Font(name="Arial", bold=False, italic=True)
        workbook.active["E1"].value = "new style"
        cell.value = 99
        for number in (1, 2):
            output = tmp_path / f"save-{number}.xlsx"
            workbook.save(output)
            check = openpyxl.load_workbook(output)
            try:
                assert check.active["A1"].value == 99
                assert check.active["A1"].font.bold is True
                assert check.active["E1"].font.bold is False
                assert check.active["E1"].font.italic is True
                assert check.active["A1"].fill.degree == 45
                assert check.active["A1"].border.bottom.style == "thin"
                assert check.active["A1"].alignment.horizontal == "right"
            finally:
                check.close()
    finally:
        workbook.close()


def test_bulk_iteration_returns_merged_subordinates(tmp_path: Path) -> None:
    source = tmp_path / "merged.xlsx"
    initial = openpyxl.Workbook()
    initial.active["A1"] = "anchor"
    initial.active.merge_cells("A1:B2")
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source, data_only=True)
    try:
        rows = list(workbook.active.iter_rows())
        assert rows[0][0].value == "anchor"
        assert isinstance(rows[0][1], wolfxl.cell.MergedCell)
        assert rows[0][1].value is None
        assert rows[1][1].value is None
    finally:
        workbook.close()


def test_cell_iterator_keeps_edits_to_future_rows(tmp_path: Path) -> None:
    source = tmp_path / "future-edit.xlsx"
    initial = openpyxl.Workbook()
    for value in (1, 2, 3):
        initial.active.append([value])
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source, modify=True, data_only=True)
    try:
        iterator = workbook.active.iter_rows()
        assert next(iterator)[0].value == 1
        workbook.active["A3"] = 99
        workbook.active["A3"].font = Font(italic=True)
        assert next(iterator)[0].value == 2
        third = next(iterator)[0]
        assert third.value == 99 and third.font.italic
    finally:
        workbook.close()


def test_column_cell_iteration_prefills_styles_and_values(tmp_path: Path) -> None:
    source = tmp_path / "columns.xlsx"
    initial = openpyxl.Workbook()
    for row in range(1, 4):
        initial.active.append([row, row * 10])
        initial.active.cell(row, 2).font = openpyxl.styles.Font(bold=True)
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source, data_only=True)
    try:
        columns = list(workbook.active.iter_cols())
        assert [[cell.value for cell in column] for column in columns] == [[1, 2, 3], [10, 20, 30]]
        assert all(cell.font.bold for cell in columns[1])
        assert columns[1][0].font is columns[1][2].font
    finally:
        workbook.close()



def test_close_detaches_retained_window_and_native_reader(tmp_path: Path) -> None:
    source = tmp_path / "close-window.xlsx"
    initial = openpyxl.Workbook()
    initial.active["A1"] = 1
    initial.active["E30000"] = 2
    initial.active["A1"].font = openpyxl.styles.Font(bold=True)
    initial.active["E30000"].font = openpyxl.styles.Font(italic=True)
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source)
    worksheet = workbook.active
    reader = CountingReader(workbook._rust_reader)
    reader_ref = weakref.ref(reader)
    workbook._rust_reader = reader
    assert worksheet["A1"].font.bold
    window = worksheet._style_payload_cache
    assert isinstance(window, StylePayloadWindow)
    assert window._reader is reader
    del reader
    untouched = worksheet["E30000"]
    workbook.close()
    workbook.close()  # Cleanup must remain idempotent.
    gc.collect()
    assert reader_ref() is None
    assert worksheet._style_payload_cache is None
    assert worksheet._style_component_cache is None
    assert window._reader is None
    assert window.get((30000, 5)) is _STYLE_PAYLOAD_CACHE_DISABLED
    assert untouched.font.italic is False  # Fresh native reads stop at close.
    assert worksheet["A1"].font.bold  # Already resolved Cell values remain cached.


def test_close_clears_windows_before_replacing_native_reader(tmp_path: Path) -> None:
    source = tmp_path / "replace-reader.xlsx"
    initial = openpyxl.Workbook()
    initial.active["A1"] = 1
    initial.active["E30000"] = 2
    initial.active["A1"].font = openpyxl.styles.Font(bold=True)
    initial.save(source)
    initial.close()
    workbook = wolfxl.load_workbook(source)
    assert workbook.active["A1"].font.bold
    window = workbook.active._style_payload_cache
    workbook.close()

    replacement = wolfxl.load_workbook(source)
    reader = CountingReader(replacement._rust_reader)
    workbook._rust_reader = reader
    try:
        # This private-reader swap is a lifecycle probe: a fresh untouched Cell
        # must not consult the old window, even when the worksheet is retained.
        assert workbook.active["E30000"].font.bold is False
        assert reader.windows == ["A20001:E30000"]
        assert workbook.active._style_payload_cache is not window
        assert window._reader is None
    finally:
        workbook.close()
        replacement.close()
