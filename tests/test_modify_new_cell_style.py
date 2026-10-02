"""New cells written in modify mode take the style an absent cell has there.

Excel gives a cell typed into an empty position the style of its row when the
row is formatted (``customFormat="1"``), else the style of its column
(``<col style>``), else the workbook default.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

import wolfxl


def _make_styled_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "header"
    ws["A2"] = "existing"
    ws["E2"] = "styled"
    ws["E2"].font = Font(underline="single")
    ws.row_dimensions[2].font = Font(bold=True)
    ws.column_dimensions["B"].font = Font(italic=True)
    wb.save(path)


def _font(path: Path, coord: str) -> Font:
    wb = openpyxl.load_workbook(path)
    try:
        return wb["Data"][coord].font
    finally:
        wb.close()


def test_new_cells_inherit_row_then_column_style(tmp_path: Path) -> None:
    src = tmp_path / "styled.xlsx"
    _make_styled_workbook(src)

    wb = wolfxl.load_workbook(src, modify=True)
    ws = wb["Data"]
    ws["B2"] = "row and column"
    ws.cell(row=2, column=3, value="row only")
    ws["B5"] = "column only"
    ws.cell(row=5, column=4, value="neither")
    wb.save(src)
    wb.close()

    both = _font(src, "B2")
    assert both.b and not both.i
    row_only = _font(src, "C2")
    assert row_only.b and not row_only.i
    col_only = _font(src, "B5")
    assert col_only.i and not col_only.b
    neither = _font(src, "D5")
    assert not neither.b and not neither.i


def test_existing_cells_keep_their_style_in_styled_row(tmp_path: Path) -> None:
    src = tmp_path / "styled.xlsx"
    _make_styled_workbook(src)

    wb = wolfxl.load_workbook(src, modify=True)
    ws = wb["Data"]
    ws["A2"] = "changed"
    ws["E2"] = "changed too"
    ws["B2"] = "new"
    wb.save(src)
    wb.close()

    unstyled = _font(src, "A2")
    assert not unstyled.b and not unstyled.i
    own_style = _font(src, "E2")
    assert own_style.u == "single" and not own_style.b
