"""Range hyperlinks (``<hyperlink ref="A3:A5">``) attach to every covered cell on load."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

import wolfxl

_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
    "</Types>"
)
_ROOT_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
    "</Relationships>"
)
_WORKBOOK = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>'
)
_WORKBOOK_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
    "</Relationships>"
)
_SHEET_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="http://www.apache.org/" TargetMode="External"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="http://www.apache.org/" TargetMode="External"/>'
    "</Relationships>"
)


def _cell(ref: str, text: str | None) -> str:
    if text is None:
        return f'<c r="{ref}"/>'
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def _sheet_xml() -> str:
    rows = []
    for r in range(1, 6):
        cells = "".join(_cell(f"{col}{r}", f"{col}{r}") for col in "ABCD")
        rows.append(f'<row r="{r}">{cells}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<dimension ref="A1:D5"/>'
        f"<sheetData>{''.join(rows)}</sheetData>"
        "<hyperlinks>"
        '<hyperlink ref="A1" r:id="rId1"/>'
        '<hyperlink ref="A3:A5" r:id="rId2" display="http://www.apache.org"/>'
        '<hyperlink ref="C1:D2" location="Sheet1!A1" display="Top" tooltip="Jump to top"/>'
        "</hyperlinks></worksheet>"
    )


@pytest.fixture
def shared_hyperlink_path(tmp_path: Path) -> Path:
    """Workbook shaped like POI sharedhyperlink.xlsx plus an internal range link."""
    path = tmp_path / "shared_hyperlink.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _CONTENT_TYPES)
        zf.writestr("_rels/.rels", _ROOT_RELS)
        zf.writestr("xl/workbook.xml", _WORKBOOK)
        zf.writestr("xl/_rels/workbook.xml.rels", _WORKBOOK_RELS)
        zf.writestr("xl/worksheets/sheet1.xml", _sheet_xml())
        zf.writestr("xl/worksheets/_rels/sheet1.xml.rels", _SHEET_RELS)
    return path


_COORDS = [f"{col}{row}" for row in range(1, 6) for col in "ABCD"]


def _snapshot(cell: object) -> tuple[object, ...] | None:
    link = cell.hyperlink  # type: ignore[attr-defined]
    if link is None:
        return None
    return (link.ref, link.target, link.location, link.display, link.tooltip)


def test_range_hyperlink_attaches_to_every_covered_cell(shared_hyperlink_path: Path) -> None:
    ws = wolfxl.load_workbook(shared_hyperlink_path)["Sheet1"]

    for coord in ("A1", "A3", "A4", "A5"):
        link = ws[coord].hyperlink
        assert link is not None, coord
        assert link.target == "http://www.apache.org/", coord
        assert link.ref == coord, coord
    assert ws["A2"].hyperlink is None
    assert ws["B3"].hyperlink is None


def test_range_hyperlink_matches_openpyxl_for_every_cell(shared_hyperlink_path: Path) -> None:
    expected_ws = openpyxl.load_workbook(shared_hyperlink_path)["Sheet1"]
    actual_ws = wolfxl.load_workbook(shared_hyperlink_path)["Sheet1"]

    for coord in _COORDS:
        assert _snapshot(actual_ws[coord]) == _snapshot(expected_ws[coord]), coord


def test_range_hyperlink_map_is_one_entry_per_covered_cell(shared_hyperlink_path: Path) -> None:
    expected_ws = openpyxl.load_workbook(shared_hyperlink_path)["Sheet1"]
    expected_refs = {c.coordinate for row in expected_ws.iter_rows() for c in row if c.hyperlink}
    link_map = wolfxl.load_workbook(shared_hyperlink_path)["Sheet1"]._get_hyperlinks_map()

    assert set(link_map) == expected_refs
    assert {key: link.ref for key, link in link_map.items()} == {key: key for key in link_map}
    assert "A3:A5" not in link_map
    # Per-cell copies are independent objects, not one shared range link.
    assert link_map["A3"] is not link_map["A4"]


def test_read_only_load_exposes_range_hyperlinks_like_eager(shared_hyperlink_path: Path) -> None:
    eager_ws = wolfxl.load_workbook(shared_hyperlink_path)["Sheet1"]
    read_only_ws = wolfxl.load_workbook(shared_hyperlink_path, read_only=True)["Sheet1"]

    for coord in ("A1", "A3", "A4", "A5"):
        assert _snapshot(read_only_ws[coord]) == _snapshot(eager_ws[coord]), coord
    assert read_only_ws["A3"].hyperlink.target == "http://www.apache.org/"


def test_range_hyperlink_survives_modify_save_roundtrip(
    shared_hyperlink_path: Path, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    wb = wolfxl.load_workbook(shared_hyperlink_path, modify=True)
    wb["Sheet1"]["B5"].value = "touched"
    wb.save(out)

    expected_ws = openpyxl.load_workbook(shared_hyperlink_path)["Sheet1"]
    saved_ws = wolfxl.load_workbook(out)["Sheet1"]
    for coord in ("A1", "A3", "A4", "A5", "C1", "D2"):
        assert _snapshot(saved_ws[coord]) == _snapshot(expected_ws[coord]), coord
