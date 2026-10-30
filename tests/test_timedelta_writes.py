"""Durations write as openpyxl writes them: a serial day count under ``[hh]:mm:ss``."""

from __future__ import annotations

from datetime import timedelta
from pathlib import Path

import openpyxl
import pytest

import wolfxl

DURATION = timedelta(hours=27, minutes=30)


def _written(path: Path) -> tuple[object, str, str]:
    book = openpyxl.load_workbook(path)
    try:
        cell = book.active["A1"]
        return cell.value, cell.data_type, cell.number_format
    finally:
        book.close()


def _append(book):
    book.active.append((DURATION, 2))


def _assign(book):
    book.active["A1"] = DURATION


def _write_rows(book):
    book.active.write_rows([[DURATION, 2]])


@pytest.mark.parametrize("write", [_append, _assign, _write_rows])
def test_durations_write_as_openpyxl_writes_them(tmp_path: Path, write) -> None:
    reference = openpyxl.Workbook()
    reference.active["A1"] = DURATION
    reference.save(tmp_path / "reference.xlsx")

    book = wolfxl.Workbook()
    write(book)
    book.save(tmp_path / "actual.xlsx")

    assert _written(tmp_path / "actual.xlsx") == _written(tmp_path / "reference.xlsx")
    assert _written(tmp_path / "actual.xlsx") == (DURATION, "d", "[hh]:mm:ss")


def test_write_only_durations_keep_their_format(tmp_path: Path) -> None:
    book = wolfxl.Workbook(write_only=True)
    book.create_sheet().append([DURATION])
    book.save(tmp_path / "write-only.xlsx")

    assert _written(tmp_path / "write-only.xlsx") == (DURATION, "d", "[hh]:mm:ss")


def test_copied_duration_rows_stay_durations(tmp_path: Path) -> None:
    source = openpyxl.Workbook()
    source.active["A1"] = 1.25
    source.active["A1"].number_format = "[h]:mm:ss"
    source.active["B1"] = 2
    source.save(tmp_path / "source.xlsx")

    loaded = wolfxl.load_workbook(tmp_path / "source.xlsx")
    copy = wolfxl.Workbook()
    try:
        for row in loaded.active.iter_rows(values_only=True):
            copy.active.append(row)
    finally:
        loaded.close()
    copy.save(tmp_path / "copy.xlsx")

    assert _written(tmp_path / "copy.xlsx") == (timedelta(days=1, hours=6), "d", "[hh]:mm:ss")
