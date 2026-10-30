"""Community cell metadata parity on the native streaming path."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest

import wolfxl
from wolfxl import _streaming


@pytest.fixture
def cached_workbook(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["number", "text", "boolean", "error", "missing", "error text", "empty text"])
    sheet.append(["=1+29", '="hello"', "=TRUE()", "=1/0", "=1", '="#N/A"', '=""'])
    sheet["A1"].hyperlink = "https://example.com/table"
    sheet["C1"].hyperlink = openpyxl.worksheet.hyperlink.Hyperlink(
        ref="C1", location="Sheet!A2", tooltip="Internal link", display="Jump"
    )
    sheet["A4"] = "tail"
    original = BytesIO()
    book.save(original)
    output = BytesIO()
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    ET.register_namespace("", namespace["s"])
    ET.register_namespace(
        "r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    with ZipFile(original) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                caches = [
                    ("n", "30.0"),
                    ("str", "hello"),
                    ("b", "1"),
                    ("e", "#DIV/0!"),
                    ("n", None),
                    ("str", "#N/A"),
                    ("str", None),
                ]
                for cell, (kind, value) in zip(
                    root.findall("s:sheetData/s:row[@r='2']/s:c", namespace), caches
                ):
                    cell.set("t", kind)
                    cached = cell.find("s:v", namespace)
                    if cached is None:
                        cached = ET.SubElement(cell, "{" + namespace["s"] + "}v")
                    cached.text = value
                data = ET.tostring(root)
            target.writestr(entry, data)
    path = tmp_path / "cached.xlsx"
    path.write_bytes(output.getvalue())
    return path


def _source(path, file_object):
    return BytesIO(path.read_bytes()) if file_object else path


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("file_object", [False, True])
@pytest.mark.parametrize("data_only", [False, True])
def test_streaming_hyperlinks(cached_workbook, monkeypatch, read_only, file_object, data_only):
    # Force the same auto-stream branch used by large eager workbooks.
    monkeypatch.setattr(_streaming, "AUTO_STREAM_ROW_THRESHOLD", 0)
    expected = openpyxl.load_workbook(cached_workbook)
    actual = wolfxl.load_workbook(
        _source(cached_workbook, file_object), read_only=read_only, data_only=data_only
    )
    try:
        rows = actual.active.iter_rows()
        for actual_row, expected_row in zip(rows, expected.active.iter_rows(), strict=True):
            for actual_cell, expected_cell in zip(actual_row, expected_row, strict=True):
                link = actual_cell.hyperlink
                reference = expected_cell.hyperlink
                if reference is None:
                    assert link is None
                else:
                    assert (link.ref, link.target, link.location, link.tooltip, link.display) == (
                        reference.ref,
                        reference.target,
                        reference.location,
                        reference.tooltip,
                        reference.display,
                    )
    finally:
        actual.close()
        expected.close()


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("file_object", [False, True])
def test_streaming_cached_formula_types(cached_workbook, monkeypatch, read_only, file_object):
    monkeypatch.setattr(_streaming, "AUTO_STREAM_ROW_THRESHOLD", 0)
    expected = openpyxl.load_workbook(cached_workbook, data_only=True)
    actual = wolfxl.load_workbook(
        _source(cached_workbook, file_object), read_only=read_only, data_only=True
    )
    try:
        actual_row = next(actual.active.iter_rows(min_row=2, max_row=2))
        expected_row = next(expected.active.iter_rows(min_row=2, max_row=2))
        assert [(cell.value, cell.data_type) for cell in actual_row] == [
            (cell.value, cell.data_type) for cell in expected_row
        ]
    finally:
        actual.close()
        expected.close()


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("file_object", [False, True])
def test_empty_formula_cache_indexed_and_iterated(cached_workbook, read_only, file_object):
    expected = openpyxl.load_workbook(cached_workbook, read_only=read_only, data_only=True)
    actual = wolfxl.load_workbook(
        _source(cached_workbook, file_object), read_only=read_only, data_only=True
    )
    try:
        reference = expected.active["G2"]
        assert (reference.value, reference.data_type) == (None, "str")
        for cell in (
            actual.active["G2"],
            next(actual.active.iter_rows(min_row=2, max_row=2, min_col=7, max_col=7))[0],
        ):
            assert (cell.value, cell.data_type) == (reference.value, reference.data_type)
        assert list(
            actual.active.iter_rows(min_row=2, max_row=2, min_col=7, max_col=7, values_only=True)
        ) == [(None,)]
    finally:
        actual.close()
        expected.close()


@pytest.fixture
def oversized_dimension_workbook(tmp_path):
    book = openpyxl.Workbook()
    book.active["A1"] = "first"
    book.active["C3"] = "last"
    original = BytesIO()
    book.save(original)
    output = BytesIO()
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(original) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                root.find("s:dimension", namespace).set("ref", "A1:D50001")
                sheet_data = root.find("s:sheetData", namespace)
                ET.SubElement(sheet_data, "{" + namespace["s"] + "}row", r="50002")
                data = ET.tostring(root)
            target.writestr(entry, data)
    path = tmp_path / "oversized-dimension.xlsx"
    path.write_bytes(output.getvalue())
    return path


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("file_object", [False, True])
@pytest.mark.parametrize("values_only", [False, True])
def test_trailing_rows_follow_load_mode(
    oversized_dimension_workbook, read_only, file_object, values_only
):
    expected = openpyxl.load_workbook(oversized_dimension_workbook, read_only=read_only)
    actual = wolfxl.load_workbook(
        _source(oversized_dimension_workbook, file_object), read_only=read_only
    )
    try:

        def values(sheet):
            for row in sheet.iter_rows(values_only=values_only):
                yield tuple(row) if values_only else tuple(cell.value for cell in row)

        for result, reference in zip(values(actual.active), values(expected.active), strict=True):
            assert result == reference
        # Explicit bounds still request padding beyond the used cells.
        assert list(actual.active.iter_rows(min_row=4, max_row=5, max_col=2, values_only=True)) == [
            (None, None),
            (None, None),
        ]
    finally:
        actual.close()
        expected.close()


@pytest.mark.parametrize("file_object", [False, True])
@pytest.mark.parametrize("values_only", [False, True])
def test_eager_stream_extends_to_merged_range(tmp_path, file_object, values_only):
    book = openpyxl.Workbook()
    book.active["A1"] = "first"
    book.active["C3"] = "last"
    book.active.merge_cells("D2:E50002")
    path = tmp_path / "merged-tail.xlsx"
    book.save(path)
    expected = openpyxl.load_workbook(path)
    actual = wolfxl.load_workbook(_source(path, file_object))
    try:
        assert _streaming.should_auto_stream(actual.active)

        def values(sheet):
            for row in sheet.iter_rows(values_only=values_only):
                yield tuple(row) if values_only else tuple(cell.value for cell in row)

        for result, reference in zip(values(actual.active), values(expected.active), strict=True):
            assert result == reference
    finally:
        actual.close()
        expected.close()


_INLINE_CELLS = (
    '<c r="A1" t="inlineStr"/>',
    '<c r="B1" t="inlineStr"><is><t>x</t></is></c>',
    '<c r="C1" t="inlineStr"><is/></c>',
    '<c r="D1" t="inlineStr"><v>ignored</v></c>',
    '<c r="E1" t="inlineStr"><f>1+1</f><v>2</v></c>',
    '<c r="F1" t="inlineStr"></c>',
    '<c r="G1" t="inlineStr"><is><r><t>a</t></r><r><t>b</t></r></is></c>',
)


def _one_row_workbook(path, cells):
    """Save a workbook whose sheet holds exactly ``cells`` in row 1."""
    book = openpyxl.Workbook()
    book.active["A2"] = 1
    original = BytesIO()
    book.save(original)
    output = BytesIO()
    last = openpyxl.utils.get_column_letter(len(cells))
    with ZipFile(original) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                text = data.decode()
                start = text.index("<sheetData>")
                end = text.index("</sheetData>") + len("</sheetData>")
                row = '<row r="1">' + "".join(cells) + "</row>"
                text = text[:start] + f"<sheetData>{row}</sheetData>" + text[end:]
                data = text.replace('<dimension ref="A2:A2" />', f'<dimension ref="A1:{last}1"/>')
                data = data.encode()
            target.writestr(entry, data)
    path.write_bytes(output.getvalue())
    return path


@pytest.fixture
def inline_string_workbook(tmp_path):
    return _one_row_workbook(tmp_path / "inline.xlsx", _INLINE_CELLS)


@pytest.mark.parametrize("data_only", [False, True])
@pytest.mark.parametrize(("read_only", "stream"), [(False, False), (False, True), (True, False)])
def test_inline_string_without_text_keeps_inline_type(
    inline_string_workbook, monkeypatch, read_only, stream, data_only
):
    # openpyxl reports `inlineStr` until an `<is>` child turns the cell into a string.
    if stream:
        monkeypatch.setattr(_streaming, "AUTO_STREAM_ROW_THRESHOLD", 0)
    expected = openpyxl.load_workbook(
        inline_string_workbook, read_only=read_only, data_only=data_only
    )
    actual = wolfxl.load_workbook(inline_string_workbook, read_only=read_only, data_only=data_only)
    try:
        reference = [(cell.value, cell.data_type) for cell in next(expected.active.iter_rows())]
        assert reference[0] == (None, "inlineStr")
        assert [(cell.value, cell.data_type) for cell in next(actual.active.iter_rows())] == reference
        if not read_only:
            indexed = [actual.active.cell(1, column) for column in range(1, len(reference) + 1)]
            assert [(cell.value, cell.data_type) for cell in indexed] == reference
    finally:
        actual.close()
        expected.close()


_VALUELESS_CELLS = (
    '<c r="A1" t="s"/>',
    '<c r="B1" t="str"/>',
    '<c r="C1" t="b"/>',
    '<c r="D1" t="e"/>',
    '<c r="E1" t="e"></c>',
    '<c r="F1" t="e"><v></v></c>',
    '<c r="G1" t="d"/>',
    '<c r="H1"/>',
    '<c r="I1" t="e"><f>1/0</f></c>',
    '<c r="J1" t="e"><f>1/0</f><v>#DIV/0!</v></c>',
    '<c r="K1" t="b"><f>TRUE()</f></c>',
    '<c r="L1" t="d"><f>NOW()</f></c>',
    '<c r="M1" t="s"><f>A2</f></c>',
    '<c r="N1" t="b"><v></v></c>',
    '<c r="O1" t="d"></c>',
)


@pytest.mark.parametrize("data_only", [False, True])
@pytest.mark.parametrize("values_only", [False, True])
@pytest.mark.parametrize(("read_only", "stream"), [(False, False), (False, True), (True, False)])
def test_valueless_cells_keep_their_type(
    tmp_path, monkeypatch, read_only, stream, values_only, data_only
):
    # openpyxl keeps the `t=` type of a cell without a value, and a formula
    # cell reads as its formula even when its cached result is an error.
    path = _one_row_workbook(tmp_path / "valueless.xlsx", _VALUELESS_CELLS)
    if stream:
        monkeypatch.setattr(_streaming, "AUTO_STREAM_ROW_THRESHOLD", 0)
    expected = openpyxl.load_workbook(path, read_only=read_only, data_only=data_only)
    actual = wolfxl.load_workbook(path, read_only=read_only, data_only=data_only)

    def first_row(book):
        row = next(book.active.iter_rows(values_only=values_only))
        return list(row) if values_only else [(cell.value, cell.data_type) for cell in row]

    try:
        reference = first_row(expected)
        assert reference[0] == (None if values_only else (None, "s"))
        assert first_row(actual) == reference
        if not read_only and not values_only:
            indexed = [actual.active.cell(1, column) for column in range(1, len(reference) + 1)]
            assert [(cell.value, cell.data_type) for cell in indexed] == reference
    finally:
        actual.close()
        expected.close()
