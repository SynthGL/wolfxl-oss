from __future__ import annotations

import copy
import importlib
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "benchmarks"))
bench = importlib.import_module("performance_contract")
fixtures = importlib.import_module("performance_fixtures")
validation = importlib.import_module("performance_validation")


def test_fixture_is_repeatable_and_reports_actual_populated_columns(tmp_path, monkeypatch):
    clock = SimpleNamespace(datetime=SimpleNamespace(now=lambda **_: datetime(2020, 1, 1)), timezone=timezone)
    monkeypatch.setattr(openpyxl.writer.excel, "datetime", clock)
    zip_clock = SimpleNamespace(time=lambda: 0, localtime=lambda _: (2020, 1, 1, 0, 0, 0, 2, 1, -1))
    monkeypatch.setattr(zipfile, "time", zip_clock)
    path, metadata = fixtures.prepare_fixture(tmp_path, rows=12, kind="plain")
    again, same = fixtures.prepare_fixture(tmp_path, rows=12, kind="plain")
    assert path == again and metadata == same
    assert metadata["requested_plain_columns"] == 8
    assert metadata["populated_plain_columns"] == 5
    assert metadata["populated_value_cells"] == 60
    # openpyxl overwrites modified at save time; a later rebuild must still match.
    clock.datetime.now = lambda **_: datetime(2040, 1, 1)
    zip_clock.localtime = lambda _: (2040, 1, 1, 0, 0, 4, 6, 1, -1)
    other, rebuilt = fixtures.prepare_fixture(tmp_path / "other", rows=12, kind="plain")
    assert rebuilt["sha256"] == metadata["sha256"]
    assert other.read_bytes() == path.read_bytes()
    with zipfile.ZipFile(other) as archive:
        assert all(info.date_time == (2000, 1, 1, 0, 0, 0) for info in archive.infolist())
        assert all(info.external_attr == 0o600 << 16 and info.create_system == 3 for info in archive.infolist())
    path.write_bytes(b"changed")
    with pytest.raises(ValueError, match="Cached fixture changed"):
        fixtures.prepare_fixture(tmp_path, rows=12, kind="plain")


def test_bounded_verifier_uses_one_iterator_and_explicit_bounds(tmp_path, monkeypatch):
    path, _ = fixtures.prepare_fixture(tmp_path, rows=12, kind="plain")
    edits = fixtures.edits_for("top", 12)
    wb = openpyxl.load_workbook(path)
    for sheet, row, col, value in edits:
        wb[sheet].cell(row, col, value)
    wb.save(path)
    wb.close()
    cls = openpyxl.worksheet._read_only.ReadOnlyWorksheet
    original = cls.iter_rows
    calls = []
    def observed(self, *args, **kwargs):
        calls.append(kwargs)
        return original(self, *args, **kwargs)
    monkeypatch.setattr(cls, "iter_rows", observed)
    result = validation.bounded_verify(path, edits)
    assert result["verified_cells"] == 2
    assert calls == [{"values_only": True, "min_row": 2, "max_row": 3,
                      "min_col": 2, "max_col": 3}]
    with pytest.raises(AssertionError, match="!="):
        validation.bounded_verify(path, [("Data", 2, 2, "wrong")])


@pytest.mark.parametrize("case", fixtures.modification_cases())
def test_full_validation_catches_unrelated_value_and_style_changes(tmp_path, case):
    path, _ = fixtures.prepare_fixture(tmp_path, rows=12, kind=fixtures.fixture_kind(case))
    edits = fixtures.edits_for(case, 12)
    expected = validation.expected_digest(path, edits)
    output = tmp_path / "modified.xlsx"
    sample = bench.measure_edit(openpyxl, "openpyxl", path, output, edits)
    validation.full_verify(output, expected)
    assert sample["phase_seconds"]["total"] >= sample["phase_seconds"]["operation"]
    wb = openpyxl.load_workbook(output)
    wb["Data"]["A1"] = "unrelated-corruption"
    wb.save(output)
    wb.close()
    with pytest.raises(AssertionError, match="signature differs"):
        validation.full_verify(output, expected)
    bench.measure_edit(openpyxl, "openpyxl", path, output, edits)
    wb = openpyxl.load_workbook(output)
    wb["Data"]["A1"].number_format = "0.000000"
    wb.save(output)
    wb.close()
    with pytest.raises(AssertionError, match="signature differs"):
        validation.full_verify(output, expected)


def test_sparse_read_only_empty_cells_are_not_counted_as_styled(tmp_path):
    path, _ = fixtures.prepare_fixture(tmp_path, rows=20, kind="styled_sparse")
    _, eager = bench.styled_signature(openpyxl, path, read_only=False)
    _, streaming = bench.styled_signature(openpyxl, path, read_only=True)
    assert eager == streaming
    assert eager["styled_cells"] == 5


def test_comparison_refuses_changed_fixture_and_uses_direct_cumulative_ratio():
    before = {"contract": "v2", "status": "complete", "workload_config": {"cases": ["one"]}, "fixtures": {"one": {"sha256": "123"}},
              "harness_sha256": "456", "label": "baseline",
              "environment": dict.fromkeys(("python", "platform", "machine", "processor", "cpu_count"), "same"),
              "packages": {"openpyxl": {"version": "3.1.5", "python_source_sha256": "789", "import_path": "/old"}},
              "results": [{"case": "one", "engine": engine,
                           "phase_medians_seconds": {"operation": value, "total": value + 1}}
                          for engine, value in (("wolfxl", 8), ("openpyxl", 12))]}
    after = copy.deepcopy(before)
    after["label"] = "final"
    after["packages"]["openpyxl"]["import_path"] = "/new"
    after["results"][0]["phase_medians_seconds"] = {"operation": 2, "total": 3}
    result = validation.comparison(before, after)["comparisons"][0]
    assert result["cumulative_engine_speedup"] == 4
    assert result["final_speedup_vs_openpyxl"] == 6
    after["fixtures"]["one"]["sha256"] = "changed"
    with pytest.raises(ValueError, match="different fixtures"):
        validation.comparison(before, after)


def test_all_cases_tiny_end_to_end(tmp_path):
    args = bench.parse_args(["--output", str(tmp_path / "receipt.json"), "--fixture-dir", str(tmp_path / "fixtures"),
                             "--edit-rows", "20", "--styled-rows", "20", "--small-rows", "20",
                             "--rounds", "1", "--warmups", "1"])
    receipt = bench.run(args)
    assert receipt["status"] == "complete"
    assert len(receipt["results"]) == 36
    assert receipt["native"]["sha256"]
    assert receipt["packages"]["wolfxl"]["python_source_sha256"]
    assert all(len(item["raw_samples"]) == 1 for item in receipt["results"])
    assert all(item["deep_validation"]["zip_crc"] == "pass" for item in receipt["results"]
               if item["case"].startswith("edit_"))
