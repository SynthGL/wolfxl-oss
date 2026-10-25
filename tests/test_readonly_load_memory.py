"""Input ownership and deferred SST decoding without platform-sensitive RSS limits."""
from __future__ import annotations

import gc
import io
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import openpyxl
import pytest
import wolfxl
from wolfxl._rust import NativeXlsxBook

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def _package(*, padding: int = 0, sst_xml: str | None = None) -> bytes:
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
            '</Types>'
        ),
        "_rels/.rels": (
            f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" '
            f'Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        ),
        "xl/workbook.xml": (
            f'<workbook xmlns="{NS}" xmlns:r="{REL}"><sheets>'
            '<sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" '
            f'Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{REL}/sharedStrings" '
            'Target="sharedStrings.xml"/></Relationships>'
        ),
        "xl/worksheets/sheet1.xml": (
            f'<worksheet xmlns="{NS}"><dimension ref="A1:C2"/><sheetData>'
            '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c>'
            '<c r="C1"><v>42</v></c></row><row r="2"><c r="A2" t="s"><v>2</v></c>'
            '<c r="B2" t="s"><v>0</v></c><c r="C2" t="b"><v>1</v></c></row>'
            '</sheetData></worksheet>'
        ),
        "xl/sharedStrings.xml": (
            f'<sst xmlns="{NS}" count="4" uniqueCount="3">'
            '<si><t>shared &amp; text</t></si>'
            '<si><r><t>rich </t></r><r><t>text</t></r></si>'
            '<si><t>東京都</t><rPh sb="0" eb="3"><t>とうきょうと</t></rPh></si></sst>'
        ),
    }
    if sst_xml is not None:
        parts["xl/sharedStrings.xml"] = sst_xml
    buffer = io.BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)
        if padding:
            package.writestr("padding.bin", b"x" * padding, compress_type=ZIP_STORED)
    return buffer.getvalue()


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("source", ["path", "bytes", "fileobj", "bytearray", "memoryview"])
def test_shared_strings_match_eager_and_streaming(tmp_path: Path, source: str, read_only: bool) -> None:
    data = _package()
    path = tmp_path / "shared.xlsx"
    path.write_bytes(data)
    sources = {
        "path": path,
        "bytes": data,
        "fileobj": io.BytesIO(data),
        "bytearray": bytearray(data),
        "memoryview": memoryview(data),
    }
    oracle = openpyxl.load_workbook(io.BytesIO(data), read_only=read_only, data_only=True)
    expected = list(oracle.active.iter_rows(values_only=True))
    oracle.close()
    wb = wolfxl.load_workbook(sources[source], read_only=read_only, data_only=True)
    try:
        if source == "fileobj":
            sources[source].close()
        assert list(wb.active.iter_rows(values_only=True)) == expected
        assert wb.active["A1"].value == "shared & text"
        assert wb.active["A2"].value == "東京都"
    finally:
        wb.close()


@pytest.mark.parametrize("source", ["path", "bytes", "fileobj"])
def test_non_streaming_native_access_after_readonly_open(tmp_path: Path, source: str) -> None:
    data = _package()
    path = tmp_path / "shared.xlsx"
    path.write_bytes(data)
    src = {"path": path, "bytes": data, "fileobj": io.BytesIO(data)}[source]
    wb = wolfxl.load_workbook(src, read_only=True, data_only=True)
    try:
        # These bypass the streaming reader and must initialize the native SST.
        reader = wb._rust_reader
        assert reader.read_cell_value("Data", "A1", True) == {"type": "string", "value": "shared & text"}
        assert reader.read_sheet_values_plain("Data", None, True) == [
            ("shared & text", "rich text", 42), ("東京都", "shared & text", True)
        ]
        assert list(wb.active.iter_rows(values_only=True)) == [
            ("shared & text", "rich text", 42), ("東京都", "shared & text", True)
        ]
    finally:
        wb.close()


def test_native_book_retains_original_immutable_bytes_without_copy() -> None:
    # A reference-count assertion is deterministic, unlike RSS/tracemalloc:
    # tracemalloc does not see Rust allocations and RSS depends on the allocator.
    # The old data.to_vec() implementation never retained the Python bytes owner.
    data = _package(padding=8 * 1024 * 1024)
    refs = sys.getrefcount(data)
    reader = NativeXlsxBook.open_from_bytes(data)
    assert sys.getrefcount(data) == refs + 1
    del data
    gc.collect()
    assert reader.read_cell_value("Data", "A1", True)["value"] == "shared & text"


def test_mutable_input_is_snapshotted_once() -> None:
    data = bytearray(_package())
    wb = wolfxl.load_workbook(data, read_only=True, data_only=True)
    data[:] = b"x" * len(data)
    try:
        assert next(wb.active.iter_rows(values_only=True))[0] == "shared & text"
    finally:
        wb.close()


@pytest.mark.parametrize("source", ["path", "bytes", "fileobj"])
@pytest.mark.parametrize("read_only", [False, True])
def test_invalid_sst_keeps_load_time_error(tmp_path: Path, source: str, read_only: bool) -> None:
    data = _package(sst_xml=f'<sst xmlns="{NS}"><si><t>&undefined;</t></si></sst>')
    path = tmp_path / "invalid-sst.xlsx"
    path.write_bytes(data)
    src = {"path": path, "bytes": data, "fileobj": io.BytesIO(data)}[source]
    with pytest.raises(OSError, match="^native xlsx open failed: XML reference:") as error:
        wolfxl.load_workbook(src, read_only=read_only)
    assert "unrecognized entity `undefined`" in str(error.value)
