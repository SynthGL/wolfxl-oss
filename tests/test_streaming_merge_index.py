"""Dense merged sheets retain semantics without scanning every rectangle."""
from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import openpyxl
from openpyxl.styles import Font, PatternFill

import wolfxl
from wolfxl._streaming_styles import StreamingStyleCache
from wolfxl._worksheet_merged_index import MergedRangeIndex
from wolfxl.utils.cell import range_boundaries, get_column_letter


def _priorities(refs: list[str]) -> dict[tuple[int, int, int, int], int]:
    result = {}
    for rank, ref in enumerate(refs):
        c1, r1, c2, r2 = range_boundaries(ref)
        result.setdefault((r1, c1, r2, c2), rank)
    return result


def _linear(refs: list[str], row: int, col: int) -> tuple[int, int, int, int] | None:
    for bounds in _priorities(refs):
        r1, c1, r2, c2 = bounds
        if r1 <= row <= r2 and c1 <= col <= c2:
            return bounds
    return None


def _count_candidates(index: MergedRangeIndex) -> list[int]:
    visited = [0]

    class CountedTuple(tuple):
        def __getitem__(self, item):
            visited[0] += 1
            return super().__getitem__(item)

    pending = [index._root]
    while pending:
        node = pending.pop()
        if node is None:
            continue
        for name in ("by_start", "by_end", "by_column"):
            object.__setattr__(node, name, CountedTuple(getattr(node, name)))
        pending.extend((node.left, node.right))
    return visited


def test_many_vertical_merges_do_not_scan_unrelated_rows() -> None:
    refs = [f"A{4*i+1}:B{4*i+2}" for i in range(1024)]
    index = MergedRangeIndex(refs)
    visits = _count_candidates(index)
    queries = [(4*i+1, 1) for i in range(0, 1024, 17)] + [(4*i+3, 3) for i in range(0, 1024, 19)]
    order = _priorities(refs)
    for row, col in queries:
        assert index.bounds_for_cell(row, col, order) == _linear(refs, row, col)
    assert visits[0] <= 12 * len(queries)


def test_many_horizontal_merges_bisect_disjoint_columns() -> None:
    refs = [f"{get_column_letter(3*i+1)}1:{get_column_letter(3*i+2)}2" for i in range(1024)]
    index = MergedRangeIndex(refs)
    visits = _count_candidates(index)
    order = _priorities(refs)
    for col in range(1, 3073, 29):
        assert index.bounds_for_cell(1, col, order) == _linear(refs, 1, col)
    assert visits[0] <= len(range(1, 3073, 29))


def test_overlapping_ranges_preserve_first_match_across_nodes() -> None:
    refs = ["A4:A7", "A1:A5", "B2:D6", "C1:C8", "A4:A7"]
    refs.extend(f"F{100+4*i}:G{101+4*i}" for i in range(16))
    index = MergedRangeIndex(refs)
    for ordered in (refs, list(reversed(refs))):
        priorities = _priorities(ordered)
        for row in range(1, 9):
            for col in range(1, 5):
                assert index.bounds_for_cell(row, col, priorities) == _linear(ordered, row, col)


def test_large_spans_and_invalid_bounds_remain_bounded() -> None:
    index = MergedRangeIndex(["A1:B1048576", "D1:E1048576", "B2:A1", "invalid"])
    assert len(index._entries) == 2
    assert index._root.left is None and index._root.right is None
    assert index.bounds_for_cell(1048576, 2) == (1, 1, 1048576, 2)
    assert index.bounds_for_cell(1048577, 2) is None


def test_empty_merge_hot_path_does_not_touch_indexes() -> None:
    class UnusedIndexes(dict):
        def get(self, *args):
            raise AssertionError("empty sheets must not access an interval index")

    cache = StreamingStyleCache(object())
    cache.merged_bounds["Empty"] = ()
    cache.merged_indexes = UnusedIndexes()
    assert cache.merge_at(SimpleNamespace(title="Empty"), 1, 1) is None


def test_many_merge_streamed_styles_match_openpyxl(tmp_path: Path) -> None:
    source = openpyxl.Workbook()
    ws = source.active
    for group in range(128):
        row = 2 * group + 1
        anchor = ws.cell(row, 1, group + 0.25)
        anchor.font = Font(bold=True)
        anchor.fill = PatternFill(fill_type="solid", fgColor="FF112233")
        anchor.number_format = "0.00"
        ws.cell(row, 4, group)
        ws.merge_cells(start_row=row, end_row=row+1, start_column=1, end_column=2)
    path = tmp_path / "many-merges.xlsx"
    source.save(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    reference = openpyxl.load_workbook(path)
    for row in wb.active.iter_rows():
        for cell in row:
            expected = reference.active[cell.coordinate]
            assert cell.value == expected.value
            assert bool(cell.font.bold) == bool(expected.font.bold)
            assert cell.fill.patternType == expected.fill.patternType
            assert cell.number_format == expected.number_format
    cache = wb._streaming_style_cache
    assert len(cache.merged_indexes[wb.active.title]._entries) == 128
    wb.close()
    assert wb._streaming_style_cache is None
    reference.close()


def test_dense_merge_mutation_rebuilds_streaming_index(tmp_path: Path) -> None:
    from wolfxl._streaming_styles import workbook_style_cache

    source = openpyxl.Workbook()
    for group in range(12):
        row = 4 * group + 1
        source.active.merge_cells(f"A{row}:B{row+1}")
    path = tmp_path / "dense-mutation.xlsx"
    source.save(path)
    wb = wolfxl.load_workbook(path, modify=True)
    ws = wb.active
    before = workbook_style_cache(wb)
    assert before.merge_at(ws, 2, 2) == (1, 1, 2, 2)
    assert len(before.merged_indexes[ws.title]._entries) == 12
    ws.unmerge_cells("A1:B2")
    assert wb._streaming_style_cache is None
    ws.merge_cells("C1:D2")
    after = workbook_style_cache(wb)
    assert after is not before
    assert after.merge_at(ws, 2, 2) is None
    assert after.merge_at(ws, 2, 4) == (1, 3, 2, 4)
    assert len(after.merged_indexes[ws.title]._entries) == 12
    wb.close()
