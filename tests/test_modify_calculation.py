"""Recalculate-on-open flags must survive a narrow modify-mode save."""

from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest
from defusedxml.ElementTree import fromstring

from wolfxl import load_workbook


@pytest.mark.parametrize("existing_calc_pr", [True, False])
def test_full_calc_on_load_preserves_untouched_parts(
    tmp_path: Path, existing_calc_pr: bool
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    original = openpyxl.Workbook()
    original.active["A1"] = 2
    original.active["B1"] = "=A1*3"
    original.create_sheet("Untouched")["A1"] = "Keep this sheet"
    original.calculation.fullCalcOnLoad = False
    original.calculation.calcMode = "manual"
    original.calculation.calcId = 191029
    if not existing_calc_pr:
        original.calculation = None
    original.save(source)
    original.close()

    with ZipFile(source) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
    wb = load_workbook(source, modify=True)
    wb.calculation.fullCalcOnLoad = True
    wb.active["A1"] = 7
    wb.save(output)
    wb.close()

    with ZipFile(output) as archive:
        after = {name: archive.read(name) for name in archive.namelist()}
    assert after.keys() == before.keys()
    assert {name for name in before if before[name] != after[name]} == {
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
    }
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    calc = fromstring(after["xl/workbook.xml"]).find("s:calcPr", namespace)
    assert calc is not None
    assert calc.get("fullCalcOnLoad") == "1"
    source_calc = fromstring(before["xl/workbook.xml"]).find("s:calcPr", namespace)
    source_attributes = dict(source_calc.attrib) if source_calc is not None else {}
    assert calc.attrib == {**source_attributes, "fullCalcOnLoad": "1"}
    if existing_calc_pr:
        assert calc.get("calcMode") == "manual"
        assert calc.get("calcId") == "191029"
    reopened = openpyxl.load_workbook(output)
    assert reopened.active["A1"].value == 7
    assert reopened.active["B1"].value == "=A1*3"
    reopened.close()


@pytest.mark.parametrize(
    "fixture",
    [
        "real-excel-chart-cf-basic.xlsx",
        "real-excel-p1-comments-validation-protection.xlsx",
        "real-excel-external-link-basic.xlsx",
    ],
)
def test_scalar_edit_preserves_all_other_parts(tmp_path: Path, fixture: str) -> None:
    source = Path(__file__).parent / "fixtures" / "external_oracle" / fixture
    output = tmp_path / fixture
    with ZipFile(source) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
    wb = load_workbook(source, modify=True)
    sheet = wb.active
    sheet["A1"] = "Updated scalar"
    wb.save(output)
    wb.close()
    with ZipFile(output) as archive:
        after = {name: archive.read(name) for name in archive.namelist()}
    assert after.keys() == before.keys()
    changed = {name for name in before if before[name] != after[name]}
    assert len(changed) == 1
    assert next(iter(changed)).startswith("xl/worksheets/sheet")
    reopened = openpyxl.load_workbook(output)
    assert reopened.active["A1"].value == "Updated scalar"
    reopened.close()


def test_reading_calculation_is_a_byte_identical_noop(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    original = openpyxl.Workbook()
    original.active["A1"] = "=1+1"
    original.save(source)
    original.close()
    wb = load_workbook(source, modify=True)
    assert wb.calculation.fullCalcOnLoad is True
    wb.save(output)
    wb.close()
    assert output.read_bytes() == source.read_bytes()


def test_calculation_only_update_and_attribute_removal(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    original = openpyxl.Workbook()
    original.calculation.fullCalcOnLoad = False
    original.calculation.calcMode = "manual"
    original.save(source)
    original.close()
    with ZipFile(source) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
    wb = load_workbook(source, modify=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcMode = None
    wb.save(output)
    wb.close()
    with ZipFile(output) as archive:
        after = {name: archive.read(name) for name in archive.namelist()}
    assert after.keys() == before.keys()
    assert {name for name in before if before[name] != after[name]} == {
        "xl/workbook.xml"
    }
    reopened = openpyxl.load_workbook(output)
    assert reopened.calculation.fullCalcOnLoad is True
    assert reopened.calculation.calcMode is None
    reopened.close()
