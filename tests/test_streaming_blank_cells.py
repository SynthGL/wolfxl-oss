"""Absent XML cells retain the coordinate-aware streaming contract."""
from pathlib import Path
from copy import copy

import openpyxl
import pytest
from openpyxl.styles import Border, Font, GradientFill, Side

import wolfxl
from wolfxl._streaming import StreamingCell


def test_missing_cells_keep_coordinates_and_authored_empty_styles(tmp_path: Path) -> None:
    path = tmp_path / "sparse.xlsx"
    source = openpyxl.Workbook()
    source.active["A2"] = 2
    source.active["B2"].font = Font(bold=True)
    source.active["D4"] = 4
    source.save(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    try:
        rows = list(wb.active.iter_rows(min_row=1, max_row=6, min_col=1, max_col=5))
        assert len(rows) == 6
        assert all(len(row) == 5 for row in rows)
        assert rows[1][1].value is None
        assert rows[1][1].font.bold
        for row_index, row in enumerate(rows, 1):
            for column, cell in enumerate(row, 1):
                assert isinstance(cell, StreamingCell)
                assert (cell.row, cell.column) == (row_index, column)
                assert cell.coordinate == f"{openpyxl.utils.get_column_letter(column)}{row_index}"
                assert cell.parent is wb.active
                if (row_index, column) not in ((2, 1), (2, 2), (4, 4)):
                    assert cell.value is None
                    assert cell.data_type == "n"
                    assert not cell.font.bold
                    assert cell.fill.patternType is None
                    assert cell.number_format == "General"
                    assert cell.alignment.horizontal is None
    finally:
        wb.close()


@pytest.mark.parametrize("attribute", ["value", "font", "fill", "border", "alignment", "number_format", "typo"])
def test_missing_cells_refuse_mutations_and_observe_close(tmp_path: Path, attribute: str) -> None:
    path = tmp_path / "blank.xlsx"
    source = openpyxl.Workbook()
    source.active["A1"] = 1
    source.save(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    cell = next(wb.active.iter_rows(min_row=2, max_row=2, max_col=1))[0]
    assert cell.coordinate == "A2"
    assert "A2" in repr(cell)
    assert cell.number_format == "General"
    with pytest.raises(RuntimeError, match="read_only=True.*A2"):
        setattr(cell, attribute, "changed")
    wb.close()
    assert cell.number_format is None
    assert cell.value is None
    assert not cell.font.bold
    assert cell.fill.patternType is None
    assert cell.border.left.style is None


def test_missing_cell_shallow_copy_keeps_coordinate_snapshot(tmp_path: Path) -> None:
    path = tmp_path / "copy-blank.xlsx"
    source = openpyxl.Workbook()
    source.active["A1"] = 1
    source.save(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    try:
        cell = next(wb.active.iter_rows(min_row=2, max_row=2, max_col=1))[0]
        cloned = copy(cell)
        assert cloned is not cell
        assert cloned == cell
        assert cloned.parent is cell.parent
        assert cloned.coordinate == "A2"
        assert cloned.font == cell.font
        assert cloned.fill == cell.fill
        assert cloned.number_format == cell.number_format
        with pytest.raises(RuntimeError, match="read_only=True"):
            cloned.value = "changed"
    finally:
        wb.close()


def test_missing_cells_keep_default_source_border(tmp_path: Path) -> None:
    path = tmp_path / "border-zero.xlsx"
    source = openpyxl.Workbook()
    source._cell_styles[0].borderId = source._borders.add(
        Border(left=Side(style="thin", color="FF123456")))
    source.active["A1"] = 1
    source.save(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    try:
        cell = next(wb.active.iter_rows(min_row=2, max_row=2, max_col=1))[0]
        assert cell.border.left.style == "thin"
        assert cell.border.left.color.rgb == "FF123456"
    finally:
        wb.close()


def test_gradient_at_source_style_zero_keeps_fresh_mutable_fill(tmp_path: Path) -> None:
    path = tmp_path / "gradient-zero.xlsx"
    source = openpyxl.Workbook()
    source._cell_styles[0].fillId = source._fills.add(
        GradientFill(degree=30, stop=("FF000000", "FFFFFFFF")))
    source.active["A1"] = 1
    source.save(path)
    wb = wolfxl.load_workbook(path, read_only=True)
    try:
        cell = next(wb.active.iter_rows(min_row=2, max_row=2, max_col=1))[0]
        assert cell.value is None
        assert cell.fill.degree == 30
        changed = cell.fill
        changed.degree = 90
        changed.stop[0].color.rgb = "FF123456"
        assert cell.fill.degree == 30
        assert cell.fill.stop[0].color.rgb == "FF000000"
    finally:
        wb.close()


def test_missing_cell_border_tracks_reader_replacement_and_merge_state(tmp_path: Path) -> None:
    from wolfxl._streaming import stream_iter_rows

    path = tmp_path / "merge-state.xlsx"
    source = openpyxl.Workbook()
    source.active["A1"] = "anchor"
    source.active["A1"].border = Border(bottom=Side(style="double"))
    source.save(path)
    wb = wolfxl.load_workbook(path, modify=True)
    try:
        cell = next(stream_iter_rows(wb.active, min_row=2, max_row=2, min_col=2, max_col=2))[0]
        assert cell.border.bottom.style is None
        class ReaderReplacement:
            def __init__(self, inner):
                self.inner = inner

            def __getattr__(self, name):
                return getattr(self.inner, name)

        wb._rust_reader = ReaderReplacement(wb._rust_reader)
        assert cell.number_format == "General"
        assert cell.border.bottom.style is None
        wb.active.merge_cells("A1:B2")
        assert cell.border.bottom.style == "double"
        assert not cell.font.bold
        wb.active.unmerge_cells("A1:B2")
        assert cell.border.bottom.style is None
        assert cell == StreamingCell(wb.active, 2, 2, None, None, "blank")
    finally:
        wb.close()


def test_native_blank_batches_preserve_retained_cells_and_bounds(tmp_path, monkeypatch):
    import wolfxl._rust as native

    path = tmp_path / "blank-batches.xlsx"
    source = openpyxl.Workbook()
    source.active["B1"] = 1
    source.active["D300"] = 300
    source.save(path)
    calls = []
    original = native.streaming_blank_rows

    def record(*args):
        calls.append(args[2:])
        return original(*args)

    monkeypatch.setattr(native, "streaming_blank_rows", record)
    wb = wolfxl.load_workbook(path, read_only=True)
    try:
        rows = list(wb.active.iter_rows(min_col=2, max_col=5, max_row=310))
        assert len(rows) == 310
        assert rows[0][0].value == 1
        assert rows[299][2].value == 300
        assert all(cell.parent is wb.active for row in rows for cell in row)
        assert [(row[0].row, row[-1].column) for row in rows] == [
            (index, 5) for index in range(1, 311)
        ]
        assert rows[1][0].coordinate == "B2"
        assert rows[128][0].coordinate == "B129"
        assert rows[309][-1].coordinate == "E310"
        assert len({id(cell) for row in rows for cell in row}) == 310 * 4
        assert calls
        assert all(count <= 128 and count * (last - first + 1) <= 4096
                   for _, count, first, last in calls)
    finally:
        wb.close()


def test_native_blank_batch_does_not_retain_worksheet():
    import gc
    import weakref
    import wolfxl._rust as native
    from wolfxl._streaming import StreamingBlankCell

    class Owner:
        pass

    owner = Owner()
    reference = weakref.ref(owner)
    rows = native.streaming_blank_rows(StreamingBlankCell, owner, 300, 3, 2, 4)
    assert [[cell.coordinate for cell in row] for row in rows] == [
        ["B300", "C300", "D300"],
        ["B301", "C301", "D301"],
        ["B302", "C302", "D302"],
    ]
    del owner
    gc.collect()
    assert reference() is not None
    del rows
    gc.collect()
    assert reference() is None


@pytest.mark.parametrize("bounds", [
    (1, 1025, 1, 1), (1, 1, 1, 32769), (1, 2, 1, 20000),
    (1, 1, 3, 1), (2**63 - 1, 1, 1, 1),
])
def test_native_blank_batch_rejects_invalid_or_unbounded_requests(bounds):
    import wolfxl._rust as native
    from wolfxl._streaming import StreamingBlankCell

    with pytest.raises(ValueError):
        native.streaming_blank_rows(StreamingBlankCell, None, *bounds)


def test_blank_batch_fallback_and_zero_width(monkeypatch):
    import wolfxl._rust as native
    from wolfxl._streaming import StreamingBlankCell, _blank_rows

    assert native.streaming_blank_rows(StreamingBlankCell, None, 1, 2, 4, 3) == [(), ()]
    monkeypatch.setattr(native, "streaming_blank_rows", None)
    rows = list(_blank_rows(None, StreamingBlankCell, 300, 302, 2, 3))
    assert [[cell.coordinate for cell in row] for row in rows] == [
        ["B300", "C300"], ["B301", "C301"]
    ]
