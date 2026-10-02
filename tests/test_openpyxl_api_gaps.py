"""openpyxl parity for sheet-scoped names, table autoFilters, print titles
after a sheet removal, and ranged chart anchors."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName as XDefinedName
from openpyxl.worksheet.table import Table as XTable

import wolfxl
from wolfxl.chart import BarChart, Reference
from wolfxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from wolfxl.workbook.defined_name import DefinedName


def _workbook_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        return zf.read("xl/workbook.xml").decode("utf-8")


def test_loaded_sheet_scoped_name_keeps_local_sheet_id(tmp_path: Path) -> None:
    path = tmp_path / "local.xlsx"
    op = openpyxl.Workbook()
    op.active.title = "First"
    second = op.create_sheet("Second")
    second.defined_names.add(XDefinedName("LocalName", attr_text="Second!$B$5"))
    op.save(path)

    wb = wolfxl.load_workbook(path)
    name = wb["Second"].defined_names["LocalName"]
    assert name.attr_text == "Second!$B$5"
    assert name.localSheetId == 1
    assert "LocalName" not in wb["First"].defined_names
    assert "LocalName" not in wb.defined_names


def test_worksheet_defined_names_add_persists_after_sheet_removal(
    tmp_path: Path,
) -> None:
    path = tmp_path / "written.xlsx"
    wb = wolfxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Other")
    ws = wb.create_sheet("Data")
    ws.defined_names.add(DefinedName("LocalName", attr_text="Data!$B$5"))
    wb.save(path)

    assert '<definedName name="LocalName" localSheetId="1">' in _workbook_xml(path)
    reloaded = openpyxl.load_workbook(path)
    assert reloaded["Data"].defined_names["LocalName"].attr_text == "Data!$B$5"
    assert "LocalName" not in reloaded["Other"].defined_names


def test_print_titles_target_the_right_sheet_after_removal(tmp_path: Path) -> None:
    path = tmp_path / "titles.xlsx"
    wb = wolfxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Report")
    ws["A1"] = "header"
    ws.print_title_rows = "1:2"
    wb.save(path)

    assert wolfxl.load_workbook(path)["Report"].print_title_rows == "$1:$2"
    assert openpyxl.load_workbook(path)["Report"].print_title_rows == "$1:$2"


def test_loaded_table_exposes_its_autofilter(tmp_path: Path) -> None:
    path = tmp_path / "tables.xlsx"
    op = openpyxl.Workbook()
    ws = op.active
    for row in (["Region", "Sales"], ["East", 1], ["West", 2]):
        ws.append(row)
    ws.add_table(XTable(displayName="Sales", ref="A1:B3"))
    op.save(path)

    expected = openpyxl.load_workbook(path).active.tables["Sales"].autoFilter
    loaded = wolfxl.load_workbook(path).active.tables["Sales"].autoFilter
    assert expected is not None and loaded is not None
    assert loaded.ref == expected.ref == "A1:B3"


def test_add_chart_honours_two_cell_anchor(tmp_path: Path) -> None:
    path = tmp_path / "chart.xlsx"
    wb = wolfxl.Workbook()
    ws = wb.active
    for value in (1, 2, 3):
        ws.append([value])
    chart = BarChart()
    chart.add_data(Reference(ws, range_string="Sheet!A1:A3"))
    chart.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=1, row=1),
        to=AnchorMarker(col=7, row=11),
    )
    ws.add_chart(chart)
    wb.save(path)

    anchor = openpyxl.load_workbook(path).active._charts[0].anchor
    assert type(anchor).__name__ == "TwoCellAnchor"
    assert (anchor._from.col, anchor._from.row) == (1, 1)
    assert (anchor.to.col, anchor.to.row) == (7, 11)
