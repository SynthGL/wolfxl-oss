"""Pin the existing Cell-free value/style-ID ingestion recipe."""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import pytest

import wolfxl
from wolfxl.styles import Font


OPTIONS = dict(include_format=True, include_extended_format=False,
               include_coordinate=False, include_style_id=True)


def _source(path: Path, epoch: datetime) -> None:
    initial = openpyxl.Workbook()
    initial.epoch = epoch
    worksheet = initial.active
    worksheet["A1"] = datetime(2026, 10, 8, 7, 30, 15)
    worksheet["A2"] = timedelta(hours=27)
    worksheet["A3"] = 3.25
    worksheet["A3"].font = openpyxl.styles.Font(bold=True)
    worksheet["A4"] = "=A3*2"
    worksheet["A5"] = "merged anchor"
    worksheet["B5"].border = openpyxl.styles.Border(bottom=openpyxl.styles.Side(style="thin"))
    worksheet.merge_cells("A5:B5")
    worksheet["A6"] = datetime(2026, 10, 8, 7, 30, 15, 123456)
    initial.save(path)
    initial.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    formula_cell = root.find(".//s:c[@r='A4']", ns)
    value = formula_cell.find("s:v", ns)
    value.text = "6.5"
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


@pytest.mark.parametrize("epoch", [openpyxl.utils.datetime.WINDOWS_EPOCH, openpyxl.utils.datetime.MAC_EPOCH])
def test_bounded_records_return_values_and_source_style_ids_without_cells(tmp_path: Path, epoch: datetime) -> None:
    source = tmp_path / "records.xlsx"
    _source(source, epoch)
    workbook = wolfxl.load_workbook(source, data_only=True)
    try:
        worksheet = workbook.active
        records = []
        for first_row in (1, 3, 5):
            batch = list(worksheet.iter_cell_records(min_row=first_row, max_row=first_row + 1,
                                                    min_col=1, max_col=2, include_empty=True, **OPTIONS))
            assert len(batch) == 4
            records.extend(batch)
        by_position = {(record["row"], record["column"]): record for record in records}
        assert by_position[1, 1]["value"] == datetime(2026, 10, 8, 7, 30, 15)
        assert isinstance(by_position[2, 1]["value"], datetime)  # Native elapsed representation.
        assert by_position[2, 1]["number_format"] == "[hh]:mm:ss"
        assert by_position[3, 1]["value"] == 3.25
        assert by_position[3, 1]["style_id"] > 0
        assert "bold" not in by_position[3, 1]
        assert by_position[4, 1]["value"] == 6.5
        assert by_position[4, 1]["formula"] == "=A3*2"
        assert by_position[6, 1]["value"].microsecond == 0
        assert by_position[5, 2]["value"] is None
        assert "style_id" not in by_position[5, 2]
        assert "number_format" not in by_position[5, 2]
        assert all("coordinate" not in record for record in records)
        assert worksheet._cells == {}
    finally:
        workbook.close()


def test_record_formula_mode_and_format_gate_are_explicit(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    _source(source, openpyxl.utils.datetime.WINDOWS_EPOCH)
    workbook = wolfxl.load_workbook(source, data_only=True)
    try:
        worksheet = workbook.active
        record = next(worksheet.iter_cell_records(min_row=4, max_row=4, min_col=1, max_col=1,
                                                 data_only=False, include_cached_formula_value=True, **OPTIONS))
        assert record["value"] == "=A3*2" and record["cached_value"] == 6.5
        unformatted = next(worksheet.iter_cell_records(min_row=3, max_row=3, min_col=1, max_col=1,
                                                      include_format=False, include_style_id=True))
        assert "style_id" not in unformatted
        assert worksheet._cells == {}
    finally:
        workbook.close()


def test_record_style_ids_remain_source_derived_for_pending_style_edits(tmp_path: Path) -> None:
    source = tmp_path / "pending.xlsx"
    _source(source, openpyxl.utils.datetime.WINDOWS_EPOCH)
    workbook = wolfxl.load_workbook(source, modify=True)
    try:
        worksheet = workbook.active
        before = next(worksheet.iter_cell_records(min_row=3, max_row=3, min_col=1, max_col=1, **OPTIONS))
        worksheet["A3"].value = 99
        worksheet["A3"].font = Font(italic=True)
        after = next(worksheet.iter_cell_records(min_row=3, max_row=3, min_col=1, max_col=1, **OPTIONS))
        assert after["value"] == 99
        assert after["style_id"] == before["style_id"]
    finally:
        workbook.close()
