"""Plain and non-plain cell assignments in one write-mode sheet all reach the file."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import wolfxl
from wolfxl.styles import Font


def _saved_values(path: Path) -> dict[str, object]:
    book = openpyxl.load_workbook(path)
    try:
        return {
            cell.coordinate: cell.value
            for row in book.active.iter_rows()
            for cell in row
            if cell.value is not None
        }
    finally:
        book.close()


@pytest.mark.parametrize("styled", [False, True], ids=["plain", "styled"])
@pytest.mark.parametrize("plain_count", [1, 2])
def test_plain_and_date_assignments_all_save(tmp_path: Path, plain_count: int, styled: bool) -> None:
    book = wolfxl.Workbook()
    sheet = book.active
    expected: dict[str, object] = {}
    for col in range(1, plain_count + 1):
        sheet.cell(row=1, column=col, value=col)
        expected[sheet.cell(row=1, column=col).coordinate] = col
    sheet["A2"] = date(2024, 1, 2)
    expected["A2"] = datetime(2024, 1, 2)
    if styled:
        sheet["A1"].font = Font(bold=True)
    book.save(tmp_path / "mixed.xlsx")

    assert _saved_values(tmp_path / "mixed.xlsx") == expected
