"""Tests for wolfxl.calc WorkbookEvaluator."""

from __future__ import annotations

import datetime
import os
import tempfile
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from wolfxl.calc._evaluator import WorkbookEvaluator
from wolfxl.calc._functions import ExcelError
from wolfxl.utils.datetime import CALENDAR_MAC_1904

import wolfxl


def _make_sum_chain_workbook() -> wolfxl.Workbook:
    """Create a workbook: A1=10, A2=20, A3=SUM(A1:A2), A4=A3*2."""
    wb = wolfxl.Workbook()
    ws = wb.active
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = "=SUM(A1:A2)"
    ws["A4"] = "=A3*2"
    return wb


def _roundtrip(wb: wolfxl.Workbook) -> tuple[wolfxl.Workbook, str]:
    """Save and reload a workbook. Caller must delete the temp file."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    return wolfxl.load_workbook(path), path


class TestLoadAndCalculate:
    def test_sum_chain_write_mode(self) -> None:
        wb = _make_sum_chain_workbook()
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!A3"] == 30.0
        assert results["Sheet!A4"] == 60.0

    def test_sum_chain_after_roundtrip(self) -> None:
        wb = _make_sum_chain_workbook()
        wb2, path = _roundtrip(wb)
        try:
            ev = WorkbookEvaluator()
            ev.load(wb2)
            results = ev.calculate()
            assert results["Sheet!A3"] == 30.0
            assert results["Sheet!A4"] == 60.0
        finally:
            wb2.close()
            os.unlink(path)

    def test_if_conditional(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["B1"] = "=IF(A1>50,A1*2,0)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 200

    def test_if_false_branch(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["B1"] = "=IF(A1>50,A1*2,0)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 0

    def test_nested_functions(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 3
        ws["A2"] = -5
        ws["A3"] = 7
        ws["B1"] = "=SUM(A1:A3)"
        ws["B2"] = "=ABS(A2)"
        ws["B3"] = "=MAX(B1,B2)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 5.0
        assert results["Sheet!B2"] == 5.0
        assert results["Sheet!B3"] == 5.0

    def test_literal_formula(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=42"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!A1"] == 42.0

    def test_direct_ref(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["B1"] = "=A1"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 100

    def test_binary_operations(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 3
        ws["B1"] = "=A1+A2"
        ws["B2"] = "=A1-A2"
        ws["B3"] = "=A1*A2"
        ws["B4"] = "=A1/A2"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 13.0
        assert results["Sheet!B2"] == 7.0
        assert results["Sheet!B3"] == 30.0
        assert abs(results["Sheet!B4"] - 10 / 3) < 1e-10

    def test_iferror(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 0
        ws["B1"] = "=IFERROR(A1,0)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 10


class TestCrossSheet:
    def test_cross_sheet_sum(self) -> None:
        wb = wolfxl.Workbook()
        ws1 = wb.active
        ws1["A1"] = 100
        ws1["A2"] = 200
        ws2 = wb.create_sheet("Summary")
        ws2["A1"] = "=SUM(Sheet!A1:A2)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Summary!A1"] == 300.0


class TestRecalculate:
    def test_perturbation_propagates(self) -> None:
        wb = _make_sum_chain_workbook()
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()

        result = ev.recalculate({"Sheet!A1": 15})
        assert result.propagated_cells == 2  # A3 and A4 changed
        assert result.total_formula_cells == 2
        assert result.propagation_ratio == 1.0
        assert result.max_chain_depth > 0

        # Verify new values
        delta_map = {d.cell_ref: d for d in result.deltas}
        assert delta_map["Sheet!A3"].new_value == 35.0  # 15+20
        assert delta_map["Sheet!A4"].new_value == 70.0  # 35*2

    def test_hardcoded_no_propagation(self) -> None:
        """A workbook with all hardcoded values should have propagation_ratio=0."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30  # hardcoded, not formula
        ws["A4"] = 60  # hardcoded, not formula

        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()

        result = ev.recalculate({"Sheet!A1": 15})
        assert result.propagation_ratio == 0.0
        assert result.propagated_cells == 0

    def test_mixed_propagation(self) -> None:
        """Some formulas, some hardcoded."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = "=SUM(A1:A2)"  # formula - will propagate
        ws["A4"] = 60  # hardcoded - won't propagate
        ws["A5"] = "=A3+A4"  # formula, depends on A3 (propagates) and A4 (static)

        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()

        result = ev.recalculate({"Sheet!A1": 15})
        assert result.propagated_cells == 2  # A3 and A5 changed
        assert result.total_formula_cells == 2
        assert result.propagation_ratio == 1.0

    def test_tolerance(self) -> None:
        """Small perturbation within tolerance should show no delta."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10.0
        ws["A2"] = "=A1"

        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()

        # Perturb by exactly 0 (same value)
        result = ev.recalculate({"Sheet!A1": 10.0})
        assert result.propagated_cells == 0

    def test_recalc_result_structure(self) -> None:
        wb = _make_sum_chain_workbook()
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()

        result = ev.recalculate({"Sheet!A1": 11})
        assert isinstance(result.perturbations, dict)
        assert isinstance(result.deltas, tuple)
        assert all(isinstance(d, wolfxl.calc.CellDelta) for d in result.deltas)
        assert isinstance(result.propagation_ratio, float)


class TestDeterminism:
    def test_100_rounds_identical(self) -> None:
        """Same perturbation 100 times must produce identical results."""
        wb = _make_sum_chain_workbook()
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()

        results = []
        for _ in range(100):
            # Reset to original values
            ev._cell_values["Sheet!A1"] = 10
            ev._cell_values["Sheet!A2"] = 20
            ev.calculate()
            r = ev.recalculate({"Sheet!A1": 11})
            results.append(r)

        # All results should be identical
        first = results[0]
        for r in results[1:]:
            assert r.propagated_cells == first.propagated_cells
            assert r.total_formula_cells == first.total_formula_cells
            assert len(r.deltas) == len(first.deltas)
            for d1, d2 in zip(first.deltas, r.deltas):
                assert d1.cell_ref == d2.cell_ref
                assert d1.new_value == d2.new_value
                assert d1.old_value == d2.old_value


class TestComplexExpressions:
    """Complex nested formulas that the regex-based evaluator couldn't handle."""

    def test_function_times_number(self) -> None:
        """=SUM(A1:A2)*2 — function result as binary operand."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["B1"] = "=SUM(A1:A2)*2"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 60.0

    def test_number_plus_function(self) -> None:
        """=5+SUM(A1:A2) — number + function call."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["B1"] = "=5+SUM(A1:A2)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 35.0

    def test_function_minus_function(self) -> None:
        """=SUM(A1:A2)-SUM(A3:A4) — two function calls in binary op."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["A2"] = 200
        ws["A3"] = 50
        ws["A4"] = 75
        ws["B1"] = "=SUM(A1:A2)-SUM(A3:A4)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 175.0

    def test_round_of_product(self) -> None:
        """=ROUND(SUM(A1:A3)*1.1,2) — binary expression inside function arg."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30
        ws["B1"] = "=ROUND(SUM(A1:A3)*1.1,2)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 66.0

    def test_round_sum_times_if(self) -> None:
        """=ROUND(SUM(A1:A3)*IF(A4>0,1.1,1.0),2) — the poster-child complex case."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30
        ws["A4"] = 1
        ws["B1"] = "=ROUND(SUM(A1:A3)*IF(A4>0,1.1,1.0),2)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 66.0

    def test_if_with_function_condition_and_args(self) -> None:
        """=IF(SUM(A1:A3)>50,SUM(A1:A3)*2,0) — functions in all IF positions."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30
        ws["B1"] = "=IF(SUM(A1:A3)>50,SUM(A1:A3)*2,0)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 120.0

    def test_operator_precedence(self) -> None:
        """=A1+A2*A3 must respect multiplication-first precedence."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 2
        ws["A2"] = 3
        ws["A3"] = 4
        ws["B1"] = "=A1+A2*A3"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 14.0  # 2+(3*4), not (2+3)*4

    @pytest.mark.parametrize(
        ("formula", "expected"),
        [
            ("=A1^3", 8.0),
            ("=A1*A2^2", 18.0),  # ^ binds tighter than *
            ("=2^3^2", 64.0),  # left-associative, like Excel
            ("=-2^2", 4.0),  # unary minus binds tighter than ^
            ("=100000*(1+A3)^(A4-1)", 100000.0),  # zero exponent from a cell
            ("=2^(A4-1)", 1.0),
            ("=0^-1", ExcelError.DIV0),
            ("=(0-8)^0.5", ExcelError.NUM),
            ("=10^400", ExcelError.NUM),
            ("=(0-2)^99999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999999", ExcelError.NUM),
        ],
    )
    def test_exponent_operator(self, formula: str, expected: object) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 2
        ws["A2"] = 3
        ws["A3"] = 0.045
        ws["A4"] = 1
        ws["B1"] = formula
        ev = WorkbookEvaluator()
        ev.load(wb)
        assert ev.calculate()["Sheet!B1"] == expected

    @pytest.mark.parametrize("operator", ["^", "-", "*"])
    def test_operator_inside_quoted_sheet_name(self, operator: str) -> None:
        wb = wolfxl.Workbook()
        wb.active["A1"] = "='A" + operator + "B'!A1+1"
        other = wb.create_sheet("A" + operator + "B")
        other["A1"] = 41
        ev = WorkbookEvaluator()
        ev.load(wb)
        assert ev.calculate()["Sheet!A1"] == 42

    @pytest.mark.parametrize(
        ("formula", "expected"),
        [
            ("=EDATE(A1,3)", 45777),  # Jan 31 + 3 months clamps to Apr 30
            ("=EOMONTH(A1,3)", 45777),
            ("=YEAR(A1)", 2025),
            ("=MONTH(A2)", 3),
            ("=DAY(A2)", 15),
            ("=DAYS(A2,A1)", 43),
            ("=HOUR(A2)", 18),
            ("=MINUTE(A2)", 30),
            ("=SECOND(A2)", 45),
        ],
    )
    def test_date_functions_read_date_cells(self, formula: str, expected: object) -> None:
        """Date-formatted cells hold dates, which Excel treats as serial numbers."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = datetime.date(2025, 1, 31)
        ws["A2"] = datetime.datetime(2025, 3, 15, 18, 30, 45)
        ws["B1"] = formula
        ev = WorkbookEvaluator()
        ev.load(wb)
        assert ev.calculate()["Sheet!B1"] == expected

    @pytest.mark.parametrize(
        ("formula", "expected"),
        [
            ("=EDATE(A1,3)", 44315),  # Apr 30, 2025 in the 1904 system
            ("=EOMONTH(A1,3)", 44315),
            ("=DATE(2025,1,31)", 44226),
            ("=EDATE(A1,0)=DATE(2025,1,31)", True),
            ("=YEAR(A1)", 2025),
            ("=DAY(EDATE(A1,1))", 28),
            ("=DAYS(DATE(2025,3,15),A1)", 43),
            ("=DATE(1904,1,1)", 0),  # the 1904 epoch is serial 0
            ("=DATE(1903,12,31)", ExcelError.NUM),
        ],
    )
    def test_date_functions_use_1904_date_system(self, formula: str, expected: object) -> None:
        """Serials follow the workbook's 1904 date system, as Excel's do."""
        from wolfxl.utils.datetime import CALENDAR_MAC_1904

        wb = wolfxl.Workbook()
        wb.epoch = CALENDAR_MAC_1904
        ws = wb.active
        ws["A1"] = datetime.date(2025, 1, 31)
        ws["B1"] = formula
        ev = WorkbookEvaluator()
        ev.load(wb)
        assert ev.calculate()["Sheet!B1"] == expected

    def test_recalculate_uses_1904_date_system(self) -> None:
        """Perturbed date inputs recalculate in the workbook's 1904 date system."""
        from wolfxl.utils.datetime import CALENDAR_MAC_1904

        wb = wolfxl.Workbook()
        wb.epoch = CALENDAR_MAC_1904
        ws = wb.active
        ws["A1"] = 1
        ws["B1"] = "=EDATE(DATE(2025,1,31),A1)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        assert ev.calculate()["Sheet!B1"] == 44254  # Feb 28, 2025
        result = ev.recalculate({"Sheet!A1": 3})
        deltas = {d.cell_ref: d.new_value for d in result.deltas}
        assert deltas["Sheet!B1"] == 44315  # Apr 30, 2025

    def test_parenthesized_expression(self) -> None:
        """=(A1+A2)*A3 — parens override default precedence."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 2
        ws["A2"] = 3
        ws["A3"] = 4
        ws["B1"] = "=(A1+A2)*A3"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 20.0  # (2+3)*4

    def test_if_result_times_number(self) -> None:
        """=IF(A1>0,A1,0)*2 — function result used in binary operation."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["B1"] = "=IF(A1>0,A1,0)*2"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 20.0

    def test_comparison_at_top_level(self) -> None:
        """=A1>B1 should return a boolean."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["B1"] = 50
        ws["C1"] = "=A1>B1"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!C1"] is True

    def test_multi_term_arithmetic(self) -> None:
        """=A1+A2+A3-A4 — three additive ops, left-to-right associativity."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30
        ws["A4"] = 5
        ws["B1"] = "=A1+A2+A3-A4"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 55.0

    def test_complex_perturbation_propagation(self) -> None:
        """Perturbation through complex formulas still propagates correctly."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["A2"] = 200
        ws["B1"] = "=SUM(A1:A2)*2"       # 600
        ws["B2"] = "=IF(B1>500,B1*1.1,0)"  # 660
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()
        result = ev.recalculate({"Sheet!A1": 110})
        assert result.propagation_ratio == 1.0
        delta_map = {d.cell_ref: d for d in result.deltas}
        assert delta_map["Sheet!B1"].new_value == 620.0  # (110+200)*2
        assert abs(delta_map["Sheet!B2"].new_value - 682.0) < 0.01  # 620*1.1


class TestEdgeCases:
    def test_load_required_before_calculate(self) -> None:
        ev = WorkbookEvaluator()
        with pytest.raises(RuntimeError, match="Call load"):
            ev.calculate()

    def test_load_required_before_recalculate(self) -> None:
        ev = WorkbookEvaluator()
        with pytest.raises(RuntimeError, match="Call load"):
            ev.recalculate({"Sheet1!A1": 1})

    def test_empty_workbook(self) -> None:
        wb = wolfxl.Workbook()
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results == {}

    def test_division_by_zero(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 0
        ws["B1"] = "=A1/A2"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == "#DIV/0!"


class TestTextComparison:
    """Verify that string comparisons work in formulas (PR review fix)."""

    def test_string_equality(self) -> None:
        """=IF(A1="OK",1,0) with string value in A1."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "OK"
        ws["B1"] = '=IF(A1="OK",1,0)'
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 1

    def test_string_inequality(self) -> None:
        """=IF(A1="OK",1,0) with different string value."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "FAIL"
        ws["B1"] = '=IF(A1="OK",1,0)'
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 0

    def test_string_comparison_case_insensitive(self) -> None:
        """Excel string comparisons are case-insensitive."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "ok"
        ws["B1"] = '=IF(A1="OK",1,0)'
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 1

    def test_string_not_equal(self) -> None:
        """=IF(A1<>"OK",1,0) with string comparison."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "FAIL"
        ws["B1"] = '=IF(A1<>"OK",1,0)'
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 1


class TestQuotedCommasInArgs:
    """Verify that commas inside string literals don't split args (PR review fix)."""

    def test_if_with_comma_in_string(self) -> None:
        """=IF(TRUE,"a,b","c") should not split on the comma inside quotes."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = '=IF(TRUE,"a,b","c")'
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!A1"] == "a,b"

    def test_if_false_branch_with_comma(self) -> None:
        """=IF(FALSE,"a","b,c") picks the false branch correctly."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = '=IF(FALSE,"a","b,c")'
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!A1"] == "b,c"


class TestScientificNotation:
    """Verify that scientific notation numeric literals parse correctly (PR review fix)."""

    def test_1e3(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 5
        ws["B1"] = "=A1+1E3"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 1005.0

    def test_negative_exponent(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["B1"] = "=A1*2.5e-1"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 25.0


# ---------------------------------------------------------------------------
# Error value propagation through the evaluator
# ---------------------------------------------------------------------------


class TestErrorPropagation:
    """Verify ExcelError values propagate through arithmetic and comparisons."""

    def test_division_by_zero_returns_excel_error(self) -> None:
        """=1/0 should produce ExcelError.DIV0, not a string."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=1/0"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert isinstance(results["Sheet!A1"], ExcelError)
        assert results["Sheet!A1"] == ExcelError.DIV0

    def test_error_propagates_through_addition(self) -> None:
        """=1/0 in A1, =A1+5 in B1 should propagate #DIV/0! to B1."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=1/0"
        ws["B1"] = "=A1+5"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!A1"] == ExcelError.DIV0
        assert results["Sheet!B1"] == ExcelError.DIV0

    def test_error_propagates_through_multiplication(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=1/0"
        ws["B1"] = "=A1*10"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == ExcelError.DIV0

    def test_error_propagates_through_comparison(self) -> None:
        """Comparisons with error operands return the error."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=1/0"
        ws["B1"] = "=A1>5"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == ExcelError.DIV0

    def test_error_chain_three_deep(self) -> None:
        """A1=1/0, B1=A1+1, C1=B1*2 - error flows through entire chain."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=1/0"
        ws["B1"] = "=A1+1"
        ws["C1"] = "=B1*2"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!A1"] == ExcelError.DIV0
        assert results["Sheet!B1"] == ExcelError.DIV0
        assert results["Sheet!C1"] == ExcelError.DIV0

    def test_iferror_catches_propagated_error(self) -> None:
        """IFERROR should catch errors from formula chains."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=1/0"
        ws["B1"] = "=IFERROR(A1, 0)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!A1"] == ExcelError.DIV0
        assert results["Sheet!B1"] == 0

    def test_sum_over_error_cells_skips_errors(self) -> None:
        """SUM aggregation skips error cells (Excel range behavior)."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = "=1/0"
        ws["A3"] = 20
        ws["B1"] = "=SUM(A1:A3)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 30.0

    def test_concatenation_with_error_propagates(self) -> None:
        """& operator should propagate errors too."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = "=1/0"
        ws["B1"] = '=A1&" suffix"'
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == ExcelError.DIV0

    def test_recalculate_propagates_new_error(self) -> None:
        """Recalc: perturbing denominator to zero propagates error."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 5
        ws["B1"] = "=A1/A2"
        ws["C1"] = "=B1+100"
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev.calculate()
        recalc = ev.recalculate({"Sheet!A2": 0})
        # Both B1 and C1 should now be #DIV/0!
        assert ev._cell_values["Sheet!B1"] == ExcelError.DIV0
        assert ev._cell_values["Sheet!C1"] == ExcelError.DIV0
        assert recalc.propagated_cells == 2


# ---------------------------------------------------------------------------
# OFFSET function tests
# ---------------------------------------------------------------------------


class TestOFFSET:
    def test_single_cell_offset(self) -> None:
        """OFFSET(A1, 2, 0) returns value of A3."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30
        ws["B1"] = "=OFFSET(A1,2,0)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 30

    def test_offset_with_col(self) -> None:
        """OFFSET(A1, 0, 1) returns value of B1."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["B1"] = 99
        ws["C1"] = "=OFFSET(A1,0,1)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!C1"] == 99

    def test_sum_of_offset_range(self) -> None:
        """SUM(OFFSET(A1, 0, 0, 5, 1)) sums A1:A5."""
        wb = wolfxl.Workbook()
        ws = wb.active
        for i in range(1, 6):
            ws[f"A{i}"] = i * 10
        ws["B1"] = "=SUM(OFFSET(A1,0,0,5,1))"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == 150.0  # 10+20+30+40+50

    def test_offset_out_of_bounds(self) -> None:
        """OFFSET with negative result row returns #REF!."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["B1"] = "=OFFSET(A1,-5,0)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!B1"] == ExcelError.REF

    def test_offset_dynamic_height(self) -> None:
        """OFFSET with height from a cell reference."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30
        ws["B1"] = 3
        ws["C1"] = "=SUM(OFFSET(A1,0,0,B1,1))"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert results["Sheet!C1"] == 60.0


class TestNamedRanges:
    """Named range resolution in the evaluator."""

    def test_single_cell_named_range(self) -> None:
        """A named range pointing to a single cell resolves its value."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["B1"] = "=Revenue*2"
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev._named_ranges["REVENUE"] = "Sheet!A1"
        results = ev.calculate()
        assert results["Sheet!B1"] == 200.0

    def test_range_named_range_in_sum(self) -> None:
        """SUM(NamedRange) where NamedRange refers to A1:A3."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = 30
        ws["B1"] = "=SUM(Expenses)"
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev._named_ranges["EXPENSES"] = "Sheet!A1:A3"
        results = ev.calculate()
        assert results["Sheet!B1"] == 60.0

    def test_case_insensitive(self) -> None:
        """Named range lookup is case-insensitive."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 42
        ws["B1"] = "=myName"
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev._named_ranges["MYNAME"] = "Sheet!A1"
        results = ev.calculate()
        assert results["Sheet!B1"] == 42

    def test_named_range_in_formula(self) -> None:
        """Named range used in arithmetic: =Revenue-COGS."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 1000
        ws["A2"] = 600
        ws["B1"] = "=Revenue-COGS"
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev._named_ranges["REVENUE"] = "Sheet!A1"
        ev._named_ranges["COGS"] = "Sheet!A2"
        results = ev.calculate()
        assert results["Sheet!B1"] == 400.0

    def test_named_range_cross_sheet(self) -> None:
        """Named range can refer to a cell on another sheet."""
        wb = wolfxl.Workbook()
        ws1 = wb.active
        ws2 = wb.create_sheet("Data")
        ws2["A1"] = 500
        ws1["A1"] = "=TotalSales"
        ev = WorkbookEvaluator()
        ev.load(wb)
        ev._named_ranges["TOTALSALES"] = "Data!A1"
        results = ev.calculate()
        assert results["Sheet!A1"] == 500

    def test_named_range_perturbation(self) -> None:
        """Recalculation through named range after perturbation."""
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["B1"] = "=Revenue*2"
        ev = WorkbookEvaluator()
        ev.load(wb)
        # Inject named ranges and re-register the formula so the graph
        # tracks the dependency through the named range.
        ev._named_ranges["REVENUE"] = "Sheet!A1"
        ev._graph.add_formula(
            "Sheet!B1", "=Revenue*2", "Sheet",
            named_ranges=ev._named_ranges,
        )
        ev.calculate()
        recalc = ev.recalculate({"Sheet!A1": 200}, tolerance=1e-10)
        assert any(d.cell_ref == "Sheet!B1" and d.new_value == 400.0 for d in recalc.deltas)

    def test_write_mode_has_no_named_ranges(self) -> None:
        """Write-mode workbook returns empty defined_names."""
        wb = wolfxl.Workbook()
        assert wb.defined_names == {}

    def test_loaded_workbook_defined_names(self, tmp_path) -> None:
        """Names read from a saved file drive calculation and recalculation."""
        from wolfxl.workbook.defined_name import DefinedName

        wb = wolfxl.Workbook()
        ws = wb.active
        ws.title = "My Data"
        for row, value in enumerate([10, 20, 30], start=1):
            ws[f"A{row}"] = value
        ws["B1"] = 0.5
        ws["C1"] = "=SUM(Sales)*Rate"
        wb.defined_names["Sales"] = DefinedName(
            "Sales", attr_text="'My Data'!$A$1:$A$3"
        )
        wb.defined_names["Rate"] = DefinedName("Rate", attr_text="'My Data'!$B$1")
        path = tmp_path / "names.xlsx"
        wb.save(str(path))

        loaded = wolfxl.load_workbook(str(path), modify=True)
        try:
            ev = WorkbookEvaluator()
            ev.load(loaded)
            assert ev.calculate()["My Data!C1"] == 30.0
            recalc = ev.recalculate({"My Data!A1": 110})
            assert [(d.cell_ref, d.new_value) for d in recalc.deltas] == [
                ("My Data!C1", 80.0)
            ]
        finally:
            loaded.close()


# 2024-01-01 as an Excel serial in the 1900 and 1904 date systems.
_JAN_1_2024 = {False: 45292, True: 43830}
_SHEET_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _date_input_workbook(date1904: bool) -> wolfxl.Workbook:
    wb = wolfxl.Workbook()
    if date1904:
        wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws["A1"] = datetime.date(2024, 1, 1)
    ws["A2"] = datetime.datetime(2024, 1, 1, 18, 0)
    ws["A3"] = datetime.time(6, 0)
    ws["A4"] = datetime.timedelta(hours=36)
    ws["A5"] = True
    return wb


def _calculate_column_b(wb: wolfxl.Workbook, formulas: list[str]) -> list[object]:
    ws = wb.active
    for row, formula in enumerate(formulas, start=1):
        ws[f"B{row}"] = formula
    ev = WorkbookEvaluator()
    ev.load(wb)
    results = ev.calculate()
    return [results[f"Sheet!B{row}"] for row in range(1, len(formulas) + 1)]


class TestDateCellSerials:
    """Date, datetime, time and duration cells read as Excel serial numbers."""

    @pytest.mark.parametrize("date1904", [False, True])
    def test_arithmetic_and_comparison(self, date1904: bool) -> None:
        base = _JAN_1_2024[date1904]
        expected = {
            "=A1": base,
            "=A1+1": base + 1,
            "=A1=DATE(2024,1,1)": True,
            "=A2": base + 0.75,
            "=A2-A1": 0.75,
            "=A2+A3": base + 1,
            "=A3": 0.25,
            "=A3*24": 6.0,
            "=A4": 1.5,
            "=A4*24": 36.0,
            "=A1+A4": base + 1.5,
            "=A1<A2": True,
            "=A2>A1+1": False,
            '=A1&"d"': f"{base}d",
        }
        wb = _date_input_workbook(date1904)
        assert _calculate_column_b(wb, list(expected)) == list(expected.values())

    def test_bool_and_stored_values_unchanged(self) -> None:
        wb = _date_input_workbook(False)
        wb.active["B1"] = "=A5"
        ev = WorkbookEvaluator()
        ev.load(wb)
        assert ev.calculate()["Sheet!B1"] is True
        # The evaluator's input map mirrors the workbook's raw cell values.
        assert ev._cell_values["Sheet!A1"] == datetime.date(2024, 1, 1)
        assert ev._cell_values["Sheet!A3"] == datetime.time(6, 0)
        assert ev._cell_values["Sheet!A4"] == datetime.timedelta(hours=36)
        assert wb.active["A1"].value == datetime.date(2024, 1, 1)

    @pytest.mark.parametrize("date1904", [False, True])
    def test_date_functions(self, date1904: bool) -> None:
        base = _JAN_1_2024[date1904]
        expected = {
            "=YEAR(A2)": 2024,
            "=MONTH(A1)": 1,
            "=DAY(A2)": 1,
            "=HOUR(A2)": 18,
            "=HOUR(A3)": 6,
            "=MINUTE(A3)": 0,
            "=DAYS(A2,A1)": 0,
            "=EDATE(A1,1)": base + 31,
            "=EOMONTH(A1,0)": base + 30,
            '=TEXT(A1,"mm/dd/yyyy")': "01/01/2024",
            '=TEXT(A2,"yyyy-mm-dd")': "2024-01-01",
            "=MAX(A1:A2)": base + 0.75,
            "=MIN(A1:A2)": base,
        }
        wb = _date_input_workbook(date1904)
        assert _calculate_column_b(wb, list(expected)) == list(expected.values())

    @pytest.mark.parametrize("date1904", [False, True])
    def test_lookups_and_criteria(self, date1904: bool) -> None:
        wb = wolfxl.Workbook()
        if date1904:
            wb.epoch = CALENDAR_MAC_1904
        ws = wb.active
        ws["A1"] = datetime.date(2024, 1, 1)
        ws["A2"] = datetime.datetime(2024, 1, 2)
        ws["A3"] = datetime.date(2024, 1, 3)
        for row, (label, amount) in enumerate([("a", 10), ("b", 20), ("c", 30)], start=1):
            ws[f"C{row}"] = label
            ws[f"D{row}"] = amount
        expected = {
            "=VLOOKUP(DATE(2024,1,2),A1:C3,3,FALSE)": "b",
            "=VLOOKUP(DATE(2024,1,2)+0.5,A1:C3,3)": "b",
            "=MATCH(DATE(2024,1,3),A1:A3,0)": 3,
            "=XLOOKUP(DATE(2024,1,1),A1:A3,C1:C3)": "a",
            '=COUNTIF(A1:A3,">="&DATE(2024,1,2))': 2,
            '=SUMIF(A1:A3,">"&DATE(2024,1,1),D1:D3)': 50,
            "=A2=DATE(2024,1,2)": True,
        }
        for row, formula in enumerate(expected, start=1):
            ws[f"F{row}"] = formula
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert [results[f"Sheet!F{row}"] for row in range(1, len(expected) + 1)] == list(
            expected.values()
        )

    def test_dates_before_march_1900_follow_lotus_serials(self) -> None:
        wb = wolfxl.Workbook()
        ws = wb.active
        ws["A1"] = datetime.date(1900, 1, 1)
        ws["A2"] = datetime.date(1900, 2, 28)
        ws["A3"] = datetime.date(1900, 3, 1)
        expected = {
            "=A1": 1,
            "=A2": 59,
            "=A3": 61,  # serial 60 is Excel's phantom 1900-02-29
            "=A3-A2": 2,
            "=A2=DATE(1900,2,28)": True,
            "=A3=DATE(1900,3,1)": True,
            "=YEAR(A1)": 1900,
            "=DAY(A2)": 28,
            "=MONTH(A3)": 3,
            "=DAY(A3)": 1,
        }
        assert _calculate_column_b(wb, list(expected)) == list(expected.values())

    @pytest.mark.parametrize("date1904", [False, True])
    def test_recalculate_date_perturbations(self, date1904: bool) -> None:
        base = _JAN_1_2024[date1904]
        wb = wolfxl.Workbook()
        if date1904:
            wb.epoch = CALENDAR_MAC_1904
        ws = wb.active
        ws["A1"] = "=DATE(2024,1,1)"
        ws["A2"] = datetime.date(2024, 1, 1)
        ws["A3"] = "=A2+1"
        ws["A5"] = "=A1=A2"
        ws["A6"] = datetime.time(6, 0)
        ws["B6"] = "=A6*24"
        ev = WorkbookEvaluator()
        ev.load(wb)
        results = ev.calculate()
        assert (results["Sheet!A3"], results["Sheet!A5"], results["Sheet!B6"]) == (
            base + 1,
            True,
            6.0,
        )

        perturbations = {
            "Sheet!A2": datetime.date(2024, 2, 1),
            "Sheet!A6": datetime.timedelta(hours=30),
        }
        recalc = ev.recalculate(perturbations)
        assert recalc.perturbations == perturbations
        assert {d.cell_ref: d.new_value for d in recalc.deltas} == {
            "Sheet!A3": base + 32,
            "Sheet!A5": False,
            "Sheet!B6": 30.0,
        }

        recalc = ev.recalculate({"Sheet!A2": datetime.datetime(2024, 1, 1, 12, 0)})
        assert {d.cell_ref: d.new_value for d in recalc.deltas} == {"Sheet!A3": base + 1.5}
        recalc = ev.recalculate({"Sheet!A2": datetime.date(2024, 1, 1)})
        assert {d.cell_ref: d.new_value for d in recalc.deltas} == {
            "Sheet!A3": base + 1,
            "Sheet!A5": True,
        }

    @pytest.mark.parametrize("modify", [False, True])
    @pytest.mark.parametrize("date1904", [False, True])
    def test_calculate_saves_numeric_caches(self, tmp_path, date1904: bool, modify: bool) -> None:
        """Date inputs keep caches fresh; formulas, values and formats are untouched."""
        base = _JAN_1_2024[date1904]
        source = tmp_path / "source.xlsx"
        original = openpyxl.Workbook()
        if date1904:
            original.epoch = CALENDAR_MAC_1904  # equal to openpyxl's MAC_EPOCH
        ws = original.active
        ws.title = "Sheet"
        ws["A1"] = "=DATE(2024,1,1)"
        ws["A2"] = datetime.date(2024, 1, 1)
        ws["A2"].number_format = "yyyy-mm-dd"
        ws["A3"] = "=A2+1"
        ws["A5"] = "=A1=A2"
        ws["A6"] = datetime.timedelta(hours=12)
        ws["A6"].number_format = "[h]:mm:ss"
        ws["B6"] = "=A6*24"
        ws["A7"] = datetime.time(6, 30)
        ws["A7"].number_format = "hh:mm:ss"
        ws["B7"] = "=A7*24"
        ws["A8"] = datetime.datetime(2024, 1, 1, 18)
        ws["A8"].number_format = "yyyy-mm-dd hh:mm:ss"
        ws["B8"] = "=A8-A2"
        ws["B9"] = "=A2+A7"
        original.save(source)

        wb = wolfxl.load_workbook(source, modify=modify)
        try:
            assert wb.active["A6"].value == datetime.timedelta(hours=12)
            assert wb.active["A7"].value == datetime.time(6, 30)
            results = wb.calculate()
            assert results["Sheet!A1"] == base
            assert results["Sheet!A3"] == base + 1
            assert results["Sheet!A5"] is True
            assert results["Sheet!B6"] == 12.0
            assert results["Sheet!B7"] == pytest.approx(6.5)
            assert results["Sheet!B8"] == pytest.approx(0.75)
            assert results["Sheet!B9"] == pytest.approx(base + 6.5 / 24)
            # A what-if perturbation never reaches the saved caches.
            wb.recalculate({"Sheet!A2": datetime.date(2024, 2, 1)})
            output = tmp_path / "calculated.xlsx"
            wb.save(output)
        finally:
            wb.close()

        with zipfile.ZipFile(output) as package:
            root = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
        cells = {cell.get("r"): cell for cell in root.findall(".//s:c", _SHEET_NS)}
        assert cells["A1"].find("s:v", _SHEET_NS).text == str(base)
        assert cells["A3"].find("s:f", _SHEET_NS).text == "A2+1"
        assert cells["A3"].find("s:v", _SHEET_NS).text == str(base + 1)
        assert cells["A3"].get("t") is None
        assert cells["A5"].get("t") == "b"
        assert cells["A5"].find("s:v", _SHEET_NS).text == "1"
        assert cells["B6"].find("s:v", _SHEET_NS).text == "12"
        assert float(cells["B7"].find("s:v", _SHEET_NS).text) == pytest.approx(6.5)
        assert float(cells["B8"].find("s:v", _SHEET_NS).text) == pytest.approx(0.75)
        assert float(cells["B9"].find("s:v", _SHEET_NS).text) == pytest.approx(base + 6.5 / 24)

        saved = openpyxl.load_workbook(output)
        try:
            sheet = saved["Sheet"]
            assert saved.epoch == original.epoch
            assert sheet["A2"].value == datetime.datetime(2024, 1, 1)
            assert sheet["A2"].number_format == "yyyy-mm-dd"
            assert sheet["A3"].value == "=A2+1"
            assert sheet["A6"].value == datetime.timedelta(hours=12)
            assert sheet["A6"].number_format == "[h]:mm:ss"
            assert sheet["B6"].value == "=A6*24"
            assert sheet["A7"].value == datetime.time(6, 30)
            assert sheet["A7"].number_format == "hh:mm:ss"
            assert sheet["A8"].value == datetime.datetime(2024, 1, 1, 18)
            assert sheet["A8"].number_format == "yyyy-mm-dd hh:mm:ss"
            assert sheet["B7"].value == "=A7*24"
            assert sheet["B8"].value == "=A8-A2"
            assert sheet["B9"].value == "=A2+A7"
        finally:
            saved.close()

        cached = openpyxl.load_workbook(output, data_only=True)
        try:
            sheet = cached["Sheet"]
            assert sheet["A1"].value == base
            assert sheet["A3"].value == base + 1
            assert sheet["A5"].value is True
            assert sheet["B6"].value == 12
            assert sheet["B7"].value == pytest.approx(6.5)
            assert sheet["B8"].value == pytest.approx(0.75)
            assert sheet["B9"].value == pytest.approx(base + 6.5 / 24)
        finally:
            cached.close()

    @pytest.mark.parametrize("modify", [False, True])
    @pytest.mark.parametrize("date1904", [False, True])
    def test_date_plus_time_with_optional_calc(self, tmp_path, date1904: bool, modify: bool) -> None:
        """TIME stays in the existing optional backend, not the built-in registry."""
        pytest.importorskip("formulas")
        source = tmp_path / "date-plus-time.xlsx"
        original = openpyxl.Workbook()
        if date1904:
            original.epoch = CALENDAR_MAC_1904
        original.active["A1"] = datetime.date(2024, 1, 1)
        original.active["A1"].number_format = "yyyy-mm-dd"
        original.active["B1"] = "=A1+TIME(6,30,0)"
        original.save(source)
        with wolfxl.load_workbook(source, modify=modify) as wb:
            result = wb.calculate()
            assert result["Sheet!B1"] == pytest.approx(_JAN_1_2024[date1904] + 6.5 / 24)
