"""Calculated formula results survive XLSX saves and are read back as caches.

Consumer-visible contract: after ``calculate()`` and ``save()``, reopening with
``data_only=True`` shows the computed typed values while the formula text is
kept; edits made after ``calculate()`` never leave stale results; package
parts and ZIP metadata survive, and uncomputed results remain uncached.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from wolfxl.workbook.defined_name import DefinedName

import wolfxl

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SHEET1 = "xl/worksheets/sheet1.xml"
READERS = (wolfxl.load_workbook, openpyxl.load_workbook)


def _workbook() -> wolfxl.Workbook:
    wb = wolfxl.Workbook()
    ws = wb.active
    ws["A1"] = 2
    ws["A2"] = 3
    ws["B1"] = "=SUM(A1:A2)"
    ws["B2"] = "=XLOOKUP(3,A1:A2,A1:A2)"
    ws["B3"] = '=IF(A1=2,"<&> text","")'
    ws["B4"] = "=A1=2"
    ws["B5"] = "=1/0"
    ws["B6"] = '=IF(A1=2,"","")'
    ws["B7"] = "=UNSUPPORTED_FUNCTION(A1)"
    ws["B8"] = "=A1/4"
    other = wb.create_sheet("Summary & detail")
    other["A1"] = "=Sheet!B1*2"
    return wb


def _cells(source, part: str = SHEET1) -> dict[str, ET.Element]:
    if isinstance(source, BytesIO):
        source.seek(0)
    with zipfile.ZipFile(source) as package:
        root = ET.fromstring(package.read(part))
    return {cell.get("r"): cell for cell in root.findall(".//s:c", NS)}


def _cached(cell: ET.Element) -> str | None:
    value = cell.find("s:v", NS)
    return None if value is None else (value.text or "")


def _data_only(source, reader=wolfxl.load_workbook):
    if isinstance(source, BytesIO):
        source.seek(0)
    return reader(source, data_only=True)


def _formula_workbook() -> wolfxl.Workbook:
    wb = wolfxl.Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    for row in range(1, 4):
        inputs[f"A{row}"] = row * 10
    model = wb.create_sheet("Model")
    model["A1"] = "=SUM(Inputs!A1:A3)"
    model["A2"] = "=Model!A1*2"
    model["A3"] = '=IF(Model!A1>50,"big","small")'
    return wb


def _formula_source(tmp_path, name: str = "source.xlsx"):
    """A saved workbook whose formulas are not yet cached."""
    path = tmp_path / name
    _formula_workbook().save(path)
    return path


@pytest.mark.parametrize("mode", ["write", "read", "modify", "fileobj"])
def test_calculated_values_survive_save(tmp_path, mode):
    wb = _workbook()
    if mode in {"read", "modify"}:
        source = tmp_path / "source.xlsx"
        wb.save(source)
        wb = wolfxl.load_workbook(source, modify=mode == "modify")
    result = wb.calculate()
    assert result["Sheet!B1"] == 5
    assert result["Sheet!B2"] == 3
    output = BytesIO() if mode == "fileobj" else tmp_path / "calculated.xlsx"
    wb.save(output)

    for reader in READERS:
        saved = _data_only(output, reader)
        assert [saved["Sheet"][f"B{i}"].value for i in range(1, 6)] == [
            5,
            3,
            "<&> text",
            True,
            "#DIV/0!",
        ]
        assert saved["Sheet"]["B8"].value == 0.5
        assert saved["Summary & detail"]["A1"].value == 10
        assert saved["Sheet"]["B7"].value is None
        saved.close()
        if isinstance(output, BytesIO):
            output.seek(0)
        formulas = reader(output)
        assert formulas["Sheet"]["B1"].value == "=SUM(A1:A2)"
        assert formulas["Sheet"]["B3"].value == '=IF(A1=2,"<&> text","")'
        assert formulas["Sheet"]["B8"].value == "=A1/4"
        assert formulas["Summary & detail"]["A1"].value == "=Sheet!B1*2"
        formulas.close()

    cells = _cells(output)
    assert cells["B2"].find("s:f", NS).text == "_xlfn.XLOOKUP(3,A1:A2,A1:A2)"
    assert cells["B3"].get("t") == "str"
    assert cells["B4"].get("t") == "b"
    assert cells["B5"].get("t") == "e"
    assert _cached(cells["B5"]) == "#DIV/0!"
    # Empty-string formula caches are explicit strings, not unsupported results.
    # Both readers expose the empty XML text as None.
    assert cells["B6"].get("t") == "str"
    assert _cached(cells["B6"]) == ""
    assert _cached(cells["B7"]) is None


def test_edited_input_and_formula_text_survive_modify_calculate_save(tmp_path):
    source = _formula_source(tmp_path)
    wb = wolfxl.load_workbook(source, modify=True)
    wb["Inputs"]["A30"] = 7  # unused by every formula
    assert wb.calculate()["Model!A1"] == 60
    saved_path = tmp_path / "saved.xlsx"
    wb.save(saved_path)

    for reader in READERS:
        values = _data_only(saved_path, reader)
        assert values["Inputs"]["A30"].value == 7
        assert values["Model"]["A1"].value == 60
        assert values["Model"]["A2"].value == 120
        assert values["Model"]["A3"].value == "big"
        values.close()
        formulas = reader(saved_path)
        assert formulas["Inputs"]["A30"].value == 7
        assert formulas["Model"]["A1"].value == "=SUM(Inputs!A1:A3)"
        assert formulas["Model"]["A2"].value == "=Model!A1*2"
        assert formulas["Model"]["A3"].value == '=IF(Model!A1>50,"big","small")'
        formulas.close()


def test_in_place_modify_save_persists_caches(tmp_path):
    source = _formula_source(tmp_path)
    wb = wolfxl.load_workbook(source, modify=True)
    wb["Inputs"]["A1"] = 100
    wb.calculate()
    wb.save(source)
    saved = _data_only(source, openpyxl.load_workbook)
    assert saved["Model"]["A1"].value == 150
    assert saved["Model"]["A2"].value == 300
    assert saved["Model"]["A3"].value == "big"


def _stale_input(wb):
    wb["Inputs"]["A1"] = 99


def _stale_formula(wb):
    wb["Model"]["A2"] = "=Model!A1*3"


def _stale_sheet_name(wb):
    wb["Inputs"].title = "Renamed"


def _stale_defined_name(wb):
    wb.defined_names["Rate"] = DefinedName(name="Rate", value="Inputs!$A$1")


@pytest.mark.parametrize(
    ("mode", "edit"),
    [
        ("write", _stale_input),
        ("write", _stale_formula),
        ("write", _stale_sheet_name),
        ("write", _stale_defined_name),
        ("modify", _stale_input),
        ("modify", _stale_formula),
        ("modify", _stale_defined_name),
    ],
)
def test_edits_after_calculate_do_not_save_stale_values(tmp_path, mode, edit):
    if mode == "write":
        wb = _formula_workbook()
    else:
        wb = wolfxl.load_workbook(_formula_source(tmp_path), modify=True)
    wb.calculate()
    edit(wb)
    path = tmp_path / "edited.xlsx"
    wb.save(path)
    stale = _data_only(path, openpyxl.load_workbook)
    assert stale["Model"]["A1"].value is None
    assert stale["Model"]["A2"].value is None
    assert stale["Model"]["A3"].value is None
    stale.close()
    # A fresh calculation after the edit is valid again.
    wb.calculate()
    wb.save(path)
    fresh = _data_only(path, openpyxl.load_workbook)
    assert fresh["Model"]["A1"].value is not None
    assert fresh["Model"]["A2"].value is not None


def test_stale_save_drops_caches_inherited_from_an_earlier_save(tmp_path):
    source = _formula_source(tmp_path)
    first = wolfxl.load_workbook(source, modify=True)
    first.calculate()
    cached = tmp_path / "cached.xlsx"
    first.save(cached)
    assert _data_only(cached, openpyxl.load_workbook)["Model"]["A1"].value == 60

    wb = wolfxl.load_workbook(cached, modify=True)
    wb.calculate()
    wb["Inputs"]["A1"] = 1000
    edited = tmp_path / "edited.xlsx"
    wb.save(edited)
    saved = _data_only(edited, openpyxl.load_workbook)
    assert saved["Model"]["A1"].value is None
    assert saved["Inputs"]["A1"].value == 1000
    assert openpyxl.load_workbook(edited)["Model"]["A1"].value == "=SUM(Inputs!A1:A3)"


def test_what_if_recalculation_does_not_replace_saved_inputs(tmp_path):
    wb = _workbook()
    wb.calculate()
    wb.recalculate({"Sheet!A1": 100})
    path = tmp_path / "what-if.xlsx"
    wb.save(path)
    saved = _data_only(path, openpyxl.load_workbook)
    assert saved["Sheet"]["A1"].value == 2
    assert saved["Sheet"]["B1"].value == 5
    assert saved["Summary & detail"]["A1"].value == 10


def test_xlookup_prefix_is_idempotent_and_does_not_change_literals(tmp_path):
    wb = wolfxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = '=IF(TRUE,XLOOKUP(1,A1:A1,A1:A1),"XLOOKUP(")'
    ws["B2"] = "=_xlfn.XLOOKUP(1,A1:A1,A1:A1)"
    ws["B3"] = '=IF(TRUE,"XLOOKUP(","_xlfn.XLOOKUP(")'
    ws["B4"] = "=xlookup(1,A1:A1,A1:A1)"
    path = tmp_path / "prefix.xlsx"
    wb.calculate()
    wb.save(path)
    saved = openpyxl.load_workbook(path)
    assert saved.active["B1"].value == '=IF(TRUE,_xlfn.XLOOKUP(1,A1:A1,A1:A1),"XLOOKUP(")'
    assert saved.active["B2"].value == "=_xlfn.XLOOKUP(1,A1:A1,A1:A1)"
    assert saved.active["B3"].value == ws["B3"].value
    assert saved.active["B4"].value == "=_xlfn.XLOOKUP(1,A1:A1,A1:A1)"

    # The saved ``_xlfn.`` spelling is understood when calculating again.
    wb = wolfxl.load_workbook(path, modify=True)
    assert wb.calculate()["Sheet!B2"] == 1
    wb.save(path)
    assert _data_only(path, openpyxl.load_workbook).active["B2"].value == 1


def test_xlookup_is_qualified_without_a_calculation(tmp_path):
    wb = wolfxl.Workbook()
    wb.active["A1"] = 1
    wb.active["B1"] = "=XLOOKUP(1,A1:A1,A1:A1)"
    path = tmp_path / "plain.xlsx"
    wb.save(path)
    assert openpyxl.load_workbook(path).active["B1"].value == "=_xlfn.XLOOKUP(1,A1:A1,A1:A1)"
    assert _cached(_cells(path)["B1"]) is None


def test_shared_formula_caches_and_opaque_parts_survive(tmp_path):
    source = tmp_path / "shared.xlsx"
    wb = wolfxl.Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = 3
    wb.active["B1"] = "=A1*2"
    wb.active["B2"] = "=A2*2"
    wb.save(source)
    with zipfile.ZipFile(source) as package:
        parts = {info.filename: package.read(info.filename) for info in package.infolist()}
    sheet = parts[SHEET1].decode()
    sheet = sheet.replace("<f>A1*2</f>", '<f t="shared" si="0" ref="B1:B2">A1*2</f>')
    sheet = sheet.replace("<f>A2*2</f>", '<f t="shared" si="0"/>')
    assert 'si="0"/>' in sheet
    parts[SHEET1] = sheet.encode()
    parts["custom/opaque.bin"] = b"\x00opaque payload\xff"
    with zipfile.ZipFile(source, "w") as package:
        for name, payload in parts.items():
            package.writestr(name, payload)

    wb = wolfxl.load_workbook(source, modify=True)
    assert wb.calculate() == {"Sheet!B1": 4, "Sheet!B2": 6}
    wb.save(source)
    with zipfile.ZipFile(source) as package:
        assert package.read("custom/opaque.bin") == b"\x00opaque payload\xff"
    saved = _data_only(source, openpyxl.load_workbook)
    assert saved.active["B1"].value == 4
    assert saved.active["B2"].value == 6
    cells = _cells(source)
    assert cells["B2"].find("s:f", NS).get("t") == "shared"
    assert cells["B2"].find("s:f", NS).get("si") == "0"


def test_zip_metadata_and_foreign_parts_survive_calculated_copy(tmp_path):
    seed = tmp_path / "seed.xlsx"
    _workbook().save(seed)
    source = tmp_path / "source.xlsx"
    with zipfile.ZipFile(seed) as package:
        entries = [(info, package.read(info.filename)) for info in package.infolist()]
    opaque = zipfile.ZipInfo("custom/opaque.bin", date_time=(2020, 1, 2, 3, 4, 6))
    opaque.compress_type = zipfile.ZIP_STORED
    opaque.external_attr = 0o600 << 16
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as package:
        package.comment = b"keep this comment"
        for info, payload in entries:
            package.writestr(info, payload)
        package.writestr(opaque, b"\x00opaque payload\xff")

    wb = wolfxl.load_workbook(source)
    wb.calculate()
    output = tmp_path / "output.xlsx"
    wb.save(output)

    formula_sheets = {SHEET1, "xl/worksheets/sheet2.xml"}
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        assert after.comment == b"keep this comment"
        assert [i.filename for i in after.infolist()] == [i.filename for i in before.infolist()]
        for old, new in zip(before.infolist(), after.infolist()):
            assert new.date_time == old.date_time
            assert new.compress_type == old.compress_type
            assert new.external_attr == old.external_attr
            if old.filename in formula_sheets:
                assert after.read(new.filename) != before.read(old.filename)
            else:
                assert after.read(new.filename) == before.read(old.filename)
    saved = _data_only(output, openpyxl.load_workbook)
    assert saved["Sheet"]["B1"].value == 5
    assert saved["Summary & detail"]["A1"].value == 10


def test_self_closing_cell_does_not_swallow_the_next_formula(tmp_path):
    seed = tmp_path / "seed.xlsx"
    wb = wolfxl.Workbook()
    wb.active["A1"] = 5
    wb.active["B1"] = "=1+1"
    wb.active["C1"] = "=2*3"
    wb.save(seed)
    source = tmp_path / "source.xlsx"
    with zipfile.ZipFile(seed) as before, zipfile.ZipFile(source, "w") as after:
        for info in before.infolist():
            payload = before.read(info.filename)
            if info.filename == SHEET1:
                sheet = re.sub(
                    r'<c r="A1"[^>]*>.*?</c>', '<c r="A1" s="0"/>', payload.decode(), flags=re.S
                )
                assert '<c r="A1" s="0"/>' in sheet
                payload = sheet.encode()
            after.writestr(info, payload)

    wb = wolfxl.load_workbook(source)
    assert wb.calculate() == {"Sheet!B1": 2, "Sheet!C1": 6}
    output = tmp_path / "output.xlsx"
    wb.save(output)
    saved = _data_only(output, openpyxl.load_workbook)
    assert saved.active["A1"].value is None
    assert saved.active["B1"].value == 2
    assert saved.active["C1"].value == 6


def test_uncomputed_results_do_not_retain_stale_caches(tmp_path):
    seed = tmp_path / "seed.xlsx"
    wb = wolfxl.Workbook()
    wb.active["A1"] = 2
    wb.active["B1"] = "=SUM(A1:A1)"
    wb.active["B2"] = "=UNSUPPORTED_FUNCTION(A1)"
    wb.active["B3"] = "=A1*2"
    wb.save(seed)
    source = tmp_path / "source.xlsx"
    with zipfile.ZipFile(seed) as before, zipfile.ZipFile(source, "w") as after:
        for info in before.infolist():
            payload = before.read(info.filename)
            if info.filename == SHEET1:
                sheet = payload.decode()
                sheet = re.sub(
                    r'(<c r="B2"[^>]*>).*?(</c>)',
                    r"\g<1><f>UNSUPPORTED_FUNCTION(A1)</f><v>42</v>\g<2>",
                    sheet,
                    flags=re.S,
                )
                sheet = re.sub(
                    r'(<c r="B3"[^>]*>).*?(</c>)',
                    r'\g<1><f t="array" ref="B3">A1*2</f><v>99</v>\g<2>',
                    sheet,
                    flags=re.S,
                )
                payload = sheet.encode()
            after.writestr(info, payload)

    wb = wolfxl.load_workbook(source, modify=True)
    wb["Sheet"]["A1"] = 3
    wb["Sheet"]["D1"] = "edit"
    wb.calculate()
    output = tmp_path / "output.xlsx"
    wb.save(output)
    cells = _cells(output)
    assert _cached(cells["B1"]) == "3"
    assert _cached(cells["B2"]) is None
    assert _cached(cells["B3"]) is None
    assert cells["B2"].find("s:f", NS).text == "UNSUPPORTED_FUNCTION(A1)"
    assert cells["B3"].find("s:f", NS).get("t") == "array"
