"""Docling's staged ExcelReader consumer, isolated from other alias users.

Run with docling-slim[format-xlsx,convert-core]==2.142.1 and openpyxl==3.1.5.
The slim distribution contains the same Excel backend without PDF/ML models.
"""

from __future__ import annotations

import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import subprocess
import sys

import pytest

pytest.importorskip("docling", reason="Install the pinned Docling XLSX integration dependencies")

DOCLING_VERSION = "2.142.1"
OPENPYXL_VERSION = "3.1.5"
# Generated consumer fixtures, not copies of the excluded external-oracle files.
# Their labels describe the alias input shapes; native slicer, VBA and PowerPivot
# preservation are outside this Docling cell-value contract.
ALIAS_FIXTURES = (
    "pivot-source",
    "chart-conditional-format",
    "external-link",
    "model-source",
    "table-validation-image-comment",
    "chart-formula",
    "macro-container",
    "listbox-source",
    "shared-pivot-source",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _convert(arm: str, source: Path, stream: bool) -> dict:
    import resource
    from io import BytesIO
    from importlib.machinery import EXTENSION_SUFFIXES

    assert importlib.metadata.version("docling-slim") == DOCLING_VERSION
    assert importlib.metadata.version("openpyxl") == OPENPYXL_VERSION
    if arm == "alias":
        import wolfxl

        wolfxl.install_as_openpyxl()

    import openpyxl
    import docling.backend.msexcel_backend as excel_backend
    from docling.backend.msexcel_backend import MsExcelDocumentBackend
    from docling.datamodel.backend_options import MsExcelBackendOptions
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.document import InputDocument

    assert openpyxl.__name__ == ("wolfxl" if arm == "alias" else "openpyxl")
    assert excel_backend.ExcelReader.__module__ == (
        "wolfxl.reader.excel" if arm == "alias" else "openpyxl.reader.excel"
    )
    data = BytesIO(source.read_bytes()) if stream else source
    in_doc = InputDocument(
        path_or_stream=data,
        filename=source.name,
        format=InputFormat.XLSX,
        backend=MsExcelDocumentBackend,
        backend_options=MsExcelBackendOptions(render_chart_images=False),
    )
    assert in_doc.valid, in_doc._rejection
    backend = in_doc._backend
    try:
        document = backend.convert()
        cells = [
            {
                "page": table.prov[0].page_no,
                "bbox": table.prov[0].bbox.model_dump(mode="json"),
                "rows": table.data.num_rows,
                "columns": table.data.num_cols,
                "cells": [cell.model_dump(mode="json") for cell in table.data.table_cells],
            }
            for table in document.tables
        ]
        texts = [(text.label.value, text.text) for text in document.texts]
    finally:
        backend.unload()
    # Record the child lifetime peak before collecting build/source receipts.
    peak_rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    native = None
    if arm == "alias":
        import wolfxl._rust

        native = _sha256(Path(wolfxl._rust.__file__))
        assert any(str(wolfxl._rust.__file__).endswith(suffix) for suffix in EXTENSION_SUFFIXES)
    return {
        "cells": cells,
        "texts": texts,
        "pages": len(document.pages),
        "peak_rss_bytes": peak_rss * (1 if sys.platform == "darwin" else 1024),
        "native_sha256": native,
        "input_sha256": _sha256(source),
        "reader": {
            "module": excel_backend.ExcelReader.__module__,
            "sha256": _sha256(Path(sys.modules[excel_backend.ExcelReader.__module__].__file__)),
        },
        "backend_sha256": _sha256(Path(excel_backend.__file__)),
        "python": sys.version,
        "versions": {
            name: importlib.metadata.version(name)
            for name in ("docling-slim", "docling-core", "openpyxl", "wolfxl")
        },
    }


def _run(arm: str, source: Path, *, stream: bool = False) -> dict:
    result = subprocess.run(
        [sys.executable, str(Path(__file__).resolve()), arm, str(source), str(int(stream))],
        check=False,
        capture_output=True,
        text=True,
        timeout=180,
        env={**os.environ, "PYTHONHASHSEED": "0"},
    )
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)


def _assert_parity(source: Path, *, stream: bool = False) -> tuple[dict, dict]:
    upstream = _run("upstream", source, stream=stream)
    alias = _run("alias", source, stream=stream)
    assert alias["cells"] == upstream["cells"]
    assert alias["texts"] == upstream["texts"]
    assert alias["pages"] == upstream["pages"]
    return upstream, alias


def _alias_fixture(path: Path, case: str) -> None:
    from datetime import datetime
    from io import BytesIO
    from zipfile import ZipFile

    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table
    from PIL import Image as PillowImage

    assert openpyxl.__name__ == "openpyxl"
    assert importlib.metadata.version("openpyxl") == OPENPYXL_VERSION
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in (
        ("Region", "Amount", "Active", "Date"),
        ("West", 17, True, datetime(2026, 1, 2)),
        ("East", -2.5, False, datetime(2026, 2, 3)),
    ):
        sheet.append(row)
    if case in ("pivot-source", "shared-pivot-source"):
        summary = workbook.create_sheet("Summary")
        summary.append(("Region", "Total"))
        summary.append(("West", 17))
        summary.append(("East", -2.5))
        if case == "shared-pivot-source":
            workbook.copy_worksheet(summary).title = "Second summary"
    elif case in ("chart-conditional-format", "chart-formula"):
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
        sheet.add_chart(chart, "F1")
        sheet.conditional_formatting.add(
            "B2:B3",
            CellIsRule(
                operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FF0000")
            ),
        )
        if case == "chart-formula":
            sheet["A5"] = "Uncached total"
            sheet["B5"] = "=SUM(B2:B3)"
    elif case == "external-link":
        sheet["A5"] = "External formula"
        sheet["B5"] = "='[other.xlsx]Sheet1'!A1"
        sheet["A2"].hyperlink = "https://example.com/west"
    elif case == "model-source":
        related = workbook.create_sheet("Related")
        related.append(("Region", "Description"))
        related.append(("West", "Unicode \u03b1"))
        related.append(("East", "Unicode \u03b2"))
    elif case == "table-validation-image-comment":
        sheet.add_table(Table(displayName="Sales", ref="A1:D3"))
        validation = DataValidation(type="list", formula1='"West,East"')
        sheet.add_data_validation(validation)
        validation.add("A2:A3")
        sheet["B2"].comment = Comment("A generated comment", "Fixture")
        png = BytesIO()
        PillowImage.new("RGB", (1, 1), "white").save(png, format="PNG")
        png.seek(0)
        sheet.add_image(Image(png), "F1")
    elif case == "macro-container":
        # Seed the macro archive with openpyxl's own package metadata; no VBA
        # payload is needed for this macro-enabled cell-value fixture.
        package = BytesIO()
        workbook.save(package)
        workbook.vba_archive = ZipFile(package)
    elif case == "listbox-source":
        sheet["F1"] = "Choices"
        sheet["F2"] = "West"
        sheet["F3"] = "East"
        sheet.merge_cells("A5:D5")
        sheet["A5"] = "Merged selection"
        sheet["A6"] = "#DIV/0!"
    workbook.save(path)
    if workbook.vba_archive is not None:
        workbook.vba_archive.close()
    workbook.close()


@pytest.mark.parametrize("fixture", ALIAS_FIXTURES)
def test_docling_alias_fixture_cell_values(tmp_path: Path, fixture: str) -> None:
    source = tmp_path / (fixture + (".xlsm" if fixture == "macro-container" else ".xlsx"))
    _alias_fixture(source, fixture)
    upstream, alias = _assert_parity(source)
    # All fixtures contain a real table. Equality of two empty conversions is not
    # sufficient evidence that the staged reader exposed the stored cells.
    assert {"Region", "Amount", "West", "17"} <= {
        cell["text"] for table in alias["cells"] for cell in table["cells"]
    }
    print(
        "DOCLING_FIXTURE "
        + json.dumps(
            {
                "fixture": fixture,
                "input_sha256": upstream["input_sha256"],
                "tables": len(alias["cells"]),
                "parity": True,
            },
            sort_keys=True,
        )
    )


def _sparse_wide(path: Path) -> None:
    import openpyxl
    from openpyxl.styles import PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    for coordinate, value in {
        "A1": "left",
        "B1": "value",
        "A2": "alpha",
        "B2": 17,
        "BZQ63": "right",
        "BZR63": "value",
        "BZQ64": "omega",
        "BZR64": 29,
    }.items():
        sheet[coordinate] = value
    # A formatting-only far corner must not become a dense worksheet rectangle.
    sheet["XFD1048576"].fill = PatternFill("solid", fgColor="FFFF00")
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize("stream", [False, True], ids=["path", "bytesio"])
def test_docling_sparse_wide_cell_values(tmp_path: Path, stream: bool) -> None:
    source = tmp_path / "sparse-wide.xlsx"
    _sparse_wide(source)
    upstream, alias = _assert_parity(source, stream=stream)
    assert [cell["text"] for table in alias["cells"] for cell in table["cells"]] == [
        "left",
        "value",
        "alpha",
        "17",
        "right",
        "value",
        "omega",
        "29",
    ]
    print(
        "DOCLING_RSS "
        + json.dumps({"stream": stream, "upstream": upstream, "alias": alias}, sort_keys=True)
    )


if __name__ == "__main__":
    print(json.dumps(_convert(sys.argv[1], Path(sys.argv[2]), bool(int(sys.argv[3])))))
