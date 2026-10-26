"""Dates written in modify mode save as Excel serials with a date number format.

The workbook comes from openpyxl, which writes an empty ``<numFmts count="0"/>``
collection, so the date formats must be added to that self-closing section.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

import wolfxl


def _make_workbook(path: Path, *, epoch_1904: bool = False) -> None:
    wb = openpyxl.Workbook()
    if epoch_1904:
        wb.epoch = openpyxl.utils.datetime.MAC_EPOCH
    wb.active.title = "Here"
    wb.active["A1"] = "when"
    wb.create_sheet("Keep")["A1"] = "keep me"
    wb.save(path)


@pytest.mark.parametrize("epoch_1904", [False, True], ids=["epoch-1900", "epoch-1904"])
def test_dates_round_trip_with_date_formats(tmp_path: Path, epoch_1904: bool) -> None:
    path = tmp_path / "dates.xlsx"
    _make_workbook(path, epoch_1904=epoch_1904)

    wb = wolfxl.load_workbook(path, modify=True)
    ws = wb["Here"]
    ws["A2"] = dt.date(2004, 4, 30)
    ws["A3"] = dt.datetime(2004, 5, 31, 13, 45, 10)
    ws.cell(row=4, column=1, value=dt.date(1904, 3, 1))
    ws.cell(row=5, column=1, value=dt.datetime(2026, 10, 25, 9, 0))
    wb.save(path)
    wb.close()

    back = openpyxl.load_workbook(path)
    cells = [row[0] for row in back["Here"]["A2:A5"]]
    assert [cell.value for cell in cells] == [
        dt.datetime(2004, 4, 30),
        dt.datetime(2004, 5, 31, 13, 45, 10),
        dt.datetime(1904, 3, 1),
        dt.datetime(2026, 10, 25, 9, 0),
    ]
    assert [cell.number_format for cell in cells] == [
        "yyyy-mm-dd",
        "yyyy-mm-dd h:mm:ss",
        "yyyy-mm-dd",
        "yyyy-mm-dd h:mm:ss",
    ]
    assert back["Keep"]["A1"].value == "keep me"


def test_timezone_aware_datetime_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "tz.xlsx"
    _make_workbook(path)

    wb = wolfxl.load_workbook(path, modify=True)
    wb["Here"]["A2"] = dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc)
    with pytest.raises(TypeError, match="timezones"):
        wb.save(path)
    wb.close()
