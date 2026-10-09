"""The source-negative probe avoids cell parsing and retains the exact API boundary."""

from pathlib import Path
import zipfile

import openpyxl
import pytest

import wolfxl
from wolfxl._rust import NativeXlsxBook


def _source(tmp_path: Path, xml: bytes | None = None) -> Path:
    original = tmp_path / "original.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(original)
    if xml is None:
        return original
    modified = tmp_path / "modified.xlsx"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(modified, "w") as output:
        for info in source.infolist():
            output.writestr(
                info, xml if info.filename == "xl/worksheets/sheet1.xml" else source.read(info)
            )
    return modified


def test_probe_negative_does_not_cache_away_exact_xml_errors(tmp_path: Path) -> None:
    path = _source(tmp_path, b"<worksheet><sheetData></worksheet>")
    reader = NativeXlsxBook.open(str(path))
    assert reader.read_merged_ranges_if_present("Sheet") == []
    with pytest.raises(OSError, match="native merge metadata read failed"):
        reader.read_merged_ranges("Sheet")
    with pytest.raises(ValueError, match="Unknown sheet"):
        reader.read_merged_ranges_if_present("missing")


@pytest.mark.parametrize(
    "false_positive",
    [
        b"<!-- mergeCell ref='A1:B2' -->",
        b"<![CDATA[<mergeCell ref='A1:B2'/>]]>",
        b"<t>mergeCell</t>",
    ],
)
def test_conservative_false_positive_is_parsed_exactly(
    tmp_path: Path, false_positive: bytes
) -> None:
    path = _source(tmp_path, b"<worksheet>" + false_positive + b"</worksheet>")
    reader = NativeXlsxBook.open(str(path))
    assert reader.read_merged_ranges_if_present("Sheet") == []
    assert reader.read_merged_ranges("Sheet") == []


def test_positive_probe_keeps_parser_error_boundary(tmp_path: Path) -> None:
    path = _source(tmp_path, b"<worksheet><!-- mergeCell --><sheetData></worksheet>")
    reader = NativeXlsxBook.open(str(path))
    with pytest.raises(OSError, match="native merge metadata read failed"):
        reader.read_merged_ranges_if_present("Sheet")


def test_namespaced_merge_probe_returns_exact_refs(tmp_path: Path) -> None:
    xml = b"<m:worksheet xmlns:m='urn:test'><m:mergeCells><m:mergeCell ref='A1:B2'/></m:mergeCells></m:worksheet>"
    reader = NativeXlsxBook.open(str(_source(tmp_path, xml)))
    assert reader.read_merged_ranges_if_present("Sheet") == ["A1:B2"]


def test_lazy_python_hydration_uses_probe_once_and_keeps_compact_edits(tmp_path: Path) -> None:
    wb = wolfxl.load_workbook(_source(tmp_path), modify=True)
    reader = wb._rust_reader
    calls = []

    class Proxy:
        def __getattr__(self, name):
            target = getattr(reader, name)
            if name != "read_merged_ranges_if_present":
                return target

            def recorded(sheet):
                calls.append(sheet)
                return target(sheet)

            return recorded

    wb._rust_reader = Proxy()
    try:
        ws = wb.active
        assert ws._merged_ranges_loaded is False
        ws["C2"] = "edited"
        assert ws._merged_ranges_loaded is True
        assert calls == ["Sheet"]
        assert (2, 3) not in ws._dirty
        assert ws._dirty_values[(2, 3)] == "edited"
        assert list(ws.merged_cells.ranges) == []
        ws["D3"] = "second edit"
        assert calls == ["Sheet"]
    finally:
        wb._rust_reader = reader
        wb.close()


@pytest.mark.parametrize("bytes_backed", [False, True])
@pytest.mark.parametrize("case_variant", [False, True])
def test_probe_resolves_path_bytes_and_case_variant_sheet_parts(
    tmp_path: Path, bytes_backed: bool, case_variant: bool
) -> None:
    path = _source(tmp_path)
    if case_variant:
        changed = tmp_path / "case-variant.xlsx"
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(changed, "w") as output:
            for info in source.infolist():
                name = (
                    info.filename.upper()
                    if info.filename.startswith("xl/worksheets/")
                    else info.filename
                )
                output.writestr(name, source.read(info))
        path = changed
    reader = (
        NativeXlsxBook.open_from_bytes(path.read_bytes())
        if bytes_backed
        else NativeXlsxBook.open(str(path))
    )
    assert reader.read_merged_ranges_if_present("Sheet") == []
    assert reader.read_merged_ranges("Sheet") == []


def test_hydration_keeps_legacy_and_xlsb_reader_dispatch() -> None:
    from types import SimpleNamespace

    from wolfxl._worksheet_collections import _loaded_merged_range_refs

    calls = []

    class LegacyReader:
        def read_merged_ranges(self, sheet):
            calls.append(sheet)
            return ["A1:B2"]

    ws = SimpleNamespace(
        _title="Legacy",
        _merged_ranges_loaded=False,
        _merged_ranges=set(),
        _collection_merged_ranges=set(),
        _pending_unmerged_ranges=set(),
        _workbook=SimpleNamespace(_rust_reader=LegacyReader(), _format="xlsb"),
    )
    assert _loaded_merged_range_refs(ws) == {"A1:B2"}
    assert _loaded_merged_range_refs(ws) == {"A1:B2"}
    assert calls == ["Legacy"]
