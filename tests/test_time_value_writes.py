"""Durations and times of day write as openpyxl writes them: a serial day count
under ``[hh]:mm:ss`` or ``h:mm:ss``. Timezone-aware values are rejected, as in
openpyxl."""

from __future__ import annotations

from datetime import datetime, time, timedelta, timezone
from pathlib import Path

import openpyxl
import pytest

import wolfxl

VALUES = [
    pytest.param(timedelta(hours=27, minutes=30), "[hh]:mm:ss", id="timedelta"),
    pytest.param(time(10, 30, 15, 250000), "h:mm:ss", id="time"),
]

_TZ = timezone(timedelta(hours=5))
AWARE_VALUES = [
    pytest.param(datetime(2024, 1, 15, 10, 30, tzinfo=_TZ), id="datetime"),
    pytest.param(datetime(2024, 1, 15, 10, 30, tzinfo=timezone.utc), id="utc-datetime"),
    pytest.param(time(10, 30, tzinfo=_TZ), id="time"),
]


def _written(path: Path) -> tuple[object, str, str]:
    book = openpyxl.load_workbook(path)
    try:
        cell = book.active["A1"]
        return cell.value, cell.data_type, cell.number_format
    finally:
        book.close()


def _append(book, value):
    book.active.append((value, 2))


def _assign(book, value):
    book.active["A1"] = value


def _write_rows(book, value):
    book.active.write_rows([[value, 2]])


@pytest.mark.parametrize(("value", "number_format"), VALUES)
@pytest.mark.parametrize("write", [_append, _assign, _write_rows])
def test_time_values_write_as_openpyxl_writes_them(
    tmp_path: Path, write, value, number_format
) -> None:
    reference = openpyxl.Workbook()
    reference.active["A1"] = value
    reference.save(tmp_path / "reference.xlsx")

    book = wolfxl.Workbook()
    write(book, value)
    book.save(tmp_path / "actual.xlsx")

    assert _written(tmp_path / "actual.xlsx") == _written(tmp_path / "reference.xlsx")
    assert _written(tmp_path / "actual.xlsx") == (value, "d", number_format)


@pytest.mark.parametrize(("value", "number_format"), VALUES)
def test_write_only_time_values_keep_their_format(tmp_path: Path, value, number_format) -> None:
    book = wolfxl.Workbook(write_only=True)
    book.create_sheet().append([value])
    book.save(tmp_path / "write-only.xlsx")

    assert _written(tmp_path / "write-only.xlsx") == (value, "d", number_format)


@pytest.mark.parametrize(("value", "number_format"), VALUES)
def test_modified_workbooks_write_time_values(tmp_path: Path, value, number_format) -> None:
    source = openpyxl.Workbook()
    source.active["B1"] = 1
    source.save(tmp_path / "source.xlsx")

    book = wolfxl.load_workbook(tmp_path / "source.xlsx")
    book.active["A1"] = value
    book.save(tmp_path / "modified.xlsx")
    book.close()

    assert _written(tmp_path / "modified.xlsx") == (value, "d", number_format)


def _assign_beside_plain(book, value):
    book.active["B1"] = 2
    book.active["A1"] = value


@pytest.mark.parametrize("value", AWARE_VALUES)
@pytest.mark.parametrize("write", [_append, _assign, _assign_beside_plain, _write_rows])
def test_timezone_aware_values_fail_on_save(tmp_path: Path, write, value) -> None:
    book = wolfxl.Workbook()
    write(book, value)

    with pytest.raises(TypeError, match="Excel does not support timezones"):
        book.save(tmp_path / "aware.xlsx")


@pytest.mark.parametrize("value", AWARE_VALUES)
def test_write_only_timezone_aware_values_fail_on_append(value) -> None:
    sheet = wolfxl.Workbook(write_only=True).create_sheet()

    with pytest.raises(TypeError, match="Excel does not support timezones"):
        sheet.append([value])


def test_copied_duration_rows_stay_durations(tmp_path: Path) -> None:
    source = openpyxl.Workbook()
    source.active["A1"] = 1.25
    source.active["A1"].number_format = "[h]:mm:ss"
    source.active["B1"] = 2
    source.save(tmp_path / "source.xlsx")

    loaded = wolfxl.load_workbook(tmp_path / "source.xlsx")
    copy = wolfxl.Workbook()
    try:
        for row in loaded.active.iter_rows(values_only=True):
            copy.active.append(row)
    finally:
        loaded.close()
    copy.save(tmp_path / "copy.xlsx")

    assert _written(tmp_path / "copy.xlsx") == (timedelta(days=1, hours=6), "d", "[hh]:mm:ss")
