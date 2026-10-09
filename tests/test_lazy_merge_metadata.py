"""Merged metadata must remain sparse and preserve edit/reopen semantics."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import zipfile

import openpyxl
from openpyxl.styles import Border, Font, Side
import pytest

import wolfxl


def _fixture(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Merged"
    ws["A1"] = "anchor"
    ws["A1"].font = Font(bold=True)
    ws["A1"].border = Border(left=Side(style="thin"), top=Side(style="double"))
    ws["C3"].border = Border(right=Side(style="thick"), bottom=Side(style="dashed"))
    ws.merge_cells("A1:C3")
    ws["E1"] = "other"
    ws.merge_cells("E1:F2")
    wb.create_sheet("Plain")["A1"] = 42
    wb.save(path)
    return path


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("bytes_backed", [False, True])
def test_open_reads_no_python_worksheet_xml(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, read_only: bool, bytes_backed: bool
) -> None:
    path = _fixture(tmp_path / "merged.xlsx")
    original = zipfile.ZipFile.open

    def guard(self, name, *args, **kwargs):
        filename = name.filename if isinstance(name, zipfile.ZipInfo) else name
        assert not str(filename).startswith("xl/worksheets/"), "worksheet XML scanned eagerly"
        return original(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "open", guard)
    source = BytesIO(path.read_bytes()) if bytes_backed else path
    wb = wolfxl.load_workbook(source, read_only=read_only)
    try:
        assert wb.active._merged_ranges_loaded is False
        assert wb.active._cells == {}
        assert set(wb._rust_reader.read_merged_ranges("Merged")) == {"A1:C3", "E1:F2"}
        assert wb._rust_reader.read_merged_ranges("Plain") == []
        assert wb._rust_reader.read_merged_cell_border("Merged", 4, 4) is None
    finally:
        wb.close()


def _border_sides(border):
    return tuple(
        getattr(border, name).style if getattr(border, name) else None
        for name in ("left", "right", "top", "bottom")
    )


def test_lazy_native_merge_border_matches_anchor_and_edges(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "merged.xlsx")
    expected = openpyxl.load_workbook(path)
    actual = wolfxl.load_workbook(path)
    try:
        # Native endpoint-only access must also work before any Python Cell exists.
        assert (
            actual._rust_reader.read_merged_cell_border("Merged", 1, 1)["right"]["style"] == "thick"
        )
        assert actual.active._cells == {}
        for coordinate in ("A1", "A2", "B1", "C2", "B3", "C3", "B2"):
            assert _border_sides(actual.active[coordinate].border) == _border_sides(
                expected.active[coordinate].border
            )
        assert actual.active["A1"].font.bold is True
        assert actual.active["B2"].value is None
    finally:
        actual.close()
        expected.close()


def test_merge_metadata_cache_repeated_access_and_unknown_sheet(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "merged.xlsx")
    wb = wolfxl.load_workbook(path)
    try:
        reader = wb._rust_reader
        first = reader.read_merged_ranges("Merged")
        assert reader.read_merged_ranges("Merged") == first
        assert reader.read_merged_endpoint_style_ids(
            "Merged"
        ) == reader.read_merged_endpoint_style_ids("Merged")
        with pytest.raises(ValueError, match="Unknown sheet"):
            reader.read_merged_ranges("missing")
    finally:
        wb.close()


def test_source_merges_survive_mutation_and_repeated_save(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "merged.xlsx")
    wb = wolfxl.load_workbook(path)
    try:
        ws = wb.active
        ws.unmerge_cells("A1:C3")
        ws.merge_cells("G1:H2")
        assert {str(ref) for ref in ws.merged_cells.ranges} == {"E1:F2", "G1:H2"}
        ws["B2"] = "released"
        ws["G1"] = "new"
        for index in range(2):
            output = tmp_path / f"saved-{index}.xlsx"
            wb.save(output)
            reopened = openpyxl.load_workbook(output)
            try:
                assert {str(ref) for ref in reopened.active.merged_cells.ranges} == {
                    "E1:F2",
                    "G1:H2",
                }
                assert reopened.active["B2"].value == "released"
                assert reopened.active["G1"].value == "new"
            finally:
                reopened.close()
    finally:
        wb.close()


def test_merge_metadata_resolves_case_variant_worksheet_parts(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "original.xlsx")
    path = tmp_path / "case-variant.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, "w") as changed:
        for info in original.infolist():
            name = (
                info.filename.upper()
                if info.filename.startswith("xl/worksheets/")
                else info.filename
            )
            changed.writestr(name, original.read(info.filename))
    wb = wolfxl.load_workbook(path)
    try:
        assert set(wb._rust_reader.read_merged_ranges("Merged")) == {"A1:C3", "E1:F2"}
        assert wb._rust_reader.read_merged_cell_border("Merged", 1, 1)["right"]["style"] == "thick"
    finally:
        wb.close()


def _unmerged_styled_fixture(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Unmerged"
    ws["A1"] = "anchor"
    ws["A1"].border = Border(left=Side(style="thin", color="FF123456"), top=Side(style="double"))
    ws["B2"].border = Border(right=Side(style="thick"), bottom=Side(style="dashed"))
    wb.save(path)
    return path


@pytest.mark.parametrize("preaccess", [False, True])
@pytest.mark.parametrize("unrelated_style_edit", [False, True])
def test_new_live_merge_keeps_lazy_anchor_and_endpoint_styles(
    tmp_path: Path, preaccess: bool, unrelated_style_edit: bool
) -> None:
    path = _unmerged_styled_fixture(tmp_path / "unmerged.xlsx")
    expected = openpyxl.load_workbook(path)
    expected.active.merge_cells("A1:B2")
    wb = wolfxl.load_workbook(path)
    try:
        if preaccess:
            assert wb.active["A1"].border.left.style == "thin"
        if unrelated_style_edit:
            wb.active["A1"].font = wolfxl.styles.Font(italic=True)
        wb.active.merge_cells("A1:B2")
        assert _border_sides(wb.active["A1"].border) == _border_sides(expected.active["A1"].border)
        assert _border_sides(wb.active["B2"].border) == _border_sides(expected.active["B2"].border)
        saved = tmp_path / f"new-merge-{preaccess}.xlsx"
        wb.save(saved)
        reopened = openpyxl.load_workbook(saved)
        try:
            assert _border_sides(reopened.active["A1"].border) == _border_sides(
                expected.active["A1"].border
            )
        finally:
            reopened.close()
    finally:
        wb.close()
        expected.close()


def test_merged_border_color_does_not_alias_authoring_style_table(tmp_path: Path) -> None:
    path = _unmerged_styled_fixture(tmp_path / "colors.xlsx")
    wb = wolfxl.load_workbook(path, read_only=True)
    # Populate a live read view without mutating its authoring source.
    wb.active._merged_ranges = {"A1:B2"}
    wb.active._merged_ranges_loaded = True
    try:
        original = [
            border.left.color.rgb if border.left.color is not None else None
            for border in wb._borders
        ]
        border = wb.active["A1"].border
        border.left.color.rgb = "FFABCDEF"
        assert [
            border.left.color.rgb if border.left.color is not None else None
            for border in wb._borders
        ] == original
    finally:
        wb.close()


def test_close_releases_merge_cache_and_reader_reference(tmp_path: Path) -> None:
    import weakref

    class ReaderProxy:
        def __init__(self, native):
            self.native = native

        def __getattr__(self, name):
            return getattr(self.native, name)

    path = _fixture(tmp_path / "close.xlsx")
    wb = wolfxl.load_workbook(path)
    proxy = ReaderProxy(wb._rust_reader)
    reference = weakref.ref(proxy)
    wb._rust_reader = proxy
    del proxy
    assert wb.active["B1"].border.top.style == "double"
    assert wb._merged_border_metadata_cache is not None
    wb.close()
    assert wb._merged_border_metadata_cache is None
    assert reference() is None
    wb.close()


def test_live_merge_keeps_explicitly_authored_anchor_border(tmp_path: Path) -> None:
    path = _unmerged_styled_fixture(tmp_path / "authored.xlsx")
    wb = wolfxl.load_workbook(path)
    try:
        assert wb.active["A1"].border.left.style == "thin"
        authored = wolfxl.styles.Border(
            left=wolfxl.styles.Side(style="dotted"),
            right=wolfxl.styles.Side(style="double"),
            top=wolfxl.styles.Side(style="dashDot"),
            bottom=wolfxl.styles.Side(style="hair"),
        )
        wb.active["A1"].border = authored
        wb.active["A1"].font = wolfxl.styles.Font(italic=True)
        wb.active.merge_cells("A1:B2")
        assert wb.active["A1"].border == authored
        saved = tmp_path / "authored-merged.xlsx"
        wb.save(saved)
        reopened = openpyxl.load_workbook(saved)
        try:
            assert _border_sides(reopened.active["A1"].border) == _border_sides(authored)
        finally:
            reopened.close()
    finally:
        wb.close()
