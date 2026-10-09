from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.styles import Font

import wolfxl


def _parts(path: Path) -> dict[str, bytes]:
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _without_targets(xml: bytes) -> bytes:
    # This fixture has two ordinary, non-nested cells. Removing their complete
    # spans lets the assertion cover every other worksheet byte in one check.
    for coordinate in (b"B3", b"C3"):
        pattern = rb'<c\b[^>]*\br="' + coordinate + rb'"[^>]*>.*?</c>'
        xml, count = re.subn(pattern, b"", xml, flags=re.DOTALL)
        assert count == 1
    return xml


def test_late_value_edits_preserve_xml_and_other_parts_across_two_saves(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    original = openpyxl.Workbook()
    worksheet = original.active
    worksheet.title = "Data"
    worksheet.append(["Formula", "Text", "Number"])
    worksheet.append(["=SUM(C2:C3)", "kept", 2])
    worksheet.append(["styled", "old", 3])
    worksheet["A3"].font = Font(bold=True, color="FF123456")
    original.create_sheet("Untouched")["A1"] = "Keep this sheet"
    original.save(source)
    original.close()

    parts = _parts(source)
    sheet_path = "xl/worksheets/sheet1.xml"
    parts[sheet_path] = parts[sheet_path].replace(
        b"<sheetData>", b"<sheetData>\n<!-- preserve prefix -->\n", 1
    ).replace(b"</sheetData>", b"\n<!-- preserve tail -->\n</sheetData>", 1)
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    original_source = source.read_bytes()

    workbook = wolfxl.load_workbook(source, modify=True)
    workbook["Data"]["B3"] = "changed"
    workbook["Data"]["C3"] = 12345
    workbook.save(first)
    first_bytes = first.read_bytes()
    workbook["Data"]["B3"] = "second"
    workbook["Data"]["C3"] = 24690
    workbook.save(second)
    workbook.close()

    for output, expected in ((first, ("changed", 12345)), (second, ("second", 24690))):
        after = _parts(output)
        assert after.keys() == parts.keys()
        assert {
            name for name in parts if parts[name] != after[name]
        } == {sheet_path}
        assert _without_targets(after[sheet_path]) == _without_targets(parts[sheet_path])
        reopened = openpyxl.load_workbook(output)
        assert (reopened["Data"]["B3"].value, reopened["Data"]["C3"].value) == expected
        assert reopened["Data"]["A2"].value == "=SUM(C2:C3)"
        assert reopened["Data"]["A3"].font.bold
        assert reopened["Untouched"]["A1"].value == "Keep this sheet"
        reopened.close()
    assert source.read_bytes() == original_source
    assert first.read_bytes() == first_bytes
