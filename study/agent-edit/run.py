"""Same one-cell edit: the reference recipe (openpyxl + recalc.py) vs WolfXL modify mode.

Each output is compared with its source for:
- package parts removed (xl/calcChain.xml excluded; Excel rebuilds it),
- primary parts per feature family (rels files are not counted as features),
- sheet features across all worksheet parts, including a digest of the x14
  extension elements so a changed extension is not mistaken for an intact one,
- cached-value regressions: formula cells whose source cached value was not an
  error and whose output cached value is an error.

Usage: python run.py [skill|wolfxl ...]   (default: both workflows)
"""

import hashlib
import json
import posixpath
import re
import shutil
import subprocess
import sys
import warnings
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET

import openpyxl

from fetch import (
    CONTOSO_URL,
    RECALC,
    ROOT,
    SKILLS_COMMIT,
    fetch_contoso,
    fetch_reference,
)

FIX = ROOT.parents[1] / "tests/fixtures/external_oracle"
PUBLIC_SOURCES = [
    FIX / "real-excel-pivot-chart-slicers.xlsx",
    FIX / "real-excel-timeline-slicer.xlsx",
    FIX / "real-excel-normalized-pivot-cf-table.xlsx",
    FIX / "real-excel-p1-comments-validation-protection.xlsx",
    FIX / "real-excel-chart-cf-basic.xlsx",
    FIX / "real-excel-macro-basic.xlsm",
    ROOT / "x14-validation-sparkline.xlsx",
]


def sources():
    contoso = fetch_contoso()
    return PUBLIC_SOURCES[:5] + ([contoso] if contoso else []) + PUBLIC_SOURCES[5:]


def recalculate(out):
    result = subprocess.run(
        [sys.executable, str(RECALC), str(out), "90"],
        capture_output=True,
        text=True,
        cwd=RECALC.parent,
    )
    try:
        recalc = json.loads(result.stdout)
    except ValueError:
        recalc = {"error": (result.stdout + result.stderr)[-400:]}
    if result.returncode and "error" not in recalc:
        recalc["error"] = f"Reference recalculator exited {result.returncode}"
    return recalc


def failed(row):
    return (
        "error" in row
        or "error" in row.get("recalc", {})
        or row.get("cached_value") != 2
    )


CELL, FORMULA = "Z1", "=1+1"
IGNORED_PARTS = {"xl/calcChain.xml"}
# One primary part per feature instance; relationship parts are excluded.
FAMILIES = {
    "vba": r"xl/vbaProject\.bin",
    "pivot": r"xl/pivotTables/pivotTable\d+\.xml",
    "pivot_cache": r"xl/pivotCache/pivotCacheDefinition\d+\.xml",
    "chart": r"xl/charts/chart\d+\.xml",
    "slicer": r"xl/slicers/slicer\d+\.xml",
    "slicer_cache": r"xl/slicerCaches/slicerCache\d+\.xml",
    "timeline": r"xl/timelines/timeline\d+\.xml",
    "timeline_cache": r"xl/timelineCaches/timelineCache\d+\.xml",
    "ext_link": r"xl/externalLinks/externalLink\d+\.xml",
    "drawing": r"xl/drawings/drawing\d+\.xml",
    "table": r"xl/tables/table\d+\.xml",
    "comment": r"xl/comments\d+\.xml",
    "power_pivot": r"xl/model/item\.data",
}
FAMILY_RES = {k: re.compile(v) for k, v in FAMILIES.items()}

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
ERRORS = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
    "#CONNECT!",
    "#BUSY!",
    "#UNKNOWN!",
    "#PYTHON!",
}


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}


def families(names):
    return {k: sum(1 for n in names if r.fullmatch(n)) for k, r in FAMILY_RES.items()}


def worksheet_parts(z):
    """Sheet name -> worksheet part path, from the workbook relationships."""
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    names = set(z.namelist())
    out = {}
    for s in wb.iter(f"{{{MAIN}}}sheet"):
        t = targets.get(s.get(f"{{{REL}}}id"), "")
        p = t.lstrip("/") if t.startswith("/") else posixpath.normpath("xl/" + t)
        if p.startswith("xl/worksheets/") and p in names:
            out[s.get("name")] = p
    return out


def scan(path):
    """Sheet features across all worksheet parts, and every formula cell's cached value."""
    counts = Counter()
    x14 = []
    caches = {}
    with zipfile.ZipFile(path) as z:
        for sheet, part in worksheet_parts(z).items():
            with z.open(part) as fh:
                for _, el in ET.iterparse(fh, events=("end",)):
                    ns, _, local = el.tag[1:].partition("}")
                    if ns == MAIN and local == "c":
                        f = el.find(f"{{{MAIN}}}f")
                        if f is not None:
                            v = el.find(f"{{{MAIN}}}v")
                            caches[(sheet, el.get("r"))] = (
                                el.get("t"),
                                None if v is None else v.text,
                            )
                        el.clear()
                    elif ns == MAIN and local in (
                        "conditionalFormatting",
                        "dataValidation",
                    ):
                        counts[local] += 1
                    elif ns in (X14, XM):
                        if ns == X14:
                            counts["x14"] += 1
                        attrs = sorted(
                            (k, v) for k, v in el.attrib.items() if k != "id"
                        )
                        text = (el.text or "").strip() if ns == XM else ""
                        x14.append(f"{local} {attrs} {text}")
    features = {
        "x14_ext": counts["x14"],
        "cf": counts["conditionalFormatting"],
        "dv": counts["dataValidation"],
        "x14_digest": hashlib.sha256("\n".join(sorted(x14)).encode()).hexdigest()[:16],
    }
    return features, caches


def is_error(cached):
    t, v = cached
    return t == "e" or (v or "").strip() in ERRORS


def cached_value_regressions(src_caches, out_caches):
    """Formula cells whose cached value was a non-error and became an error."""
    bad = []
    for key, before in src_caches.items():
        after = out_caches.get(key)
        if before[1] is None or is_error(before) or after is None:
            continue
        if is_error(after):
            bad.append(
                {
                    "sheet": key[0],
                    "cell": key[1],
                    "before": before[1],
                    "after": after[1],
                }
            )
    return bad


def cached(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True, keep_vba=path.suffix == ".xlsm")
    v = wb[sheet][CELL].value
    wb.close()
    return v


def summarize(src, out, sheet, extra):
    a, b = parts(src), parts(out)
    fa, fb = families(a), families(b)
    feat_a, caches_a = scan(src)
    feat_b, caches_b = scan(out)
    regressions = cached_value_regressions(caches_a, caches_b)
    return {
        **extra,
        "parts_missing": sorted(set(a) - set(b) - IGNORED_PARTS),
        "parts_unchanged": sum(1 for n in a if b.get(n) == a[n]),
        "families_lost": {k: [fa[k], fb[k]] for k in fa if fb[k] < fa[k]},
        "sheet_features": [feat_a, feat_b],
        "sheet_features_changed": sorted(k for k in feat_a if feat_a[k] != feat_b[k]),
        "formula_cells_checked": len(caches_a),
        "cached_value_regressions": len(regressions),
        "cached_value_regression_samples": regressions[:10],
        "cached_value": cached(out, sheet),
    }


def skill_workflow(src, out):
    shutil.copy(src, out)
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(out, keep_vba=out.suffix == ".xlsm")
        sheet = wb.sheetnames[0]
        wb[sheet][CELL] = FORMULA
        wb.save(out)
        wb.close()
    edit_step = {
        "parts_missing": sorted(set(parts(src)) - set(parts(out)) - IGNORED_PARTS),
        "sheet_features": scan(out)[0],
    }
    recalc = recalculate(out)
    return sheet, {
        "openpyxl_warnings": sorted({str(x.message) for x in w}),
        "edit_step": edit_step,
        "recalc": recalc,
    }


def wolfxl_workflow(src, out):
    import wolfxl

    shutil.copy(src, out)
    wb = wolfxl.load_workbook(out, modify=True, keep_vba=out.suffix == ".xlsm")
    sheet = wb.sheetnames[0]
    wb[sheet][CELL] = FORMULA
    report = wb.calculate()
    wb.save(out)
    wb.close()
    return sheet, {"calculate": type(report).__name__}


WORKFLOWS = {"skill": skill_workflow, "wolfxl": wolfxl_workflow}


def environment():
    """Tool identities recorded with every run."""

    def out(cmd):
        r = subprocess.run(cmd, capture_output=True, text=True, check=False)
        return r.stdout.strip() or r.stderr.strip()

    return {
        "openpyxl": openpyxl.__version__,
        "libreoffice": out(["soffice", "--version"]),
        "anthropics_skills": SKILLS_COMMIT,
        "python": sys.version.split()[0],
    }


def main(selected):
    unknown = set(selected) - set(WORKFLOWS)
    if unknown:
        raise SystemExit(f"Unknown workflow(s): {', '.join(sorted(unknown))}")
    if "skill" in selected:
        fetch_reference()
    inputs = sources()
    out_dir = ROOT / "out"
    out_dir.mkdir(exist_ok=True)
    results = {
        **environment(),
        "edit": f"{CELL} {FORMULA}",
        "workflows": selected,
        "files": {},
        "source_sha256": {
            src.name: hashlib.sha256(src.read_bytes()).hexdigest() for src in inputs
        },
        "contoso_url": CONTOSO_URL,
        "contoso_available": len(inputs) == 8,
    }
    if "wolfxl" in selected:
        import wolfxl

        results["wolfxl"] = wolfxl.__version__
    for src in inputs:
        row = {}
        for name in selected:
            out = out_dir / f"{src.stem}.{name}{src.suffix}"
            try:
                sheet, extra = WORKFLOWS[name](src, out)
                row[name] = summarize(src, out, sheet, extra)
            except Exception as e:  # recorded, not hidden
                row[name] = {"error": f"{type(e).__name__}: {e}"[:400]}
        results["files"][src.name] = row
    (ROOT / "results.json").write_text(json.dumps(results, indent=2, default=str))
    for f, row in results["files"].items():
        for name, r in row.items():
            if "error" in r:
                print(f"{f:55} {name:6} ERROR {r['error'][:120]}")
                continue
            print(
                f"{f:55} {name:6} missing={len(r['parts_missing']):3} lost={r['families_lost']} "
                f"changed={r['sheet_features_changed']} regressions={r['cached_value_regressions']} "
                f"cached={r['cached_value']}"
            )
    return int(
        any(failed(row) for rows in results["files"].values() for row in rows.values())
    )


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:] or list(WORKFLOWS)))
